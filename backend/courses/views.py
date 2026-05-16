from rest_framework.decorators import action, api_view, permission_classes, parser_classes
from rest_framework.permissions import IsAuthenticated, BasePermission
from rest_framework.response import Response
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser
from rest_framework import status
from rest_framework.views import APIView
from rest_framework import viewsets
from django.db import transaction
from django.db.models import Count, Q
from django.shortcuts import get_object_or_404
from django.http import HttpResponse
from django.utils import timezone
from django.core.files.base import ContentFile
import json
import random
import secrets
import csv
import io
import re
import base64
import logging
import unicodedata
from collections import OrderedDict, defaultdict
from datetime import timedelta, datetime
from .serializers import ActivityCommentSerializer, CategorySerializer, CourseSerializer

from .models import (
    ActivityType,
    GradingComponentScore,
    Lesson,
    LessonCompletion,
    LessonImage,
    Module,
    GradingScheme,
    CourseActivity,
    CourseComment,
    InstructorFeedback,
    ActivitySubmission,
    SubmissionAttachment,
    AttendanceSession,
    AttendanceRecord,
    QuizAttempt,
    QuizAttemptAnswer,
    QuizAttemptScoreAudit,
    ClassworkDraft,
    QuestionBankItem,
    QuizSecurityEvent,
    QuizAttemptAcknowledgement,
    EnrollmentRequest,
)
from .serializers import (
    CourseSerializer,
    LessonSerializer,
    ModuleSerializer,
    CourseActivitySerializer,
    CourseCommentSerializer,
    InstructorFeedbackSerializer,
    ActivitySubmissionSerializer,
    AttendanceSessionSerializer,
    AttendanceRecordSerializer,
    QuizAttemptSerializer,
    GradingSchemeSerializer,
    GradeSheetSerializer,
    ClassworkDraftSerializer,
    QuestionBankItemSerializer,
    QuizSecurityEventSerializer,
    MeetingSerializer,
    EnrollmentRequestSerializer,
)
from .services.grading import ACTIVITY_CATEGORY_LABELS, _slugify_identifier, compute_grade_details_for_students
from .services.lesson_extraction import extract_lesson_content
from .services.meetings import create_meeting, join_meeting, list_course_meetings
from .services.module_extraction import extract_module_structure
from users_app.models import AdminLog, Category, Course, User
from users_app.events.registry import dispatch_event

try:
    import docx  # python-docx
except Exception:
    docx = None

try:
    import pdfplumber
except Exception:
    pdfplumber = None

try:
    import xlsxwriter
except Exception:
    xlsxwriter = None

logger = logging.getLogger(__name__)
HEURISTIC_RECOVERY_WARNING = "Recovered questions using heuristic fallback for unnumbered content."
LESSON_ALLOWED_EXTENSIONS = {".pdf", ".doc", ".docx", ".ppt", ".pptx", ".txt"}
LESSON_MAX_FILE_SIZE = 20 * 1024 * 1024  # 20MB
QUIZ_SECURITY_FORCE_SUBMIT_THRESHOLD = 3
QUIZ_SECURITY_INACTIVITY_SECONDS = 300
QUIZ_SECURITY_RACE_RECONCILE_SECONDS = 3
QUIZ_SECURITY_MAX_EVENTS_PER_ATTEMPT = 100


def _build_lesson_serializer_context(request, lessons_queryset=None):
    context = {"request": request}
    user = getattr(request, "user", None)
    if not user or not getattr(user, "is_authenticated", False) or getattr(user, "role", "") != "student":
        return context

    lesson_ids = []
    if lessons_queryset is not None:
        try:
            lesson_ids = list(lessons_queryset.values_list("id", flat=True))
        except Exception:
            lesson_ids = [getattr(item, "id", None) for item in lessons_queryset if getattr(item, "id", None)]

    completion_map = {
        lesson_id: True
        for lesson_id in LessonCompletion.objects.filter(student=user, lesson_id__in=lesson_ids).values_list("lesson_id", flat=True)
    } if lesson_ids else {}
    context["lesson_completion_map"] = completion_map
    return context


class IsInstructorRole(BasePermission):
    def has_permission(self, request, view):
        return bool(
            getattr(request, "user", None)
            and request.user.is_authenticated
            and getattr(request.user, "role", "") in {"instructor", "admin"}
        )


class IsAdminRole(BasePermission):
    def has_permission(self, request, view):
        user = getattr(request, "user", None)
        return bool(
            user
            and user.is_authenticated
            and (
                getattr(user, "role", "") == "admin"
                or getattr(user, "is_staff", False)
                or getattr(user, "is_superuser", False)
            )
        )


class IsStudentRole(BasePermission):
    def has_permission(self, request, view):
        return bool(getattr(request, "user", None) and request.user.is_authenticated and getattr(request.user, "role", "") == "student")


class IsCourseOwner(BasePermission):
    def has_object_permission(self, request, view, obj):
        course = obj if isinstance(obj, Course) else getattr(obj, "course", None)
        return bool(course and getattr(course, "instructor_id", None) == getattr(request.user, "id", None))


class IsCourseEnrolled(BasePermission):
    def has_object_permission(self, request, view, obj):
        course = obj if isinstance(obj, Course) else getattr(obj, "course", None)
        if not course:
            return False
        return bool(course.students.filter(id=getattr(request.user, "id", None)).exists())


class _PermissionContext:
    def __init__(self, user):
        self.user = user


def _allow(permission, user, obj=None):
    request = _PermissionContext(user)
    if not permission.has_permission(request, None):
        return False
    if obj is None:
        return True
    return permission.has_object_permission(request, None, obj)


def _as_bool(value, default=False):
    if isinstance(value, bool):
        return value
    if value in (None, ""):
        return default
    if isinstance(value, (int, float)):
        return bool(value)
    if isinstance(value, str):
        lowered = value.strip().lower()
        if lowered in {"1", "true", "yes", "on"}:
            return True
        if lowered in {"0", "false", "no", "off"}:
            return False
    return bool(value)


def _is_admin_user(user):
    return bool(
        user
        and getattr(user, "is_authenticated", False)
        and (
            getattr(user, "role", "") == "admin"
            or getattr(user, "is_staff", False)
            or getattr(user, "is_superuser", False)
        )
    )


def _is_instructor_user(user):
    return bool(user and getattr(user, "is_authenticated", False) and getattr(user, "role", "") == "instructor")


def _can_manage_course(user, course):
    return bool(course and (_is_admin_user(user) or getattr(course, "instructor_id", None) == getattr(user, "id", None)))


def _can_view_course(user, course):
    if _can_manage_course(user, course):
        return True
    return bool(
        course
        and getattr(user, "role", "") == "student"
        and course.students.filter(id=getattr(user, "id", None)).exists()
    )


def _pending_enrollment_requests_for_user(user):
    queryset = EnrollmentRequest.objects.select_related("course", "student", "reviewed_by").filter(
        status=EnrollmentRequest.STATUS_PENDING
    )
    if _is_admin_user(user):
        return queryset
    return queryset.filter(course__instructor=user)


def _log_admin_action(action, performed_by=None, target_user=None, description=""):
    try:
        AdminLog.objects.create(
            action=action,
            performed_by=performed_by if getattr(performed_by, "is_authenticated", False) else None,
            target_user=target_user,
            description=description or "",
        )
    except Exception:
        logger.warning("Failed to persist admin course action log.", exc_info=True)


INSTRUCTOR_ROLE_PERMISSION = IsInstructorRole()
ADMIN_ROLE_PERMISSION = IsAdminRole()
STUDENT_ROLE_PERMISSION = IsStudentRole()
COURSE_OWNER_PERMISSION = IsCourseOwner()
COURSE_ENROLLED_PERMISSION = IsCourseEnrolled()


def _anti_cheat_runtime_config(activity):
    return {
        "enabled": bool(getattr(activity, "anti_cheat_enabled", False)),
        "warn_on_tab_switch": bool(getattr(activity, "anti_cheat_tab_switch", False)),
        "detect_multiple_tabs": bool(getattr(activity, "anti_cheat_multi_tab", False)),
        "disable_copy_paste": bool(getattr(activity, "anti_cheat_disable_copy_paste", False)),
        "require_fullscreen": bool(getattr(activity, "anti_cheat_fullscreen_required", False)),
        "violation_threshold": QUIZ_SECURITY_FORCE_SUBMIT_THRESHOLD,
        "inactivity_seconds": QUIZ_SECURITY_INACTIVITY_SECONDS,
    }


# -----------------------------
# Category Management
# -----------------------------
@api_view(["GET", "POST"])
@permission_classes([IsAuthenticated])
def category_list_create(request):
    if request.method == "GET":
        categories = Category.objects.order_by("name")
        serializer = CategorySerializer(categories, many=True)
        return Response(serializer.data)

    if not _allow(ADMIN_ROLE_PERMISSION, request.user):
        return Response({"error": "Only admin can manage categories."}, status=403)

    serializer = CategorySerializer(data=request.data)
    if serializer.is_valid():
        category = serializer.save()
        _log_admin_action(
            "Category created",
            performed_by=request.user,
            description=f"Created category '{category.name}'.",
        )
        return Response(serializer.data, status=201)
    return Response(serializer.errors, status=400)


@api_view(["GET", "PUT", "PATCH", "DELETE"])
@permission_classes([IsAuthenticated])
def category_detail(request, category_id):
    category = get_object_or_404(Category, id=category_id)

    if request.method == "GET":
        return Response(CategorySerializer(category).data)

    if not _allow(ADMIN_ROLE_PERMISSION, request.user):
        return Response({"error": "Only admin can manage categories."}, status=403)

    if request.method in {"PUT", "PATCH"}:
        serializer = CategorySerializer(
            category,
            data=request.data,
            partial=request.method == "PATCH",
        )
        if serializer.is_valid():
            serializer.save()
            _log_admin_action(
                "Category updated",
                performed_by=request.user,
                description=f"Updated category '{category.name}'.",
            )
            return Response(serializer.data)
        return Response(serializer.errors, status=400)

    if category.courses.exists():
        return Response({"error": "Cannot delete a category that is assigned to courses."}, status=400)

    deleted_name = category.name
    category.delete()
    _log_admin_action(
        "Category deleted",
        performed_by=request.user,
        description=f"Deleted category '{deleted_name}'.",
    )
    return Response(status=204)


@api_view(["GET", "POST"])
@permission_classes([IsAuthenticated])
@parser_classes([MultiPartParser, FormParser, JSONParser])
def admin_course_management(request):
    if not _is_admin_user(request.user):
        return Response({"error": "Only admin can manage all courses."}, status=403)

    if request.method == "GET":
        courses = (
            Course.objects.select_related("instructor", "category")
            .annotate(students_count=Count("students", distinct=True), lessons_count=Count("lessons", distinct=True))
            .order_by("-id")
        )
        serializer = CourseSerializer(courses, many=True, context={"request": request})
        payload = []
        for item, course in zip(serializer.data, courses):
            item["instructor_name"] = (
                f"{course.instructor.first_name} {course.instructor.last_name}".strip() or course.instructor.username
            )
            item["instructor_id"] = course.instructor_id
            payload.append(item)
        return Response(payload)

    instructor_id = request.data.get("instructor_id")
    instructor = User.objects.filter(id=instructor_id, role="instructor").first()
    if not instructor:
        return Response({"error": "A valid instructor_id is required."}, status=400)

    serializer = CourseSerializer(data=request.data, context={"request": request})
    serializer.is_valid(raise_exception=True)
    course = serializer.save(instructor=instructor)
    _log_admin_action(
        "Course created",
        performed_by=request.user,
        target_user=instructor,
        description=f"Created course '{course.title}' and assigned it to {instructor.username}.",
    )
    data = CourseSerializer(course, context={"request": request}).data
    data["instructor_name"] = f"{instructor.first_name} {instructor.last_name}".strip() or instructor.username
    data["instructor_id"] = instructor.id
    return Response(data, status=201)


@api_view(["GET", "PATCH", "PUT", "DELETE"])
@permission_classes([IsAuthenticated])
@parser_classes([MultiPartParser, FormParser, JSONParser])
def admin_course_detail(request, course_id):
    if not _is_admin_user(request.user):
        return Response({"error": "Only admin can manage all courses."}, status=403)

    course = get_object_or_404(Course.objects.select_related("instructor", "category"), id=course_id)

    if request.method == "GET":
        data = CourseSerializer(course, context={"request": request}).data
        data["instructor_name"] = f"{course.instructor.first_name} {course.instructor.last_name}".strip() or course.instructor.username
        data["instructor_id"] = course.instructor_id
        return Response(data)

    if request.method in {"PATCH", "PUT"}:
        next_instructor_id = request.data.get("instructor_id")
        if next_instructor_id not in (None, ""):
            next_instructor = User.objects.filter(id=next_instructor_id, role="instructor").first()
            if not next_instructor:
                return Response({"error": "Assigned instructor must have instructor role."}, status=400)
        else:
            next_instructor = course.instructor

        serializer = CourseSerializer(
            course,
            data=request.data,
            partial=request.method == "PATCH",
            context={"request": request},
        )
        serializer.is_valid(raise_exception=True)
        updated_course = serializer.save(instructor=next_instructor)
        if "archived" in request.data:
            updated_course.is_archived = _as_bool(request.data.get("archived"))
            updated_course.save(update_fields=["is_archived"])
        _log_admin_action(
            "Course updated",
            performed_by=request.user,
            target_user=next_instructor,
            description=f"Updated course '{updated_course.title}'.",
        )
        data = CourseSerializer(updated_course, context={"request": request}).data
        data["instructor_name"] = f"{next_instructor.first_name} {next_instructor.last_name}".strip() or next_instructor.username
        data["instructor_id"] = next_instructor.id
        return Response(data)

    course_title = course.title
    instructor = course.instructor
    course.delete()
    _log_admin_action(
        "Course deleted",
        performed_by=request.user,
        target_user=instructor,
        description=f"Deleted course '{course_title}'.",
    )
    return Response({"message": "Course deleted successfully."}, status=200)


@api_view(["DELETE"])
@permission_classes([IsAuthenticated])
def admin_content_delete(request, content_type, object_id):
    if not _is_admin_user(request.user):
        return Response({"error": "Only admin can remove content."}, status=403)

    content_map = {
        "course-comment": (CourseComment, "Course comment"),
        "activity-comment": (ActivityComment, "Activity comment"),
        "lesson": (Lesson, "Lesson"),
        "activity": (CourseActivity, "Activity"),
    }
    config = content_map.get(content_type)
    if not config:
        return Response({"error": "Unsupported content type."}, status=400)

    model, label = config
    obj = get_object_or_404(model, id=object_id)
    course = getattr(obj, "course", None) or getattr(getattr(obj, "activity", None), "course", None)
    course_title = getattr(course, "title", "")
    description = f"Removed {label.lower()} #{object_id}"
    if course_title:
        description = f"{description} from course '{course_title}'."
    obj.delete()
    _log_admin_action("Content removed", performed_by=request.user, description=description)
    return Response({"message": f"{label} deleted successfully."}, status=200)


# -----------------------------
# Instructor Courses (GET + POST)
# -----------------------------
@api_view(["GET", "POST"])
@permission_classes([IsAuthenticated])
@parser_classes([MultiPartParser, FormParser])
def instructor_courses_list(request):
    instructor = request.user
    if not getattr(instructor, "is_authenticated", False) or getattr(instructor, "role", "") != "instructor":
        return Response({"error": "Only instructors can manage their own course workspace."}, status=403)

    if request.method == "GET":
        courses = Course.objects.filter(instructor=instructor).select_related("category").annotate(
            students_count=Count("students", distinct=True),
            lessons_count=Count("lessons", distinct=True),
        )
        serializer = CourseSerializer(courses, many=True, context={"request": request})
        return Response(serializer.data)

    elif request.method == "POST":
        serializer = CourseSerializer(data=request.data, context={"request": request})
        if serializer.is_valid():
            serializer.save(instructor=instructor)
            return Response(serializer.data, status=201)
        return Response(serializer.errors, status=400)


# -----------------------------
# Lessons
# -----------------------------
@api_view(["GET"])
@permission_classes([IsAuthenticated])
def lessons_list(request, course_id):
    course = Course.objects.filter(id=course_id).first()
    if not course:
        return Response({"error": "Course not found"}, status=404)
    can_access = _allow(COURSE_OWNER_PERMISSION, request.user, course) or _allow(COURSE_ENROLLED_PERMISSION, request.user, course)
    if not can_access:
        return Response({"error": "Course not found or access denied"}, status=404)

    lessons = Lesson.objects.filter(course=course)
    serializer = LessonSerializer(lessons, many=True, context=_build_lesson_serializer_context(request, lessons))
    return Response(serializer.data)


@api_view(["POST"])
@permission_classes([IsAuthenticated])
@parser_classes([MultiPartParser, FormParser])
def add_lesson(request, course_id):
    if not _is_instructor_user(request.user):
        return Response({"error": "Only instructor can create lessons"}, status=403)

    course = get_object_or_404(Course, id=course_id)
    if not _can_manage_course(request.user, course):
        return Response({"error": "Course not found or access denied"}, status=404)

    upload = request.FILES.get("file")
    extraction_payload = None
    extraction_warning = None
    if upload:
        is_valid, error_message = _validate_lesson_upload(upload)
        if not is_valid:
            return Response({"error": error_message}, status=400)
        # Extraction is best-effort: even when warnings exist, the lesson/file can still be saved.
        extraction_payload = extract_lesson_content(upload)
        if extraction_payload["warnings"]:
            extraction_warning = " ".join(extraction_payload["warnings"])

    data = request.data.copy()
    data["course"] = course_id
    if extraction_payload and not str(data.get("title") or "").strip():
        data["title"] = extraction_payload["title_suggestion"]
    if extraction_payload:
        data["extracted_text"] = extraction_payload["extracted_text"]

    serializer = LessonSerializer(data=data, context={"request": request})
    if serializer.is_valid():
        lesson = serializer.save(file=upload)
        if extraction_payload:
            _save_lesson_images(
                lesson=lesson,
                extracted_images=extraction_payload["images"],
                kept_indexes=_parse_kept_image_indexes(request.data.get("kept_image_indexes")),
            )

        response_data = LessonSerializer(lesson, context={"request": request}).data
        if extraction_warning:
            response_data["warnings"] = [extraction_warning]
        return Response(response_data, status=201)
    return Response(serializer.errors, status=400)


@api_view(["PATCH", "DELETE"])
@permission_classes([IsAuthenticated])
def lesson_detail(request, lesson_id, course_id=None):
    lesson = get_object_or_404(Lesson, id=lesson_id)

    if course_id is not None and lesson.course_id != course_id:
        return Response({"error": "Lesson not found"}, status=404)

    if not _can_manage_course(request.user, lesson.course):
        return Response({"error": "Only instructors or admins can modify lessons"}, status=403)

    if request.method == "PATCH":
        serializer = LessonSerializer(
            lesson,
            data=request.data,
            partial=True,
            context={"request": request},
        )
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data, status=200)
        return Response(serializer.errors, status=400)

    lesson.delete()
    return Response({"message": "Lesson deleted successfully"}, status=200)


@api_view(["DELETE"])
@permission_classes([IsAuthenticated])
def module_detail(request, module_id, course_id=None):
    module = get_object_or_404(Module, id=module_id)

    if course_id is not None and module.course_id != course_id:
        return Response({"error": "Module not found"}, status=404)

    if not _can_manage_course(request.user, module.course):
        return Response({"error": "Only instructors or admins can modify modules"}, status=403)

    module.delete()
    return Response({"message": "Module deleted successfully"}, status=200)


def _get_course_for_module_api(course_id, user):
    course = get_object_or_404(Course, id=course_id)
    if _can_manage_course(user, course):
        return course
    if user.role == "student" and course.students.filter(id=user.id).exists():
        return course
    return None


@api_view(["GET", "POST"])
@permission_classes([IsAuthenticated])
def course_modules(request, course_id):
    course = _get_course_for_module_api(course_id, request.user)
    if not course:
        return Response({"error": "Course not found or access denied"}, status=404)

    if request.method == "GET":
        modules = Module.objects.filter(course=course).prefetch_related("lessons")
        serializer = ModuleSerializer(modules, many=True, context={"request": request})
        return Response(serializer.data)

    if not _can_manage_course(request.user, course):
        return Response({"error": "Only instructors or admins can create modules"}, status=403)

    serializer = ModuleSerializer(data=request.data, context={"request": request})
    if serializer.is_valid():
        module = serializer.save(course=course)
        return Response(ModuleSerializer(module, context={"request": request}).data, status=201)
    return Response(serializer.errors, status=400)


@api_view(["GET", "POST"])
@permission_classes([IsAuthenticated])
@parser_classes([MultiPartParser, FormParser])
def module_lessons(request, module_id):
    module = get_object_or_404(Module, id=module_id)
    course = _get_course_for_module_api(module.course_id, request.user)
    if not course:
        return Response({"error": "Course not found or access denied"}, status=404)

    if request.method == "GET":
        lessons = Lesson.objects.filter(module=module).order_by("order", "id")
        serializer = LessonSerializer(lessons, many=True, context=_build_lesson_serializer_context(request, lessons))
        return Response(serializer.data)

    if not _can_manage_course(request.user, course):
        return Response({"error": "Only instructors or admins can create lessons"}, status=403)

    upload = request.FILES.get("file")
    extraction_payload = None
    extraction_warning = None
    if upload:
        is_valid, error_message = _validate_lesson_upload(upload)
        if not is_valid:
            return Response({"error": error_message}, status=400)
        extraction_payload = extract_lesson_content(upload)
        if extraction_payload["warnings"]:
            extraction_warning = " ".join(extraction_payload["warnings"])

    data = request.data.copy()
    data["course"] = course.id
    data["module"] = module.id
    if extraction_payload and not str(data.get("title") or "").strip():
        data["title"] = extraction_payload["title_suggestion"]
    if extraction_payload:
        data["extracted_text"] = extraction_payload["extracted_text"]
    serializer = LessonSerializer(data=data, context={"request": request})
    if serializer.is_valid():
        lesson = serializer.save(course=course, module=module, file=upload)
        if extraction_payload:
            _save_lesson_images(
                lesson=lesson,
                extracted_images=extraction_payload["images"],
                kept_indexes=_parse_kept_image_indexes(request.data.get("kept_image_indexes")),
            )

        response_data = LessonSerializer(lesson, context={"request": request}).data
        if extraction_warning:
            response_data["warnings"] = [extraction_warning]
        return Response(response_data, status=201)
    return Response(serializer.errors, status=400)


def _validate_lesson_upload(upload):
    filename = str(getattr(upload, "name", "") or "")
    extension = filename[filename.rfind(".") :].lower() if "." in filename else ""
    if extension not in LESSON_ALLOWED_EXTENSIONS:
        return False, "Unsupported file type. Allowed: PDF, DOC/DOCX, PPT/PPTX, TXT."
    size = int(getattr(upload, "size", 0) or 0)
    if size > LESSON_MAX_FILE_SIZE:
        return False, "File is too large. Maximum size is 20MB."
    return True, ""


def _parse_kept_image_indexes(raw_value):
    if raw_value in (None, "", []):
        return None
    if isinstance(raw_value, list):
        raw_items = raw_value
    else:
        try:
            parsed = json.loads(raw_value)
            raw_items = parsed if isinstance(parsed, list) else []
        except Exception:
            raw_items = []
    cleaned = set()
    for item in raw_items:
        try:
            cleaned.add(int(item))
        except (TypeError, ValueError):
            continue
    return cleaned if cleaned else None


def _save_lesson_images(lesson, extracted_images, kept_indexes=None):
    # Persist only selected extracted images so instructors can remove unwanted previews before final save.
    for index, image_payload in enumerate(extracted_images):
        if kept_indexes is not None and index not in kept_indexes:
            continue
        blob = image_payload.get("bytes")
        if not blob:
            continue
        filename = image_payload.get("name") or f"lesson_{lesson.id}_{index + 1}.png"
        image = LessonImage(lesson=lesson)
        image.image.save(filename, ContentFile(blob), save=True)


@api_view(["POST"])
@permission_classes([IsAuthenticated])
@parser_classes([MultiPartParser, FormParser])
def extract_lesson_file_preview(request, course_id):
    course = _get_course_for_module_api(course_id, request.user)
    if not course:
        return Response({"error": "Course not found or access denied"}, status=404)
    if not _can_manage_course(request.user, course):
        return Response({"error": "Only instructors or admins can extract lesson files"}, status=403)

    upload = request.FILES.get("file")
    if not upload:
        return Response({"error": "file is required"}, status=400)

    is_valid, error_message = _validate_lesson_upload(upload)
    if not is_valid:
        return Response({"error": error_message}, status=400)

    extraction_payload = extract_lesson_content(upload)
    images = []
    for index, payload in enumerate(extraction_payload["images"]):
        blob = payload.get("bytes")
        if not blob:
            continue
        encoded = base64.b64encode(blob).decode("ascii")
        images.append(
            {
                "index": index,
                "name": payload.get("name") or f"image_{index + 1}.png",
                "data_url": f"data:image/png;base64,{encoded}",
            }
        )

    return Response(
        {
            "title_suggestion": extraction_payload["title_suggestion"],
            "extracted_text": extraction_payload["extracted_text"],
            "images": images,
            "warnings": extraction_payload["warnings"],
        },
        status=200,
    )


def _parse_import_lessons(raw_value):
    if raw_value in (None, "", []):
        return []
    try:
        payload = json.loads(raw_value) if isinstance(raw_value, str) else raw_value
    except Exception:
        return []
    if not isinstance(payload, list):
        return []

    cleaned = []
    for index, item in enumerate(payload, start=1):
        if not isinstance(item, dict):
            continue
        title = str(item.get("title") or "").strip() or f"Lesson {index}"
        content = str(item.get("content") or "").strip()
        lesson_type = str(item.get("type") or "paragraph").strip().lower()
        if lesson_type not in {"slide", "section", "paragraph"}:
            lesson_type = "paragraph"
        if not title and not content:
            continue
        cleaned.append({"title": title, "content": content, "type": lesson_type})
    return cleaned


@api_view(["POST"])
@permission_classes([IsAuthenticated])
@parser_classes([MultiPartParser, FormParser])
def import_module_from_file(request, course_id):
    """
    Analyze + create pipeline for module import.
    - analyze_only=true: parse file and return structured lesson candidates
    - analyze_only=false: create module + lessons from edited lesson list
    """
    course = _get_course_for_module_api(course_id, request.user)
    if not course:
        return Response({"error": "Course not found or access denied"}, status=404)
    if not _can_manage_course(request.user, course):
        return Response({"error": "Only instructors or admins can import modules"}, status=403)

    upload = request.FILES.get("file")
    if not upload:
        return Response({"error": "file is required"}, status=400)
    is_valid, error_message = _validate_lesson_upload(upload)
    if not is_valid:
        return Response({"error": error_message}, status=400)

    analyze_only_raw = str(request.data.get("analyze_only", "true")).strip().lower()
    analyze_only = analyze_only_raw in {"1", "true", "yes"}
    module_title = str(request.data.get("module_title") or "").strip()

    structure = extract_module_structure(upload, module_title_override=module_title)
    if analyze_only:
        return Response(structure, status=200)

    edited_lessons = _parse_import_lessons(request.data.get("lessons"))
    lesson_rows = edited_lessons if edited_lessons else structure["lessons"]
    if not lesson_rows:
        return Response(
            {
                "error": "No lessons available to import.",
                "warnings": structure.get("warnings", []),
            },
            status=400,
        )

    next_order = Module.objects.filter(course=course).count() + 1
    module = Module.objects.create(
        course=course,
        title=structure["module_title"] or "Imported Module",
        order=next_order,
    )

    created_lessons = []
    for index, item in enumerate(lesson_rows, start=1):
        lesson = Lesson.objects.create(
            course=course,
            module=module,
            title=str(item.get("title") or f"Lesson {index}").strip(),
            content=str(item.get("content") or "").strip(),
            description=str(item.get("content") or "").strip(),
            order=index,
        )
        created_lessons.append(lesson)

    module_data = ModuleSerializer(module, context={"request": request}).data
    lessons_data = LessonSerializer(created_lessons, many=True, context={"request": request}).data
    return Response(
        {
            "module": module_data,
            "lessons": lessons_data,
            "warnings": structure.get("warnings", []),
        },
        status=201,
    )


# -----------------------------
# Activities
# -----------------------------
@api_view(["GET"])
@permission_classes([IsAuthenticated])
def activities_list(request, course_id):
    activities = CourseActivity.objects.filter(course_id=course_id)
    serializer = CourseActivitySerializer(activities, many=True, context={"request": request})
    return Response(serializer.data)


@api_view(["POST"])
@permission_classes([IsAuthenticated])
@parser_classes([MultiPartParser, FormParser])
def add_activity(request, course_id):
    data = request.data.copy()
    data["course"] = course_id
    serializer = CourseActivitySerializer(data=data, context={"request": request})
    if serializer.is_valid():
        serializer.save()
        return Response(serializer.data, status=201)
    else:
        logger.warning(
            "Activity serializer validation failed.",
            extra={"course_id": course_id, "user_id": getattr(request.user, "id", None)},
        )
        return Response(serializer.errors, status=400)


# -----------------------------
# Comments
# -----------------------------
@api_view(["GET"])
@permission_classes([IsAuthenticated])
def comments_list(request, course_id):
    comments = CourseComment.objects.filter(course_id=course_id)
    serializer = CourseCommentSerializer(comments, many=True, context={"request": request})
    return Response(serializer.data)


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def add_comment(request, course_id):
    data = request.data.copy()
    data["course"] = course_id
    data["user"] = request.user.id
    serializer = CourseCommentSerializer(data=data, context={"request": request})
    if serializer.is_valid():
        serializer.save()
        return Response(serializer.data, status=201)
    return Response(serializer.errors, status=400)


# -----------------------------
# Instructor Feedback
# -----------------------------
@api_view(["GET"])
@permission_classes([IsAuthenticated])
def feedback_list(request, course_id):
    feedbacks = InstructorFeedback.objects.filter(course_id=course_id)
    serializer = InstructorFeedbackSerializer(feedbacks, many=True, context={"request": request})
    return Response(serializer.data)


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def leave_feedback(request, course_id):
    data = request.data.copy()
    data["course"] = course_id
    data["student"] = request.user.id
    serializer = InstructorFeedbackSerializer(data=data, context={"request": request})
    if serializer.is_valid():
        serializer.save()
        return Response(serializer.data, status=201)
    return Response(serializer.errors, status=400)


# -----------------------------
# View Course Detail
# -----------------------------
@api_view(["GET"])
@permission_classes([IsAuthenticated])
def course_detail(request, course_id):
    try:
        course = Course.objects.select_related("category", "instructor").get(id=course_id)
    except Course.DoesNotExist:
        return Response({"error": "Course not found"}, status=404)
    can_access = _can_view_course(request.user, course)
    if not can_access:
        return Response({"error": "Course not found or access denied"}, status=404)

    students_qs = course.students.filter(role="student")
    lessons_count = Lesson.objects.filter(course=course).count()
    modules_count = Module.objects.filter(course=course).count()
    attendance_sessions_count = AttendanceSession.objects.filter(course=course).count()
    average_grade = None
    grade_map = compute_grade_details_for_students(course, list(students_qs))
    valid_grades = []
    for student in students_qs:
        details = grade_map.get(student.id, {})
        if details.get("error"):
            continue
        try:
            valid_grades.append(float(details.get("final_grade", 0.0)))
        except (TypeError, ValueError):
            continue
    if valid_grades:
        average_grade = round(sum(valid_grades) / len(valid_grades), 2)

    # Prepare data with instructor + students + course stats
    data = {
        "id": course.id,
        "title": course.title,
        "description": course.description,
        "category": (
            {"id": course.category.id, "name": course.category.name}
            if course.category_id
            else None
        ),
        "thumbnail": request.build_absolute_uri(course.thumbnail.url) if course.thumbnail else None,
        "start_date": course.start_date,
        "start_time": course.start_time,
        "scheduled_start_at": course.get_start_datetime(),
        "is_archived": course.is_archived,
        "status": course.get_status(),
        "state": course.get_status(),
        "code": course.join_code,
        "join_code": course.join_code,
        "join_code_enabled": course.join_code_enabled,
        "join_code_expiration": course.join_code_expiration,
        "modules_count": modules_count,
        "lessons_count": lessons_count,
        "students_count": students_qs.count(),
        "pending_enrollment_requests_count": course.enrollment_requests.filter(
            status=EnrollmentRequest.STATUS_PENDING
        ).count(),
        "attendance_sessions_count": attendance_sessions_count,
        "average_grade": average_grade,
        "instructor": {
            "id": course.instructor.id,
            "username": course.instructor.username,
            "role": "instructor",
            "school_id": course.instructor.school_id
        },
        "students": [
            {"id": s.id, "username": s.username, "role": "student", "school_id": s.school_id}
            for s in students_qs
        ]
    }
    return Response(data)


# -----------------------------
# Students & Instructor
# -----------------------------
@api_view(["POST"])
@permission_classes([IsAuthenticated])
def add_student(request, course_id):

    try:
        course = Course.objects.get(id=course_id)
    except Course.DoesNotExist:
        return Response(
            {"error": "Course not found or access denied"},
            status=404
        )
    if not _can_manage_course(request.user, course):
        return Response({"error": "Course not found or access denied"}, status=404)

    student_code = request.data.get("student_id")

    if not student_code:
        return Response({"error": "No student ID provided"}, status=400)

    try:
        student = User.objects.get(
            school_id=student_code,
            role="student"
        )
    except User.DoesNotExist:
        return Response({"error": "Student not found"}, status=404)

    course.students.add(student)
    dispatch_event("student_added_to_course", course=course, student=student, actor=request.user)

    return Response({"message": "Student added successfully!"})



@api_view(["GET"])
@permission_classes([IsAuthenticated])
def students_list(request, course_id):
    try:
        course = Course.objects.get(id=course_id)
    except Course.DoesNotExist:
        return Response({"error": "Course not found"}, status=404)
    can_access = _can_view_course(request.user, course)
    if not can_access:
        return Response({"error": "Course not found or access denied"}, status=404)

    students = course.students.all()
    instructor = course.instructor

    data = []

    # Add instructor first
    data.append({
        "id": instructor.id,
        "username": instructor.username,
        "school_id": instructor.school_id,
        "role": "instructor"
    })

    # Add students
    for s in students:
        data.append({
            "id": s.id,
            "username": s.username,
            "school_id": s.school_id,
            "role": "student"
        })

    return Response(data)


@api_view(["GET", "PUT", "DELETE"])
@permission_classes([IsAuthenticated])
@parser_classes([MultiPartParser, FormParser])
def activity_detail(request, course_id, activity_id):
    try:
        activity = CourseActivity.objects.get(id=activity_id, course_id=course_id)
    except CourseActivity.DoesNotExist:
        return Response({"error": "Activity not found"}, status=404)

    can_access = _allow(COURSE_OWNER_PERMISSION, request.user, activity) or _allow(COURSE_ENROLLED_PERMISSION, request.user, activity)
    if request.method == "GET":
        if not can_access:
            return Response({"error": "Activity not found or access denied"}, status=404)
        serializer = CourseActivitySerializer(activity, context={"request": request})
        return Response(serializer.data)

    if not _allow(COURSE_OWNER_PERMISSION, request.user, activity):
        return Response({"error": "Only instructor can modify activities"}, status=403)

    elif request.method in ["PUT", "PATCH"]:
        serializer = CourseActivitySerializer(activity, data=request.data, partial=True, context={"request": request})
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data)
        return Response(serializer.errors, status=400)

    elif request.method == "DELETE":
        activity.delete()
        return Response({"message": "Activity deleted!"}, status=204)


def _quiz_type():
    return ActivityType.objects.filter(name__iexact="quiz").first()


def _is_quiz_or_exam_activity(activity):
    return str(getattr(activity.activity_type, "name", "") or "").lower() == "quiz"


SECTION_TYPE_KEYWORDS = {
    "multiple_choice": ["multiple choice", "mcq", "choose the best answer"],
    "true_false": ["true or false", "true/false", "t/f"],
    "identification": ["identification", "identify"],
    "short_answer": ["short answer", "brief answer"],
    "enumeration": ["enumeration", "enumerate", "list down", "listing"],
    "coding": ["coding", "programming", "write a program", "code"],
    "essay": ["essay", "long answer", "short answer", "free response"],
    "matching": ["matching", "match the following", "match type"],
    "file_upload": ["file upload", "upload", "attachment"],
}


def _normalize_line(line):
    normalized = re.sub(r"[\u2013\u2014]", "-", str(line or ""))
    normalized = re.sub(r"\s+", " ", normalized).strip()
    return normalized


def normalize_text(raw_text):
    """
    Preprocess raw document text to improve parser tolerance:
    - remove common emoji/symbol ranges
    - normalize tabs/newlines
    - trim each line
    """
    text = str(raw_text or "")
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\t", " ")
    text = re.sub(r"[\U0001F300-\U0001FAFF\u2600-\u27BF]+", "", text)
    cleaned_chars = []
    for char in text:
        category = unicodedata.category(char)
        if category in {"So", "Sk", "Cs"} and ord(char) > 127:
            continue
        cleaned_chars.append(char)
    text = "".join(cleaned_chars)

    normalized_lines = []
    for line in text.split("\n"):
        line = _normalize_line(line)
        normalized_lines.append(line)
    return "\n".join(normalized_lines)


def _line_for_matching(line):
    cleaned = _normalize_line(line).lower()
    cleaned = re.sub(r"[^a-z0-9\s:/\-\.\)\(]", " ", cleaned)
    cleaned = re.sub(r"\s+", " ", cleaned).strip()
    return cleaned


def _map_section_type(title):
    value = _line_for_matching(title)
    for section_type, keywords in SECTION_TYPE_KEYWORDS.items():
        if any(keyword in value for keyword in keywords):
            return section_type
    return "identification"


def _is_answer_key_header(line):
    value = _line_for_matching(line)
    return bool(re.match(r"^(answer\s*key|answers|key)\b", value))


def split_document(raw_text):
    cleaned_text = normalize_text(raw_text)
    raw_lines = [str(line or "").rstrip() for line in cleaned_text.splitlines()]
    answer_key_index = None
    for index, line in enumerate(raw_lines):
        if _is_answer_key_header(line):
            answer_key_index = index
            break

    if answer_key_index is None:
        exam_lines = [line for line in raw_lines if _normalize_line(line)]
        answer_lines = []
    else:
        exam_lines = [line for line in raw_lines[:answer_key_index] if _normalize_line(line)]
        answer_lines = [line for line in raw_lines[answer_key_index + 1 :] if _normalize_line(line)]

    return exam_lines, answer_lines


def _is_question_start(line):
    text = _normalize_line(line)
    return bool(
        re.match(
            r"^(?:q\s*)?\d{1,4}(?:\s*[\.\)\-:]\s*|\s+)\S+",
            text,
            flags=re.IGNORECASE,
        )
    )


def _extract_question_number_and_text(line):
    text = _normalize_line(line)
    match = re.match(
        r"^(?:q\s*)?(\d{1,4})(?:\s*[\.\)\-:]\s*|\s+)(.+)$",
        text,
        flags=re.IGNORECASE,
    )
    if not match:
        return None, text
    return int(match.group(1)), match.group(2).strip()


def _extract_inline_choices(line):
    matches = re.findall(r"([A-Ha-h])[\)\.\:\-]\s*([^A-H\n\r]+?)(?=(?:\s+[A-Ha-h][\)\.\:\-])|$)", line)
    return [(letter.upper(), _normalize_line(text)) for letter, text in matches if _normalize_line(text)]


def _extract_vertical_choice(line):
    match = re.match(r"^([A-Ha-h])[\)\.\:\-]\s*(.+)$", _normalize_line(line))
    if not match:
        return None
    return match.group(1).upper(), _normalize_line(match.group(2))


def _extract_mcq_choices(line):
    normalized = _normalize_line(line)
    marker_pattern = re.compile(r"(^|\s)([A-Da-d])[\.\)]\s*")
    markers = list(marker_pattern.finditer(normalized))
    if not markers:
        return []

    parsed = []
    for idx, match in enumerate(markers):
        letter = match.group(2).upper()
        start = match.end()
        end = markers[idx + 1].start() if idx + 1 < len(markers) else len(normalized)
        value = _normalize_line(normalized[start:end]).rstrip(",")
        if value:
            parsed.append((letter, value))
    return parsed if len(parsed) >= 2 else []


def _normalize_true_false_answer(raw_value):
    value = _normalize_line(raw_value).lower()
    if value in {"t", "true", "1", "yes", "y", "✓", "check"}:
        return "TRUE"
    if value in {"f", "false", "0", "no", "n", "✗", "x"}:
        return "FALSE"
    return ""


def _has_explicit_true_false_marker(text):
    # Enhancement:
    # Require explicit TF cues to avoid accidental misclassification in normal sentences.
    value = _normalize_line(text).lower()
    return bool(
        re.search(r"\(\s*(?:t\s*/\s*f|true\s*/\s*false)\s*\)", value)
        or re.search(r"\btrue\s+or\s+false\b", value)
        or re.search(r"^\s*(?:t\s*/\s*f|true\s*/\s*false)\s*[:\-]?\s*$", value)
    )


def _normalize_answer_section(section_name):
    normalized = _line_for_matching(section_name).replace(" ", "")
    if "multiplechoice" in normalized:
        return "multiple_choice"
    if "true/false" in normalized or "truefalse" in normalized or "trueorfalse" in normalized:
        return "true_false"
    if "shortanswer" in normalized:
        return "short_answer"
    if "enumeration" in normalized:
        return "enumeration"
    if "coding" in normalized:
        return "coding"
    if "essay" in normalized:
        return "essay"
    if "matching" in normalized:
        return "matching"
    if "fileupload" in normalized:
        return "file_upload"
    return "identification"


def _route_question_type(section_type, question_text, has_options=False):
    text = _normalize_line(question_text).lower()
    routed = str(section_type or "").strip().lower()
    if routed in {"mcq", "multiple choice"}:
        routed = "multiple_choice"
    if routed in {"truefalse", "tf"}:
        routed = "true_false"
    if routed in {"short", "short answer"}:
        routed = "short_answer"
    if routed in {"matching_type"}:
        routed = "matching"
    if not routed:
        routed = "identification"

    if _has_explicit_true_false_marker(text):
        return "true_false"
    if re.match(r"^(explain|describe|discuss)\b", text):
        return "essay"
    if "match the following" in text or "matching type" in text:
        return "matching"
    if "upload" in text and "file" in text:
        return "file_upload"
    # Only infer multiple-choice from options when the type was not already
    # explicitly identified as another structured question type.
    if has_options and routed in {"", "multiple_choice"}:
        return "multiple_choice"
    if routed not in {
        "multiple_choice",
        "true_false",
        "short_answer",
        "identification",
        "enumeration",
        "essay",
        "coding",
        "matching",
        "file_upload",
    }:
        return "identification"
    return routed


def _is_instruction_line(line):
    normalized = _line_for_matching(line)
    return bool(
        re.match(
            r"^(write|choose|select|give|direction|directions|instruction|instructions|for teacher use)\b",
            normalized,
            flags=re.IGNORECASE,
        )
    )


def _extract_docx_text(document):
    """
    Build text while preserving simple ordered-list numbering from docx paragraphs.
    This helps avoid losing question numbers when Word auto-numbering is used.
    """
    lines = []
    counters = defaultdict(int)

    def _paragraph_number_prefix(paragraph):
        p = getattr(paragraph, "_p", None)
        if p is None:
            return ""
        p_pr = getattr(p, "pPr", None)
        if p_pr is None:
            return ""
        num_pr = getattr(p_pr, "numPr", None)
        if num_pr is None:
            return ""
        num_id = getattr(getattr(num_pr, "numId", None), "val", None)
        if num_id is None:
            return ""
        ilvl = getattr(getattr(num_pr, "ilvl", None), "val", 0)
        try:
            level = int(ilvl)
        except Exception:
            level = 0
        key = (str(num_id), level)
        counters[key] += 1
        # Reset deeper levels for the same numbering group.
        for counter_key in list(counters.keys()):
            if counter_key[0] == str(num_id) and counter_key[1] > level:
                counters[counter_key] = 0
        return f"{counters[key]}."

    for paragraph in document.paragraphs:
        raw_text = str(paragraph.text or "").strip()
        prefix = _paragraph_number_prefix(paragraph)
        if not raw_text and not prefix:
            continue
        if prefix and raw_text and not re.match(r"^(?:q\s*)?\d{1,3}\s*[\.\)\-:]", raw_text, flags=re.IGNORECASE):
            line = f"{prefix} {raw_text}"
        elif prefix and not raw_text:
            line = prefix
        else:
            line = raw_text
        lines.append(line)
    return "\n".join(lines)


def parse_sections(lines):
    markers = []
    section_name_pattern = r"(identification|short\s*answer|multiple\s*choice|true\s*(?:/|\s+or\s+)?\s*false|enumeration|coding|essay|matching|file\s*upload)"
    section_header_pattern = re.compile(
        rf"^(?:(part|section)\s*([ivxlcdm]+|\d+|[a-z])\s*[:\-]?\s*)?({section_name_pattern})?(?:\s*\(.*\))?\s*:?\s*$",
        flags=re.IGNORECASE,
    )
    mcq_choice_pattern = re.compile(r"^[A-Da-d][\.\)]\s+")
    pipe_question_pattern = re.compile(r"^\d+\|")

    for index, raw_line in enumerate(lines):
        line = _normalize_line(raw_line)
        if not line or _is_question_start(line):
            continue
        if pipe_question_pattern.match(line):
            continue
        if mcq_choice_pattern.match(line):
            logger.debug("Skipped section detection for MCQ choice: %s", line)
            continue
        if _is_answer_key_header(line):
            continue

        lower = _line_for_matching(line)
        direct_type = None
        for section_type, keywords in SECTION_TYPE_KEYWORDS.items():
            if any(keyword in lower for keyword in keywords):
                direct_type = section_type
                break

        match = section_header_pattern.match(line)
        if match and (match.group(1) or match.group(3) or direct_type):
            named_section = _normalize_line(match.group(3) or "")
            inferred = _map_section_type(named_section or line)
            markers.append({"start": index, "title": line, "type": inferred})
            continue

        if direct_type and len(lower.split()) <= 8:
            markers.append({"start": index, "title": line, "type": direct_type})

    markers = sorted(markers, key=lambda item: item["start"])
    deduped = []
    seen = set()
    for marker in markers:
        key = (marker["start"], marker["title"].lower())
        if key in seen:
            continue
        seen.add(key)
        deduped.append(marker)

    if not deduped or deduped[0]["start"] > 0:
        deduped.insert(0, {"start": 0, "title": "General Section", "type": "identification"})

    logger.debug("Parser PASS1 sections detected: %s", deduped)
    return deduped


def _parse_question_block(question_number, first_line_text, block_lines, section_type, warnings):
    body_parts = [first_line_text]
    choices_by_letter = {}
    inline_answer = ""
    starter_code_parts = []
    expected_output = ""
    test_cases = ""
    in_code_block = False

    initial_inline_choices = _extract_inline_choices(first_line_text)
    if len(initial_inline_choices) >= 2:
        cut_index = re.search(r"[A-Ha-h][\)\.\:\-]", first_line_text)
        if cut_index:
            body_parts = [first_line_text[: cut_index.start()].strip()]
        for letter, text in initial_inline_choices:
            choices_by_letter[letter] = text

    for raw_line in block_lines:
        line = _normalize_line(raw_line)
        if not line:
            continue

        answer_match = re.match(r"^(answer|ans|key)\s*[:\-]\s*(.+)$", line, flags=re.IGNORECASE)
        if answer_match:
            inline_answer = answer_match.group(2).strip()
            continue

        if re.match(r"^(starter\s*code|code)\s*[:\-]?", line, flags=re.IGNORECASE):
            in_code_block = True
            cleaned = re.sub(r"^(starter\s*code|code)\s*[:\-]?\s*", "", line, flags=re.IGNORECASE).strip()
            if cleaned:
                starter_code_parts.append(cleaned)
            continue
        if re.match(r"^expected\s*output\s*[:\-]", line, flags=re.IGNORECASE):
            expected_output = re.sub(r"^expected\s*output\s*[:\-]\s*", "", line, flags=re.IGNORECASE).strip()
            continue
        if re.match(r"^test\s*cases?\s*[:\-]", line, flags=re.IGNORECASE):
            test_cases = re.sub(r"^test\s*cases?\s*[:\-]\s*", "", line, flags=re.IGNORECASE).strip()
            continue
        if line.startswith("```"):
            in_code_block = not in_code_block
            continue
        if in_code_block:
            starter_code_parts.append(raw_line.rstrip())
            continue

        vertical_choice = _extract_vertical_choice(line)
        if vertical_choice:
            choices_by_letter[vertical_choice[0]] = vertical_choice[1]
            continue

        inline_choices = _extract_inline_choices(line)
        if len(inline_choices) >= 2:
            for letter, text in inline_choices:
                choices_by_letter[letter] = text
            continue

        if re.match(r"^[_\-\=\.\s]{3,}$", line):
            continue

        body_parts.append(line)

    question_text = _normalize_line(" ".join([part for part in body_parts if part]))
    choices = [choices_by_letter[key] for key in sorted(choices_by_letter.keys()) if choices_by_letter.get(key)]
    detected_type = section_type
    lowered = question_text.lower()
    if section_type == "identification":
        if len(choices) >= 2:
            detected_type = "multiple_choice"
        elif "true or false" in lowered or "true/false" in lowered:
            detected_type = "true_false"
    if detected_type == "true_false" and not choices:
        choices = ["True", "False"]

    normalized_answer = inline_answer.strip()
    answer_key = normalized_answer
    if detected_type == "multiple_choice" and re.match(r"^[A-Ha-h]$", normalized_answer):
        letter = normalized_answer.upper()
        answer_key = letter
        option_index = ord(letter) - ord("A")
        if 0 <= option_index < len(choices):
            normalized_answer = choices[option_index]

    acceptable_answers = []
    if detected_type == "enumeration" and normalized_answer:
        acceptable_answers = [item.strip() for item in re.split(r"[;,/]", normalized_answer) if item.strip()]
        if not acceptable_answers:
            acceptable_answers = [normalized_answer]

    if not question_text:
        warnings.append(f"Question {question_number}: missing question text.")

    question = {
        "id": question_number,
        "number": question_number,
        "question": question_text,
        "question_text": question_text,
        "type": detected_type,
        "choices": choices,
        "options": [{"id": idx + 1, "text": choice} for idx, choice in enumerate(choices)],
        "correct_answer": normalized_answer,
        "answer_key": answer_key,
        "acceptable_answers": acceptable_answers,
        "starter_code": "\n".join(starter_code_parts).strip(),
        "expected_output": expected_output,
        "test_cases": test_cases,
        "points": 1,
    }
    return question


def _coerce_choice_values(raw_choices, raw_options):
    normalized = []
    seen = set()

    if isinstance(raw_choices, list):
        for item in raw_choices:
            value = _normalize_line(item)
            value = value.rstrip(",")
            if not value:
                continue
            key = value.lower()
            if key in seen:
                continue
            seen.add(key)
            normalized.append(value)

    if isinstance(raw_options, list):
        for option in raw_options:
            value = ""
            if isinstance(option, dict):
                value = _normalize_line(option.get("text"))
            else:
                value = _normalize_line(option)
            value = value.rstrip(",")
            if not value:
                continue
            key = value.lower()
            if key in seen:
                continue
            seen.add(key)
            normalized.append(value)

    return normalized


def _normalize_question_payload(question, section_type, answer_map, inline_answer_map, warnings):
    if not isinstance(question, dict):
        warnings.append("Skipped malformed question payload (expected object).")
        return None

    number = question.get("number")
    try:
        number = int(number)
    except Exception:
        number = None

    question_text = _normalize_line(question.get("question_text") or question.get("question") or "")
    if not question_text:
        warnings.append(f"Skipped question with missing text (number={number or 'unknown'}).")
        return None

    q_type = _route_question_type(
        _normalize_line(question.get("type") or "").lower().replace(" ", "_") or section_type,
        question_text,
        has_options=bool(question.get("choices") or question.get("options")),
    )

    choices = _coerce_choice_values(question.get("choices"), question.get("options"))
    if q_type == "multiple_choice" and len(choices) < 2:
        inline_choices = _extract_mcq_choices(question_text)
        if len(inline_choices) >= 2:
            choices = [text for _, text in inline_choices]

    if q_type == "true_false":
        choices = ["True", "False"]

    options = [{"id": idx + 1, "text": text} for idx, text in enumerate(choices)]

    number_key = str(number) if number is not None else ""
    raw_answer = (
        question.get("correct_answer")
        or question.get("answer_key")
        or answer_map.get(number_key, "")
        or inline_answer_map.get(number_key, "")
    )
    raw_answer = str(raw_answer or "").strip()

    normalized_question = dict(question)
    normalized_question["number"] = number if number is not None else question.get("number")
    normalized_question["id"] = number if number is not None else question.get("id")
    normalized_question["question"] = question_text
    normalized_question["question_text"] = question_text
    normalized_question["type"] = q_type
    normalized_question["choices"] = choices
    normalized_question["options"] = options
    normalized_question.setdefault("answer_key", "")
    normalized_question.setdefault("acceptable_answers", [])
    normalized_question.setdefault("starter_code", "")
    normalized_question.setdefault("expected_output", "")
    normalized_question.setdefault("test_cases", "")
    normalized_question["points"] = question.get("points", 1)

    if q_type == "multiple_choice":
        if re.match(r"^[A-Ha-h]$", raw_answer):
            letter = raw_answer.upper()
            idx = ord(letter) - ord("A")
            normalized_question["answer_key"] = letter
            normalized_question["correct_answer"] = letter
            if 0 <= idx < len(choices):
                normalized_question["correct_answer_text"] = choices[idx]
        elif raw_answer.isdigit() and choices:
            idx = int(raw_answer) - 1
            if 0 <= idx < len(choices):
                letter = chr(ord("A") + idx)
                normalized_question["answer_key"] = letter
                normalized_question["correct_answer"] = letter
                normalized_question["correct_answer_text"] = choices[idx]
        elif raw_answer:
            normalized_question["correct_answer"] = raw_answer
            for idx, choice in enumerate(choices):
                if choice.strip().lower() == raw_answer.strip().lower():
                    letter = chr(ord("A") + idx)
                    normalized_question["answer_key"] = letter
                    normalized_question["correct_answer_text"] = choice
                    normalized_question["correct_answer"] = letter
                    break
        else:
            normalized_question["correct_answer"] = ""
    elif q_type == "true_false":
        normalized_question["correct_answer"] = _normalize_true_false_answer(raw_answer)
    elif q_type == "enumeration":
        normalized_question["correct_answer"] = raw_answer
        if raw_answer:
            normalized_question["acceptable_answers"] = [part.strip() for part in re.split(r"[;,/]", raw_answer) if part.strip()]
    elif q_type == "short_answer":
        normalized_question["correct_answer"] = raw_answer
        if raw_answer:
            normalized_question["acceptable_answers"] = [part.strip() for part in re.split(r"[|;\n]+", raw_answer) if part.strip()]
    elif q_type == "identification":
        normalized_question["correct_answer"] = raw_answer
        if raw_answer:
            normalized_question["acceptable_answers"] = [part.strip() for part in re.split(r"[|,;\n]+", raw_answer) if part.strip()]
    elif q_type == "matching":
        normalized_question["correct_answer"] = raw_answer
        normalized_question["matching_pairs"] = _parse_matching_pairs(question.get("matching_pairs") or raw_answer)
        if not normalized_question["matching_pairs"] and choices:
            pair_map = {}
            for idx, value in enumerate(choices):
                pair_map[str(idx + 1)] = value
            normalized_question["matching_pairs"] = pair_map
    elif q_type == "file_upload":
        normalized_question["correct_answer"] = ""
        normalized_question["options"] = []
        normalized_question["choices"] = []
    elif q_type in {"coding", "essay"}:
        normalized_question["correct_answer"] = raw_answer
        if raw_answer and not normalized_question.get("expected_output"):
            normalized_question["expected_output"] = raw_answer
    else:
        normalized_question["correct_answer"] = raw_answer

    return normalized_question


def _normalize_sections_payload(sections, answer_map=None, inline_answer_map=None):
    warnings = []
    answer_map = answer_map or {}
    inline_answer_map = inline_answer_map or {}
    normalized_sections = []

    if not isinstance(sections, list):
        return [], ["Malformed import payload: sections must be a list."]

    for section_index, section in enumerate(sections, start=1):
        if not isinstance(section, dict):
            warnings.append(f"Skipped malformed section payload at index {section_index}.")
            continue

        title = _normalize_line(section.get("title") or section.get("name") or f"Section {section_index}")
        s_type = _normalize_line(section.get("type") or "").lower().replace(" ", "_")
        if not s_type or s_type not in SECTION_TYPE_KEYWORDS:
            s_type = _map_section_type(title)

        raw_questions = section.get("questions")
        if not isinstance(raw_questions, list):
            raw_questions = []
            warnings.append(f"Section '{title}' has malformed questions payload; treated as empty.")

        normalized_questions = []
        for raw_question in raw_questions:
            normalized_question = _normalize_question_payload(raw_question, s_type, answer_map, inline_answer_map, warnings)
            if normalized_question is not None:
                normalized_questions.append(normalized_question)

        if not normalized_questions:
            continue

        normalized_sections.append(
            {
                "id": section.get("id", len(normalized_sections) + 1),
                "title": title,
                "type": s_type,
                "instructions": section.get("instructions", ""),
                "questions": normalized_questions,
            }
        )

    for idx, section in enumerate(normalized_sections, start=1):
        section["id"] = idx

    return normalized_sections, warnings


def _dedupe_question_numbers_within_sections(sections):
    warnings = []
    duplicate_numbers = []
    for section in sections:
        questions = section.get("questions", [])
        seen = set()
        next_number = 1
        for question in questions:
            original = question.get("number")
            try:
                number = int(original)
            except Exception:
                number = None
            if number is None or number in seen:
                while next_number in seen:
                    next_number += 1
                fixed_number = next_number
                if original is not None:
                    warnings.append(
                        f"Duplicate/invalid question number '{original}' in section '{section.get('title', 'Section')}'; "
                        f"renumbered to {fixed_number}."
                    )
                    duplicate_numbers.append(original)
                question["original_number"] = original
                question["number"] = fixed_number
                question["id"] = fixed_number
                seen.add(fixed_number)
                next_number += 1
                continue
            seen.add(number)
            question["number"] = number
            question["id"] = number
            if number >= next_number:
                next_number = number + 1
    return sections, warnings, duplicate_numbers


def detect_ocr_issues(sections):
    warnings = []
    non_sequential_sections = []
    for section in sections:
        numbers = []
        for question in section.get("questions", []):
            try:
                numbers.append(int(question.get("number")))
            except Exception:
                continue
        if len(numbers) < 2:
            continue
        if any(numbers[idx] <= numbers[idx - 1] for idx in range(1, len(numbers))):
            non_sequential_sections.append(section.get("title", "Section"))
    if non_sequential_sections:
        warnings.append("Possible OCR column merge issue: non-sequential question numbering detected.")
    return {
        "warnings": warnings,
        "non_sequential_sections": non_sequential_sections,
        "has_non_sequential": bool(non_sequential_sections),
    }


def validate_enumeration(question):
    warnings = []
    q_text = _normalize_line(question.get("question_text") or question.get("question"))
    answer = str(question.get("correct_answer") or "").strip()
    match = re.search(r"\b(?:give|list|name)\s+(\d+)\b", q_text, flags=re.IGNORECASE)
    if not match:
        return warnings
    expected_count = int(match.group(1))
    if not answer:
        return warnings
    parts = [item.strip() for item in re.split(r"[,;\n]+", answer) if item.strip()]
    if len(parts) < expected_count:
        warnings.append(f"Question {question.get('number')}: Incomplete enumeration answer (expected {expected_count}, got {len(parts)}).")
    return warnings


def _parse_matching_pairs(value):
    if isinstance(value, dict):
        return {str(k).strip(): str(v).strip() for k, v in value.items() if str(k).strip() and str(v).strip()}
    text = str(value or "").strip()
    if not text:
        return {}
    pairs = {}
    for token in re.split(r"[,\n;]+", text):
        part = token.strip()
        if not part:
            continue
        match = re.match(r"^\s*([^:=\-]+)\s*(?:=|:|\-)\s*(.+?)\s*$", part)
        if not match:
            continue
        left = match.group(1).strip()
        right = match.group(2).strip()
        if left and right:
            pairs[left] = right
    return pairs


def find_missing_numbers(sections):
    numbers = []
    for section in sections or []:
        for question in section.get("questions", []) or []:
            try:
                numbers.append(int(question.get("number")))
            except Exception:
                continue
    if not numbers:
        return []
    min_number = min(numbers)
    max_number = max(numbers)
    existing = set(numbers)
    return [value for value in range(min_number, max_number + 1) if value not in existing]


def recover_missing_questions(lines, missing_numbers, section_markers):
    recovered = []
    if not missing_numbers:
        return recovered

    missing_set = set(int(number) for number in missing_numbers)
    question_pattern = re.compile(r"^(?:q\s*)?(\d{1,4})\s*[\.\)\-:]?\s*(\S.*)$", flags=re.IGNORECASE)

    for index, raw_line in enumerate(lines or []):
        line = _normalize_line(raw_line)
        if not line or _is_answer_key_header(line):
            continue
        match = question_pattern.match(line)
        if not match:
            continue
        number = int(match.group(1))
        if number not in missing_set:
            continue
        question_text = _normalize_line(match.group(2))
        if not question_text:
            continue

        marker_idx = 0
        for idx, marker in enumerate(section_markers or []):
            start = marker.get("start", 0)
            if index >= start:
                marker_idx = idx
            else:
                break
        section_type = (section_markers or [{"type": "identification"}])[marker_idx].get("type", "identification")
        q_type = _route_question_type(section_type, question_text, has_options=False)
        recovered.append(
            {
                "section_index": marker_idx,
                "question": {
                    "id": number,
                    "number": number,
                    "question": question_text,
                    "question_text": question_text,
                    "type": q_type,
                    "choices": ["True", "False"] if q_type == "true_false" else [],
                    "options": ([{"id": 1, "text": "True"}, {"id": 2, "text": "False"}] if q_type == "true_false" else []),
                    "correct_answer": "",
                    "answer_key": "",
                    "acceptable_answers": [],
                    "starter_code": "",
                    "expected_output": "",
                    "test_cases": "",
                    "points": 1,
                },
            }
        )
        missing_set.remove(number)
        if not missing_set:
            break
    return recovered


def realign_answers_with_offset(question_numbers, answers_map, offset_range=3):
    numeric_answers = {int(k): v for k, v in (answers_map or {}).items() if str(k).isdigit()}
    question_numbers = sorted({int(number) for number in (question_numbers or [])})
    if not question_numbers or not numeric_answers:
        return answers_map or {}, 0, [], sorted(numeric_answers.keys())

    best_offset = 0
    best_score = -1
    best_bound = {}
    best_unmatched_questions = []
    best_unmatched_answers = sorted(numeric_answers.keys())

    for offset in range(-abs(offset_range), abs(offset_range) + 1):
        candidate = {}
        matched_answer_keys = set()
        for qn in question_numbers:
            source_key = qn + offset
            if source_key in numeric_answers:
                candidate[qn] = numeric_answers[source_key]
                matched_answer_keys.add(source_key)
        score = len(candidate)
        if score > best_score:
            best_score = score
            best_offset = offset
            best_bound = candidate
            best_unmatched_questions = [qn for qn in question_numbers if qn not in candidate]
            best_unmatched_answers = [key for key in sorted(numeric_answers.keys()) if key not in matched_answer_keys]

    adjusted = dict(answers_map or {})
    for key in list(adjusted.keys()):
        if str(key).isdigit():
            adjusted.pop(key, None)
    for qn, value in best_bound.items():
        adjusted[str(qn)] = value
    return adjusted, best_offset, best_unmatched_questions, best_unmatched_answers


def _apply_context_aware_type(question):
    q_text = _normalize_line(question.get("question_text") or question.get("question"))
    choices = question.get("choices", []) or []
    options = question.get("options", []) or []
    q_type = _route_question_type(
        str(question.get("type") or "identification").strip().lower(),
        q_text,
        has_options=bool(choices or options),
    )
    question["type"] = q_type

    if q_type == "true_false":
        question["choices"] = ["True", "False"]
        question["options"] = [{"id": 1, "text": "True"}, {"id": 2, "text": "False"}]
        return

    if q_type in {"essay", "coding", "file_upload"}:
        question["choices"] = []
        question["options"] = []
        return

    if q_type == "identification" and not choices and not options:
        question["type"] = "identification"


def validate_output(
    sections,
    answer_map=None,
    duplicate_numbers=None,
    detected_sections=None,
    section_detection_uncertainty=False,
    mode="balanced",
):
    warnings = []
    strict_errors = []
    answer_map = answer_map or {}
    duplicate_numbers = list(duplicate_numbers or [])
    detected_sections = list(detected_sections or [])
    question_count = 0
    answered_count = 0
    seen_numbers = set()
    unmatched_questions = []
    invalid_mcq_questions = []
    numbering_duplicates = []

    for section in sections:
        for question in section.get("questions", []):
            _apply_context_aware_type(question)
            question_count += 1
            q_number = question.get("number")
            q_text = _normalize_line(question.get("question_text") or question.get("question"))
            question["question"] = q_text
            question["question_text"] = q_text
            question.setdefault("options", [])
            question.setdefault("choices", [])
            question.setdefault("starter_code", "")
            question.setdefault("expected_output", "")
            question.setdefault("test_cases", "")
            question.setdefault("points", 1)

            if not q_text or len(q_text) < 3:
                warnings.append(f"Question {q_number}: Invalid or empty question detected.")

            if q_number in seen_numbers:
                numbering_duplicates.append(q_number)
                warnings.append(f"Duplicate question number detected after normalization: {q_number}")
            seen_numbers.add(q_number)

            q_type = str(question.get("type") or "identification").strip().lower()
            answer_value = str(question.get("correct_answer") or "").strip()
            if answer_value:
                answered_count += 1
            else:
                if q_type in {"multiple_choice", "true_false", "matching"}:
                    unmatched_questions.append(q_number)
                    warnings.append(f"Missing answer for question {q_number}")

            if q_type == "true_false" and answer_value:
                normalized_tf = _normalize_true_false_answer(answer_value)
                if normalized_tf:
                    question["correct_answer"] = normalized_tf

            if q_type == "multiple_choice":
                option_texts = []
                for option in question.get("options", []) or []:
                    if isinstance(option, dict):
                        text = _normalize_line(option.get("text"))
                    else:
                        text = _normalize_line(option)
                    if text:
                        option_texts.append(text)
                if not option_texts:
                    option_texts = [_normalize_line(choice) for choice in (question.get("choices", []) or []) if _normalize_line(choice)]
                answer_token = str(question.get("correct_answer") or "").strip()
                if answer_token:
                    is_valid = False
                    letter_match = re.match(r"^([A-Ha-h])(?:[\)\.\-:]|$)", answer_token)
                    if letter_match:
                        idx = ord(letter_match.group(1).upper()) - ord("A")
                        if 0 <= idx < len(option_texts):
                            is_valid = True
                    else:
                        normalized_answer = answer_token.lower()
                        if any(option.lower() == normalized_answer for option in option_texts):
                            is_valid = True
                    if not is_valid:
                        invalid_mcq_questions.append(q_number)
                        warnings.append(f"Question {q_number}: Answer does not match any option.")

            if q_type == "enumeration":
                warnings.extend(validate_enumeration(question))
            if q_type == "matching" and answer_value:
                if not _parse_matching_pairs(answer_value):
                    warnings.append(f"Question {q_number}: Invalid matching pair mapping.")

    unmatched_answers = []
    for key in answer_map.keys():
        if str(key).isdigit():
            if int(key) not in seen_numbers:
                unmatched_answers.append(key)

    if question_count > 0 and len(unmatched_answers) > 0:
        warnings.append(f"Extra answers detected: {len(unmatched_answers)} unmatched answer(s).")

    missing_answers = len(unmatched_questions) > 0
    too_many_unanswered = question_count > 0 and (len(unmatched_questions) / max(question_count, 1)) > 0.10
    if too_many_unanswered:
        warnings.append("Too many missing answers.")

    ocr = detect_ocr_issues(sections)
    warnings.extend(ocr["warnings"])
    if section_detection_uncertainty:
        warnings.append("Section detection uncertainty detected; parser relied on fallback section mapping.")

    if missing_answers:
        strict_errors.append("Missing answers detected.")
    if invalid_mcq_questions:
        strict_errors.append("Invalid MCQ answer mapping detected.")
    if ocr["has_non_sequential"]:
        strict_errors.append("Possible OCR issue detected.")

    if mode == "tolerant":
        essential = []
        for warning in warnings:
            if any(key in warning.lower() for key in ["no questions", "too many missing answers", "ocr", "invalid or empty question"]):
                essential.append(warning)
        warnings = essential

    debug = {
        "unmatched_questions": [q for q in unmatched_questions if q is not None],
        "unmatched_answers": unmatched_answers,
        "duplicate_numbers": list(dict.fromkeys([*duplicate_numbers, *numbering_duplicates])),
        "detected_sections": detected_sections,
    }

    flags = {
        "missing_answers": missing_answers,
        "duplicate_numbers": bool(debug["duplicate_numbers"]),
        "non_sequential_numbers": ocr["has_non_sequential"],
        "invalid_mcq_mapping": bool(invalid_mcq_questions),
        "section_detection_uncertainty": bool(section_detection_uncertainty),
        "too_many_unanswered": too_many_unanswered,
    }
    return {"warnings": warnings, "debug": debug, "flags": flags, "strict_errors": strict_errors}


def compute_confidence(flags):
    score = 100
    if flags.get("missing_answers"):
        score -= 20
    if flags.get("duplicate_numbers"):
        score -= 15
    if flags.get("non_sequential_numbers"):
        score -= 15
    if flags.get("invalid_mcq_mapping"):
        score -= 10
    if flags.get("section_detection_uncertainty"):
        score -= 10
    if flags.get("too_many_unanswered"):
        score -= 10
    return max(0, min(100, score))


def _finalize_import_result(
    sections,
    warnings,
    notes,
    answer_map,
    line_count,
    mode="balanced",
    detected_sections=None,
    duplicate_numbers=None,
    section_detection_uncertainty=False,
    debug_extra=None,
):
    validated = validate_output(
        sections,
        answer_map=answer_map,
        duplicate_numbers=duplicate_numbers,
        detected_sections=detected_sections,
        section_detection_uncertainty=section_detection_uncertainty,
        mode=mode,
    )
    merged_warnings = [*(warnings or []), *(validated.get("warnings") or [])]
    # keep warning order deterministic and compact
    deduped_warnings = []
    seen = set()
    for warning in merged_warnings:
        key = str(warning).strip().lower()
        if not key or key in seen:
            continue
        seen.add(key)
        deduped_warnings.append(warning)

    if mode == "tolerant":
        deduped_warnings = [
            warning
            for warning in deduped_warnings
            if any(
                key in str(warning).lower()
                for key in [
                    "too many missing answers",
                    "ocr",
                    "no questions",
                    "invalid or empty question",
                    "extra answers detected",
                ]
            )
        ]

    confidence_score = compute_confidence(validated.get("flags") or {})
    merged_debug = {
        "unmatched_questions": [],
        "unmatched_answers": [],
        "duplicate_numbers": [],
        "detected_sections": detected_sections or [],
    }
    merged_debug.update(validated.get("debug") or {})
    if isinstance(debug_extra, dict):
        for key, value in debug_extra.items():
            merged_debug[key] = value

    result = {
        "data": {"sections": sections},
        "sections": sections,
        "warnings": deduped_warnings,
        "notes": notes or [],
        "answer_map": answer_map or {},
        "line_count": line_count,
        "confidence_score": confidence_score,
        "debug": merged_debug,
    }
    if mode == "strict" and validated.get("strict_errors"):
        result["errors"] = validated["strict_errors"]
    return result


def _parse_json_import_payload(raw_text, mode="balanced"):
    text = str(raw_text or "").strip()
    if not text or text[0] not in {"{", "["}:
        return None
    try:
        payload = json.loads(text)
    except Exception:
        return None

    sections = []
    warnings = []
    if isinstance(payload, dict):
        candidate_sections = payload.get("sections")
        if isinstance(candidate_sections, list):
            sections = candidate_sections
        elif isinstance(payload.get("questions"), list):
            sections = [
                {
                    "id": 1,
                    "title": payload.get("title") or "General Section",
                    "type": payload.get("type") or "identification",
                    "instructions": payload.get("instructions", ""),
                    "questions": payload.get("questions"),
                }
            ]
        elif isinstance(payload.get("quiz"), dict) and isinstance(payload["quiz"].get("sections"), list):
            sections = payload["quiz"]["sections"]
        elif isinstance(payload.get("exam"), dict) and isinstance(payload["exam"].get("sections"), list):
            sections = payload["exam"]["sections"]
        else:
            warnings.append("JSON payload detected but no sections/questions were found.")
    elif isinstance(payload, list):
        sections = [{"id": 1, "title": "General Section", "type": "identification", "instructions": "", "questions": payload}]
    else:
        warnings.append("Unsupported JSON payload shape for quiz import.")

    normalized_sections, normalize_warnings = _normalize_sections_payload(sections)
    warnings.extend(normalize_warnings)
    normalized_sections, dedupe_warnings, duplicate_numbers = _dedupe_question_numbers_within_sections(normalized_sections)
    warnings.extend(dedupe_warnings)
    detected_sections = [{"title": section.get("title"), "type": section.get("type")} for section in normalized_sections]
    section_uncertain = len(detected_sections) <= 1 and (
        not detected_sections or str(detected_sections[0].get("title", "")).lower() == "general section"
    )

    detected_question_count = sum(len(section.get("questions", [])) for section in normalized_sections)
    if detected_question_count > 0:
        logger.info("Detected %s questions", detected_question_count)
        logger.info("Detected %s answers", 0)
        logger.info("Merged %s questions successfully", detected_question_count)
    return _finalize_import_result(
        normalized_sections,
        warnings=warnings,
        notes=[],
        answer_map={},
        line_count=len(text.splitlines()),
        mode=mode,
        detected_sections=detected_sections,
        duplicate_numbers=duplicate_numbers,
        section_detection_uncertainty=section_uncertain,
    )


def parse_questions(lines, section_markers):
    warnings = []
    inline_answer_map = {}
    answer_pattern = re.compile(r"^(?:(?:answer\s*[:=]|ans\s*[:=]|->|\u2192)\s*(.+)|=\s+(.+))$", flags=re.IGNORECASE)
    answer_inline_pattern = re.compile(r"(?:\b(?:answer|ans)\s*[:=]|->|\u2192|^\s*=\s+)\s*(.+?)\s*$", flags=re.IGNORECASE)
    bracket_answer_pattern = re.compile(
        r"\[\s*(?:Answer|Ans)\s*#?\s*(\d+)?\s*:\s*(.*?)\s*\]",
        flags=re.IGNORECASE,
    )

    def _normalize_inline_answer(raw_value):
        value = _normalize_line(raw_value).rstrip(",.")
        lower = value.lower()
        if lower in {"true", "false"}:
            return lower.upper()
        if re.match(r"^[a-d]$", lower):
            return lower.upper()
        return value

    def _extract_inline_from_text(text):
        match = answer_inline_pattern.search(text or "")
        if not match:
            return _normalize_line(text or ""), ""
        extracted = _normalize_inline_answer(match.group(1))
        cleaned = _normalize_line((text or "")[: match.start()])
        return cleaned, extracted

    def _extract_bracketed_answers(text, fallback_number):
        raw_text = str(text or "")
        detected = []
        for match in bracket_answer_pattern.finditer(raw_text):
            question_number = str(match.group(1) or fallback_number)
            answer_value = _normalize_inline_answer(match.group(2))
            detected.append((question_number, answer_value))
        cleaned = _normalize_line(bracket_answer_pattern.sub(" ", raw_text))
        cleaned = re.sub(r"\s{2,}", " ", cleaned).strip()
        cleaned = cleaned.rstrip(" ,;:-")
        return cleaned, detected

    def _attach_answer_to_last(raw_value):
        if last_question_number is None:
            return
        normalized = _normalize_inline_answer(raw_value)
        if not normalized:
            return
        key = str(last_question_number)
        if not inline_answer_map.get(key):
            inline_answer_map[key] = normalized
        logger.debug("Attached answer to Q%s", key)

    if not section_markers:
        section_markers = [{"start": 0, "title": "General Section", "type": "identification"}]

    sections = [
        {
            "id": marker_index + 1,
            "title": marker["title"],
            "type": marker["type"],
            "instructions": "",
            "questions": [],
        }
        for marker_index, marker in enumerate(section_markers)
    ]

    marker_idx = 0
    next_start = section_markers[1]["start"] if len(section_markers) > 1 else len(lines) + 1
    current_section = sections[marker_idx]
    current_question = None
    current_choices = {}
    current_choice_letter = None
    last_question_number = None
    mode = "question"

    def flush_choice_state():
        nonlocal current_choices, current_choice_letter
        if current_question is None:
            current_choices = {}
            current_choice_letter = None
            return
        if current_choices:
            ordered = [current_choices[key] for key in sorted(current_choices.keys()) if current_choices.get(key)]
            current_question["choices"] = ordered
            current_question["options"] = [{"id": idx + 1, "text": item} for idx, item in enumerate(ordered)]
            if current_question.get("type") == "identification":
                current_question["type"] = "multiple_choice"
        current_choices = {}
        current_choice_letter = None

    for index, raw_line in enumerate(lines):
        while marker_idx + 1 < len(section_markers) and index >= next_start:
            flush_choice_state()
            marker_idx += 1
            current_section = sections[marker_idx]
            next_start = section_markers[marker_idx + 1]["start"] if marker_idx + 1 < len(section_markers) else len(lines) + 1
            current_question = None
            mode = "question"

        line = _normalize_line(raw_line)
        if not line:
            continue

        # Global answer-line detection (state-independent):
        # Must attach via last_question_number regardless of mode/current_question.
        answer_line_match = answer_pattern.match(line)
        if answer_line_match:
            raw_value = answer_line_match.group(1) or answer_line_match.group(2)
            logger.debug("Detected answer line: %s", raw_value)
            if last_question_number is not None:
                logger.debug("Attaching to question #%s", last_question_number)
                _attach_answer_to_last(raw_value)
            continue

        bracket_only_match = bracket_answer_pattern.fullmatch(line)
        if bracket_only_match:
            raw_value = bracket_only_match.group(2)
            logger.debug("Detected answer line: %s", raw_value)
            if last_question_number is not None:
                logger.debug("Attaching to question #%s", last_question_number)
                _attach_answer_to_last(raw_value)
            continue

        # Pipe-delimited format support:
        # 1|Question?|A.Option|B.Option|C.Option|D.Option
        # Answer|A
        if "|" in line and re.match(r"^\s*answer\s*\|", line, flags=re.IGNORECASE):
            if last_question_number is not None:
                parts = line.split("|", 1)
                raw_answer = parts[1] if len(parts) > 1 else ""
                logger.debug("Detected answer line: %s", raw_answer)
                logger.debug("Attaching to question #%s", last_question_number)
                _attach_answer_to_last(raw_answer)
                logger.debug("Parsed pipe-format answer for Q%s", last_question_number)
            continue

        if "|" in line and re.match(r"^\s*\d+\|", line):
            flush_choice_state()
            parts = [item.strip() for item in line.split("|")]
            if len(parts) >= 2 and str(parts[0]).isdigit():
                question_number = int(parts[0])
                question_text = _normalize_line(parts[1])
                question_text, bracketed = _extract_bracketed_answers(question_text, question_number)
                for q_number, q_answer in bracketed:
                    inline_answer_map[str(q_number)] = q_answer
                    logger.debug("Detected bracketed inline answer for Q%s", q_number)
                question_text, extracted_inline = _extract_inline_from_text(question_text)
                if extracted_inline:
                    inline_answer_map[str(question_number)] = extracted_inline
                    logger.debug("Detected inline answer for Q%s", question_number)

                choice_parts = [item for item in parts[2:] if _normalize_line(item)]
                choices = []
                has_labeled_choices = any(
                    re.match(r"^([A-Da-d])[\.\)]?\s*(.+)$", choice, flags=re.IGNORECASE)
                    for choice in choice_parts
                )
                if has_labeled_choices:
                    for choice in choice_parts:
                        match = re.match(r"^([A-Da-d])[\.\)]?\s*(.+)$", choice, flags=re.IGNORECASE)
                        if not match:
                            continue
                        choice_text = _normalize_line(match.group(2)).rstrip(",.")
                        if choice_text:
                            choices.append(choice_text)

                q_type = _route_question_type(
                    current_section.get("type", "identification"),
                    question_text,
                    has_options=bool(choices),
                )
                current_question = {
                    "id": question_number,
                    "number": question_number,
                    "question": question_text,
                    "question_text": question_text,
                    "type": q_type,
                    "choices": choices,
                    "options": [{"id": idx + 1, "text": choice} for idx, choice in enumerate(choices)],
                    "correct_answer": "",
                    "answer_key": "",
                    "acceptable_answers": [],
                    "starter_code": "",
                    "expected_output": "",
                    "test_cases": "",
                    "points": 1,
                }
                if q_type == "true_false":
                    current_question["choices"] = ["True", "False"]
                    current_question["options"] = [
                        {"id": 1, "text": "True"},
                        {"id": 2, "text": "False"},
                    ]
                current_section["questions"].append(current_question)
                current_choices = {}
                last_question_number = question_number
                logger.debug("Set last_question_number = %s", question_number)
                logger.debug("Detected pipe-format question")
                continue

        if _is_question_start(line):
            flush_choice_state()
            number, question_text = _extract_question_number_and_text(line)
            if number is None:
                continue

            question_text, bracketed = _extract_bracketed_answers(question_text.strip(), number)

            question_text, extracted_inline = _extract_inline_from_text(question_text)
            if extracted_inline:
                inline_answer_map[str(number)] = extracted_inline
                logger.debug("Detected inline answer for Q%s", number)
            inline_choices = _extract_mcq_choices(question_text)
            if len(inline_choices) >= 2:
                first_choice_match = re.search(r"[A-Da-d][\.\)]\s*", question_text)
                if first_choice_match:
                    question_text = _normalize_line(question_text[: first_choice_match.start()])
                current_choices = {letter: text for letter, text in inline_choices}
            else:
                current_choices = {}

            q_type = _route_question_type(
                current_section.get("type", "identification"),
                question_text,
                has_options=bool(current_choices),
            )

            current_question = {
                "id": number,
                "number": number,
                "question": question_text,
                "question_text": question_text,
                "type": q_type,
                "choices": [],
                "options": [],
                "correct_answer": "",
                "answer_key": "",
                "acceptable_answers": [],
                "starter_code": "",
                "expected_output": "",
                "test_cases": "",
                "points": 1,
            }
            current_section["questions"].append(current_question)
            last_question_number = number
            logger.debug("Set last_question_number = %s", number)
            for _, q_answer in bracketed:
                if q_answer:
                    _attach_answer_to_last(q_answer)
                    logger.debug("Extracted inline bracket answer for Q%s: %s", last_question_number, q_answer)
            mode = "choices" if q_type == "multiple_choice" or len(current_choices) >= 2 else "question"
            continue

        if current_question is None:
            continue

        horizontal_choices = _extract_mcq_choices(line)
        if len(horizontal_choices) >= 2:
            for letter, text in horizontal_choices:
                current_choices[letter] = text
                current_choice_letter = letter
            mode = "choices"
            continue
        vertical_choice = _extract_vertical_choice(line)
        if vertical_choice:
            current_choices[vertical_choice[0]] = vertical_choice[1]
            current_choice_letter = vertical_choice[0]
            mode = "choices"
            continue

        if re.match(r"^[_\-\=\.\s]{3,}$", line):
            continue

        line, bracketed = _extract_bracketed_answers(line, current_question.get("number"))
        for _, q_answer in bracketed:
            if q_answer:
                _attach_answer_to_last(q_answer)
                logger.debug("Extracted inline bracket answer for Q%s: %s", last_question_number, q_answer)
        if not line:
            continue

        # OCR resilience: if option text is broken into the next line and we are in
        # choices mode, append continuation to the last option instead of question text.
        if mode == "choices" and current_choice_letter and current_choices.get(current_choice_letter):
            continuation = _normalize_line(line)
            if continuation and not _is_question_start(continuation) and not _extract_vertical_choice(continuation):
                current_choices[current_choice_letter] = _normalize_line(
                    f"{current_choices[current_choice_letter]} {continuation}"
                )
                continue

        # Prevent inline answer contamination inside mixed content lines.
        clean_part, extracted_answer = _extract_inline_from_text(line)
        if extracted_answer:
            _attach_answer_to_last(extracted_answer)
            if not clean_part:
                continue
            line = clean_part

        if current_question.get("type") == "coding":
            starter = current_question.get("starter_code", "")
            current_question["starter_code"] = f"{starter}\n{raw_line}".strip() if starter else raw_line.strip()
        else:
            separator = "\n" if current_question.get("type") == "coding" else " "
            updated_text = f"{current_question.get('question', '')}{separator}{line}".strip()
            current_question["question"] = updated_text
            current_question["question_text"] = updated_text
        mode = "question"

    flush_choice_state()

    cleaned_sections = []
    for section in sections:
        if section["questions"]:
            cleaned_sections.append(section)

    # Fallback pass 1: if sections exist but question regex missed all items,
    # retry globally with relaxed numbering capture and attach to first section.
    if not cleaned_sections:
        fallback_questions = []
        fallback_pattern = re.compile(
            r"^(?:q\s*)?(\d{1,3})\s*(?:[\.\)\-:]\s*|\s+)(.+)$",
            flags=re.IGNORECASE,
        )
        seen_numbers = set()
        for raw_line in lines:
            line = _normalize_line(raw_line)
            if not line or _is_answer_key_header(line):
                continue
            match = fallback_pattern.match(line)
            if not match:
                continue
            q_no = int(match.group(1))
            q_text = _normalize_line(match.group(2))
            if not q_text:
                continue
            # Ignore answer-key-like single token lines (e.g., "3 C", "4 TRUE")
            if re.match(r"^(?:[A-H]|TRUE|FALSE)$", q_text, flags=re.IGNORECASE):
                continue
            if q_no in seen_numbers:
                continue
            seen_numbers.add(q_no)
            fallback_questions.append(
                {
                    "id": q_no,
                    "number": q_no,
                    "question": q_text,
                    "question_text": q_text,
                    "type": "identification",
                    "choices": [],
                    "options": [],
                    "correct_answer": "",
                    "answer_key": "",
                    "acceptable_answers": [],
                    "starter_code": "",
                    "expected_output": "",
                    "test_cases": "",
                    "points": 1,
                }
            )

        if fallback_questions:
            cleaned_sections = [
                {
                    "id": 1,
                    "title": sections[0]["title"] if sections else "General Section",
                    "type": sections[0]["type"] if sections else "identification",
                    "instructions": "",
                    "questions": fallback_questions,
                }
            ]

    # Fallback pass 2: recover when source text has unnumbered lines
    # (common with Word auto-numbered lists where numbering is not exported).
    if not cleaned_sections and sections:
        heuristic_sections = []
        running_number = 1

        for marker_index, marker in enumerate(section_markers):
            start = marker["start"]
            end = section_markers[marker_index + 1]["start"] if marker_index + 1 < len(section_markers) else len(lines)
            section_lines = [_normalize_line(item) for item in lines[start:end]]
            section_title = _normalize_line(marker.get("title", "Section"))
            section_type = marker.get("type", "identification")

            if section_title.lower() == "general section":
                continue

            questions = []
            i = 0
            while i < len(section_lines):
                line = section_lines[i]
                if (
                    not line
                    or line == section_title
                    or _is_answer_key_header(line)
                    or _is_instruction_line(line)
                    or re.match(r"^[_\-\=\.\s]{3,}$", line)
                ):
                    i += 1
                    continue

                if _is_question_start(line):
                    i += 1
                    continue

                if section_type == "multiple_choice":
                    if _extract_vertical_choice(line):
                        i += 1
                        continue
                    inline = _extract_mcq_choices(line)
                    if len(inline) >= 2:
                        i += 1
                        continue

                    stem = line
                    choices_map = {}
                    j = i + 1
                    while j < len(section_lines):
                        candidate = section_lines[j]
                        if not candidate or _is_instruction_line(candidate):
                            j += 1
                            continue
                        vc = _extract_vertical_choice(candidate)
                        if vc:
                            choices_map[vc[0]] = vc[1]
                            j += 1
                            continue
                        hc = _extract_mcq_choices(candidate)
                        if len(hc) >= 2:
                            for letter, text in hc:
                                choices_map[letter] = text
                            j += 1
                            continue
                        break

                    ordered = [choices_map[key] for key in sorted(choices_map.keys()) if choices_map.get(key)]
                    if ordered:
                        questions.append(
                            {
                                "id": running_number,
                                "number": running_number,
                                "question": stem,
                                "question_text": stem,
                                "type": "multiple_choice",
                                "choices": ordered,
                                "options": [{"id": idx + 1, "text": item} for idx, item in enumerate(ordered)],
                                "correct_answer": "",
                                "answer_key": "",
                                "acceptable_answers": [],
                                "starter_code": "",
                                "expected_output": "",
                                "test_cases": "",
                                "points": 1,
                            }
                        )
                        running_number += 1
                        i = j
                        continue

                if len(line.split()) < 3:
                    i += 1
                    continue

                questions.append(
                    {
                        "id": running_number,
                        "number": running_number,
                        "question": line,
                        "question_text": line,
                        "type": section_type,
                        "choices": ["True", "False"] if section_type == "true_false" else [],
                        "options": (
                            [{"id": 1, "text": "True"}, {"id": 2, "text": "False"}]
                            if section_type == "true_false"
                            else []
                        ),
                        "correct_answer": "",
                        "answer_key": "",
                        "acceptable_answers": [],
                        "starter_code": "",
                        "expected_output": "",
                        "test_cases": "",
                        "points": 1,
                    }
                )
                running_number += 1
                i += 1

            if questions:
                heuristic_sections.append(
                    {
                        "id": len(heuristic_sections) + 1,
                        "title": section_title,
                        "type": section_type,
                        "instructions": "",
                        "questions": questions,
                    }
                )

        if heuristic_sections:
            cleaned_sections = heuristic_sections
            warnings.append(HEURISTIC_RECOVERY_WARNING)

    if section_markers and not cleaned_sections:
        warnings.append("Sections detected but no questions attached.")
    if not cleaned_sections:
        warnings.append("No questions parsed. Likely format mismatch.")
    else:
        # Keep section ids compact after filtering/remapping.
        for idx, section in enumerate(cleaned_sections, start=1):
            section["id"] = idx

    logger.debug(
        "Parser PASS2 questions detected: %s",
        [{"title": section["title"], "count": len(section["questions"])} for section in cleaned_sections],
    )
    return cleaned_sections, warnings, inline_answer_map


def parse_answers(lines):
    warnings = []
    answer_map = {}
    # Bug Fix:
    # Duplicate detection should be scoped by (section, number), not number-only.
    duplicate_counter = defaultdict(int)
    section_scoped_answers = defaultdict(dict)
    numbered_by_section = {
        "identification": [],
        "multiple_choice": [],
        "true_false": [],
        "short_answer": [],
        "enumeration": [],
        "coding": [],
        "essay": [],
        "matching": [],
        "file_upload": [],
    }
    positional_by_section = {
        "identification": [],
        "multiple_choice": [],
        "true_false": [],
        "short_answer": [],
        "enumeration": [],
        "coding": [],
        "essay": [],
        "matching": [],
        "file_upload": [],
    }
    positional_global = []
    if not lines:
        logger.debug("Parser PASS3 answers: no answer key content found after split.")
        return answer_map, warnings, {"by_section": positional_by_section, "numbered_by_section": numbered_by_section, "global": positional_global}

    section_label_pattern = re.compile(
        r"^(?:(part|section)\s*([ivxlcdm]+|\d+|[a-z])\s*[:\-]?\s*)?"
        r"(identification|short\s*answer|multiple\s*choice|true\s*(?:/|\s+or\s+)?\s*false|enumeration|coding|essay|matching|file\s*upload)"
        r"(?:\s*\([^)]*\))?\s*:?\s*$",
        flags=re.IGNORECASE,
    )
    answer_start_pattern = re.compile(r"^(?:q\s*)?(\d{1,3})\s*[\.\)\:\-]?\s*(.*)$", flags=re.IGNORECASE)
    compact_pair_pattern = re.compile(
        r"(\d{1,4})\s*[-\u2013\u2014]\s*([A-Za-z0-9#().+/:,;| ]+?)(?=\s+\d{1,4}\s*[-\u2013\u2014]|$)"
    )

    current_number = None
    current_lines = []
    current_section = ""

    def commit_current():
        nonlocal current_number, current_lines
        if current_number is None:
            return
        filtered_lines = [line.rstrip() for line in current_lines if _normalize_line(line)]
        if not filtered_lines:
            warnings.append(f"Missing answer text for question {current_number}")
            current_number = None
            current_lines = []
            return
        if current_section == "coding" and len(filtered_lines) > 1:
            answer_value = "\n".join(filtered_lines).strip()
        elif len(filtered_lines) > 1:
            answer_value = " ".join([_normalize_line(line) for line in filtered_lines]).strip()
        else:
            answer_value = _normalize_line(filtered_lines[0])
        key = str(current_number)
        section_key = current_section or "identification"
        if key in section_scoped_answers[section_key]:
            duplicate_counter[(section_key, current_number)] += 1
        else:
            section_scoped_answers[section_key][key] = answer_value
        # Keep flat mapping for backward compatibility.
        if key not in answer_map:
            answer_map[key] = answer_value
        numbered_by_section.setdefault(section_key, []).append({"number": current_number, "value": answer_value})
        current_number = None
        current_lines = []

    for raw_line in lines:
        line = raw_line.rstrip()
        normalized = _normalize_line(line)
        if not normalized:
            if current_section == "coding" and current_number is not None:
                current_lines.append("")
            continue

        # Compact answer-key support:
        # 1-B 2-C 3-A ... 31-Variable
        # Works with mixed answer types in one line.
        compact_pairs = compact_pair_pattern.findall(normalized)
        if compact_pairs:
            commit_current()
            for q_number_str, value in compact_pairs:
                q_number = int(q_number_str)
                answer_value = _normalize_line(value).rstrip(",.;")
                key = str(q_number)
                section_key = current_section or "identification"
                if key in section_scoped_answers[section_key]:
                    duplicate_counter[(section_key, q_number)] += 1
                    continue
                section_scoped_answers[section_key][key] = answer_value
                if key not in answer_map:
                    answer_map[key] = answer_value
                numbered_by_section.setdefault(section_key, []).append({"number": q_number, "value": answer_value})
            continue

        section_match = section_label_pattern.match(normalized)
        if section_match:
            commit_current()
            current_section = _normalize_answer_section(section_match.group(3) or normalized)
            continue

        match = answer_start_pattern.match(normalized)
        if match:
            commit_current()
            current_number = int(match.group(1))
            first_value = match.group(2).rstrip()
            current_lines = [first_value] if first_value else []
            continue

        if current_number is None:
            if _is_instruction_line(normalized) or re.match(r"^[_\-\=\.\s]{3,}$", normalized):
                continue
            section_key = current_section or "identification"
            cleaned_value = normalized
            positional_by_section.setdefault(section_key, []).append(cleaned_value)
            positional_global.append(cleaned_value)
            continue

        if current_number is not None:
            current_lines.append(line)

    commit_current()
    for section_number_key in sorted(duplicate_counter.keys()):
        section_name, answer_number = section_number_key
        duplicate_count = duplicate_counter[section_number_key]
        warnings.append(
            f"Duplicate answer number {answer_number} in section '{section_name}' encountered {duplicate_count} time(s); "
            "kept first value and skipped later duplicates."
        )

    logger.debug("Parser PASS3 answers detected: %s", answer_map)
    return answer_map, warnings, {
        "by_section": positional_by_section,
        "numbered_by_section": numbered_by_section,
        "global": positional_global,
        # Enhancement:
        # Section-aware answers map extends flat map without breaking callers.
        "by_number_and_section": section_scoped_answers,
    }


def _normalize_question_numbers_for_global_key(sections, answer_map):
    """
    If parser recovered questions with numbering restarted per section but answer key
    is global (1..N), remap question numbers sequentially to avoid wrong merges.
    """
    all_questions = []
    for section in sections:
        all_questions.extend(section.get("questions", []))
    if not all_questions:
        return sections

    numbers = [question.get("number") for question in all_questions if isinstance(question.get("number"), int)]
    if len(numbers) != len(all_questions):
        return sections
    if len(set(numbers)) == len(numbers):
        return sections

    numeric_keys = sorted({int(key) for key in answer_map.keys() if str(key).isdigit()})
    if not numeric_keys:
        return sections
    max_question_no = max(numbers) if numbers else 0
    has_global_numbering_signal = (
        len(numeric_keys) >= len(all_questions)
        or (numeric_keys and max(numeric_keys) > max_question_no)
    )
    if not has_global_numbering_signal:
        return sections

    running = 1
    for section in sections:
        for question in section.get("questions", []):
            question["original_number"] = question.get("number")
            question["number"] = running
            question["id"] = running
            running += 1
    return sections


def merge_results(sections, answer_map, positional_answers=None, inline_answer_map=None):
    warnings = []
    merged_count = 0
    positional_answers = positional_answers or {"by_section": {}, "global": []}
    inline_answer_map = inline_answer_map or {}
    all_question_numbers = []
    for section in sections or []:
        for question in section.get("questions", []) or []:
            try:
                all_question_numbers.append(int(question.get("number")))
            except Exception:
                continue

    strict_number_map = {str(key): value for key, value in (answer_map or {}).items() if str(key).isdigit()}
    section_number_map = positional_answers.get("by_number_and_section", {}) if isinstance(positional_answers, dict) else {}
    unmatched_questions = [number for number in all_question_numbers if str(number) not in strict_number_map and str(number) not in inline_answer_map]
    unmatched_answers = [int(key) for key in strict_number_map.keys() if int(key) not in set(all_question_numbers)]
    alignment_offset = 0
    realignment_skipped_reason = ""
    if unmatched_questions or unmatched_answers:
        # Critical Fix:
        # Skip risky realignment when answer keys are already number-based.
        direct_matches = sum(1 for number in all_question_numbers if str(number) in strict_number_map)
        match_rate = direct_matches / max(len(all_question_numbers), 1) if all_question_numbers else 1.0
        if strict_number_map:
            realignment_skipped_reason = "number_based_answers_detected"
        elif match_rate >= 0.8:
            realignment_skipped_reason = "high_match_rate"
        else:
            aligned_map, alignment_offset, _, _ = realign_answers_with_offset(all_question_numbers, strict_number_map, offset_range=3)
            strict_number_map = {str(key): value for key, value in aligned_map.items() if str(key).isdigit()}
            if alignment_offset != 0:
                warnings.append(f"Answer key realignment applied with offset {alignment_offset}.")
    answer_map = {**answer_map, **strict_number_map}

    for section in sections:
        section_type = str(section.get("type") or "identification")
        section_answers = section_number_map.get(section_type, {}) if isinstance(section_number_map, dict) else {}
        for question in section.get("questions", []):
            number_key = str(question.get("number"))
            answer_value = section_answers.get(number_key) if isinstance(section_answers, dict) else None
            if not answer_value:
                answer_value = answer_map.get(number_key)
            if not answer_value:
                inline_value = inline_answer_map.get(number_key)
                if inline_value:
                    answer_value = inline_value
                    logger.debug("Using inline fallback for Q%s", number_key)
            if not answer_value:
                if not question.get("correct_answer"):
                    if str(question.get("type") or "").lower() in {"multiple_choice", "true_false", "matching"}:
                        warnings.append(f"Missing answer for question {number_key}")
                continue

            merged_count += 1
            raw_answer = _normalize_line(answer_value)
            question_type = question.get("type")

            if question_type == "multiple_choice":
                letter_match = re.match(r"^([A-Ha-h])(?:[\)\.\-:]|$)", raw_answer)
                if letter_match:
                    letter = letter_match.group(1).upper()
                    question["answer_key"] = letter
                    question["correct_answer"] = letter
                    choice_index = ord(letter) - ord("A")
                    choices = question.get("choices", [])
                    if 0 <= choice_index < len(choices):
                        question["correct_answer_text"] = choices[choice_index]
                else:
                    question["correct_answer"] = raw_answer
            elif question_type == "true_false":
                normalized = _normalize_true_false_answer(raw_answer)
                question["correct_answer"] = normalized or raw_answer
            elif question_type == "enumeration":
                parts = [item.strip() for item in re.split(r"[;,/]", raw_answer) if item.strip()]
                question["acceptable_answers"] = parts or [raw_answer]
                question["correct_answer"] = raw_answer
            elif question_type == "short_answer":
                valid_answers = [item.strip() for item in re.split(r"[|;\n]+", answer_value) if item.strip()]
                question["acceptable_answers"] = valid_answers or [raw_answer]
                question["correct_answer"] = raw_answer
            elif question_type == "identification":
                keywords = [item.strip() for item in re.split(r"[|,;\n]+", answer_value) if item.strip()]
                question["acceptable_answers"] = keywords or question.get("acceptable_answers", [])
                question["correct_answer"] = raw_answer
            elif question_type in {"coding", "essay"}:
                question["correct_answer"] = answer_value
                if not question.get("expected_output"):
                    question["expected_output"] = answer_value
            elif question_type == "matching":
                question["correct_answer"] = raw_answer
                question["matching_pairs"] = _parse_matching_pairs(answer_value)
            elif question_type == "file_upload":
                question["correct_answer"] = ""
            else:
                question["correct_answer"] = raw_answer

    merge_debug = {
        "alignment_offset": alignment_offset,
        "unmatched_questions": unmatched_questions,
        "unmatched_answers": unmatched_answers,
        "realignment_skipped_reason": realignment_skipped_reason,
    }
    logger.debug("Parser PASS4 merge results: merged=%s warnings=%s debug=%s", merged_count, warnings, merge_debug)
    return sections, warnings, merged_count, merge_debug


def _parse_questions_from_plain_text(raw_text, mode="balanced"):
    raw_text = str(raw_text or "")
    mode = str(mode or "balanced").strip().lower()
    if mode not in {"strict", "balanced", "tolerant"}:
        mode = "balanced"

    json_parsed = _parse_json_import_payload(raw_text, mode=mode)
    if json_parsed is not None and (json_parsed.get("sections") or json_parsed.get("warnings")):
        return json_parsed

    normalized_full_text = normalize_text(raw_text)
    all_lines = [line for line in normalized_full_text.splitlines() if _normalize_line(line)]
    has_inline_answers = any(
        re.search(r"^\s*(?:\b(?:answer|ans)\s*[:=]|->|\u2192)\s*(.+?)\s*$|^\s*=\s+(.+?)\s*$", line, flags=re.IGNORECASE)
        for line in all_lines
    )

    if has_inline_answers:
        exam_lines = all_lines
        answer_lines = []
    else:
        exam_lines, answer_lines = split_document(raw_text)

    section_markers = parse_sections(exam_lines)
    sections, question_warnings, inline_answer_map = parse_questions(exam_lines, section_markers)
    missing_numbers = find_missing_numbers(sections)
    recovered_question_items = recover_missing_questions(exam_lines, missing_numbers, section_markers)
    recovered_numbers = []
    for item in recovered_question_items:
        section_index = item.get("section_index", 0)
        if 0 <= section_index < len(sections):
            sections[section_index].setdefault("questions", []).append(item["question"])
            recovered_numbers.append(item["question"].get("number"))
    if recovered_numbers:
        question_warnings.append(f"Recovered missing questions: {', '.join(str(num) for num in sorted(recovered_numbers))}.")
    if has_inline_answers:
        answer_map = dict(inline_answer_map)
        answer_warnings = []
        positional_answers = {"by_section": {}, "numbered_by_section": {}, "global": []}
    else:
        answer_map, answer_warnings, positional_answers = parse_answers(answer_lines)
        # Answer key has priority; inline answers are fallback-only.
        for key, value in inline_answer_map.items():
            answer_map.setdefault(key, value)
    sections, dedupe_warnings, duplicate_numbers = _dedupe_question_numbers_within_sections(sections)
    sections = _normalize_question_numbers_for_global_key(sections, answer_map)
    merged_sections, merge_warnings, merged_count, merge_debug = merge_results(
        sections,
        answer_map,
        positional_answers=positional_answers,
        inline_answer_map=inline_answer_map,
    )
    merged_sections, normalization_warnings = _normalize_sections_payload(
        merged_sections,
        answer_map=answer_map,
        inline_answer_map=inline_answer_map,
    )
    merged_count = sum(
        1 for section in merged_sections for question in section.get("questions", []) if question.get("correct_answer")
    )

    detected_question_count = sum(len(section.get("questions", [])) for section in merged_sections)
    detected_answer_count = len(answer_map)

    warnings = [*question_warnings, *answer_warnings, *dedupe_warnings, *merge_warnings, *normalization_warnings]
    notes = []
    if HEURISTIC_RECOVERY_WARNING in warnings and detected_question_count > 0:
        warnings = [warning for warning in warnings if warning != HEURISTIC_RECOVERY_WARNING]
        notes.append(HEURISTIC_RECOVERY_WARNING)
    if detected_question_count > 0 and detected_answer_count == 0 and not has_inline_answers:
        warnings.append("Questions detected but no answers parsed from answer key.")
    if detected_answer_count > 0 and detected_question_count == 0:
        warnings.append("Answer key parsed but exam content failed.")

    logger.info("Detected %s questions", detected_question_count)
    logger.info("Detected %s answers", detected_answer_count)
    logger.info("Merged %s questions successfully", merged_count)
    detected_sections = [{"title": marker.get("title"), "type": marker.get("type")} for marker in section_markers]
    section_uncertain = len(detected_sections) <= 1 and (
        not detected_sections or str(detected_sections[0].get("title", "")).lower() == "general section"
    )
    return _finalize_import_result(
        merged_sections,
        warnings=warnings,
        notes=notes,
        answer_map=answer_map,
        line_count=len(exam_lines) + len(answer_lines),
        mode=mode,
        detected_sections=detected_sections,
        duplicate_numbers=duplicate_numbers,
        section_detection_uncertainty=section_uncertain,
        debug_extra={
            "missing_questions": missing_numbers,
            "recovered_questions": sorted(recovered_numbers),
            "alignment_offset": merge_debug.get("alignment_offset", 0),
            "unmatched_answers": merge_debug.get("unmatched_answers", []),
            "realignment_skipped_reason": merge_debug.get("realignment_skipped_reason", ""),
        },
    )


@api_view(["GET", "POST"])
@permission_classes([IsAuthenticated])
def exam_quizzes_list(request, course_id):
    course = _get_course_with_access(course_id, request.user)
    if not course:
        return Response({"error": "Course not found or access denied"}, status=404)

    if request.method == "GET":
        queryset = (
            CourseActivity.objects.filter(
                Q(course=course) | Q(assigned_courses=course),
                activity_type__name__iexact="quiz",
            )
            .select_related("course", "activity_type")
            .prefetch_related("assigned_courses")
            .distinct()
            .order_by("-created_at")
        )
        serializer = CourseActivitySerializer(queryset, many=True, context={"request": request})
        return Response(serializer.data)

    if request.user.role != "instructor" or course.instructor_id != request.user.id:
        return Response({"error": "Only instructor can create exam or quiz."}, status=403)

    quiz_type = _quiz_type()
    if not quiz_type:
        return Response({"error": "Quiz activity type is not configured."}, status=400)

    payload = request.data.copy()
    payload["course"] = course.id
    payload["activity_type"] = quiz_type.id
    payload.setdefault("assessment_type", "quiz")
    payload.setdefault("publish_state", "draft")
    payload.setdefault("grading_type", "points")

    serializer = CourseActivitySerializer(data=payload, context={"request": request})
    if serializer.is_valid():
        activity = serializer.save()
        return Response(CourseActivitySerializer(activity, context={"request": request}).data, status=201)
    return Response(serializer.errors, status=400)


@api_view(["GET", "PUT", "PATCH", "DELETE"])
@permission_classes([IsAuthenticated])
def exam_quiz_detail(request, course_id, activity_id):
    course = _get_course_with_access(course_id, request.user)
    if not course:
        return Response({"error": "Course not found or access denied"}, status=404)

    activity = get_object_or_404(
        CourseActivity.objects.filter(
            Q(course=course) | Q(assigned_courses=course),
            activity_type__name__iexact="quiz",
        ).distinct(),
        id=activity_id,
    )

    if request.method == "GET":
        return Response(CourseActivitySerializer(activity, context={"request": request}).data)

    if request.user.role != "instructor" or activity.course.instructor_id != request.user.id:
        return Response({"error": "Only instructor can modify exam or quiz."}, status=403)

    if request.method in {"PUT", "PATCH"}:
        serializer = CourseActivitySerializer(
            activity,
            data=request.data,
            partial=(request.method == "PATCH"),
            context={"request": request},
        )
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data)
        return Response(serializer.errors, status=400)

    activity.delete()
    return Response({"message": "Exam deleted successfully"}, status=200)


@api_view(["PATCH"])
@permission_classes([IsAuthenticated])
def exam_quiz_settings(request, course_id, activity_id):
    course = _get_course_with_access(course_id, request.user)
    if not course:
        return Response({"error": "Course not found or access denied"}, status=404)
    activity = get_object_or_404(
        CourseActivity.objects.filter(course=course, activity_type__name__iexact="quiz"),
        id=activity_id,
    )
    if request.user.role != "instructor" or activity.course.instructor_id != request.user.id:
        return Response({"error": "Only instructor can update exam settings."}, status=403)

    show_score_immediately = _as_bool(request.data.get("show_score_immediately", activity.show_score_immediately), activity.show_score_immediately)
    allow_answer_review = _as_bool(request.data.get("allow_answer_review", activity.allow_answer_review), activity.allow_answer_review)
    apply_to_past_attempts = _as_bool(request.data.get("apply_to_past_attempts", False), False)

    with transaction.atomic():
        if not apply_to_past_attempts:
            attempts_qs = QuizAttempt.objects.select_for_update().filter(quiz=activity)
            for attempt in attempts_qs:
                snapshot = attempt.visibility_snapshot or {}
                snapshot.setdefault("show_score_immediately", bool(activity.show_score_immediately))
                snapshot.setdefault("allow_answer_review", bool(activity.allow_answer_review))
                attempt.visibility_snapshot = snapshot
                attempt.save(update_fields=["visibility_snapshot"])

        activity.show_score_immediately = show_score_immediately
        activity.allow_answer_review = allow_answer_review
        activity.save(update_fields=["show_score_immediately", "allow_answer_review"])

        if apply_to_past_attempts:
            attempts_qs = QuizAttempt.objects.select_for_update().filter(quiz=activity)
            for attempt in attempts_qs:
                attempt.visibility_snapshot = {
                    "show_score_immediately": show_score_immediately,
                    "allow_answer_review": allow_answer_review,
                }
                attempt.save(update_fields=["visibility_snapshot"])

    return Response(
        {
            "id": activity.id,
            "show_score_immediately": activity.show_score_immediately,
            "allow_answer_review": activity.allow_answer_review,
            "apply_to_past_attempts": apply_to_past_attempts,
        }
    )


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def exam_quiz_submission_reviews(request, course_id, activity_id):
    course = _get_course_with_access(course_id, request.user)
    if not course:
        return Response({"error": "Course not found or access denied"}, status=404)
    activity = get_object_or_404(
        CourseActivity.objects.filter(course=course, activity_type__name__iexact="quiz"),
        id=activity_id,
    )
    if request.user.role != "instructor" or activity.course.instructor_id != request.user.id:
        return Response({"error": "Only instructor can review exam submissions."}, status=403)

    attempts = (
        QuizAttempt.objects.filter(quiz=activity, submitted_at__isnull=False)
        .select_related("student", "graded_by")
        .prefetch_related("answer_records", "score_audits__actor")
        .order_by("-submitted_at", "-started_at")
    )
    return Response(
        {
            "activity_id": activity.id,
            "title": activity.title,
            "attempts": [
                {
                    "attempt_id": attempt.id,
                    "student_id": attempt.student_id,
                    "student_name": getattr(attempt.student, "username", "Student"),
                    "status": attempt.status,
                    "score": float(attempt.score or 0),
                    "override_score": float(attempt.override_score) if attempt.override_score is not None else None,
                    "display_score": float(attempt.override_score if attempt.override_score is not None else attempt.score or 0),
                    "total_points": float(attempt.total_points or 0),
                    "submitted_at": attempt.submitted_at,
                    "graded_at": attempt.graded_at,
                    "graded_by_name": getattr(attempt.graded_by, "username", None),
                    "is_overridden": bool(attempt.is_overridden),
                }
                for attempt in attempts
            ],
        }
    )


@api_view(["GET", "PATCH"])
@permission_classes([IsAuthenticated])
def exam_quiz_submission_review_detail(request, course_id, activity_id, attempt_id):
    course = _get_course_with_access(course_id, request.user)
    if not course:
        return Response({"error": "Course not found or access denied"}, status=404)
    activity = get_object_or_404(
        CourseActivity.objects.filter(course=course, activity_type__name__iexact="quiz"),
        id=activity_id,
    )
    if request.user.role != "instructor" or activity.course.instructor_id != request.user.id:
        return Response({"error": "Only instructor can review exam submissions."}, status=403)

    attempt = get_object_or_404(
        QuizAttempt.objects.select_related("student", "graded_by").prefetch_related("answer_records", "score_audits__actor"),
        id=attempt_id,
        quiz=activity,
    )
    _ensure_attempt_answer_records(attempt)

    if request.method == "GET":
        return Response(_serialize_attempt_review_payload(attempt))

    answer_updates = request.data.get("answers")
    override_total_score = request.data.get("override_total_score", None)
    note = str(request.data.get("note") or "").strip()

    with transaction.atomic():
        # Lock only the attempt row here. PostgreSQL rejects FOR UPDATE queries
        # that include the nullable graded_by outer join.
        attempt = QuizAttempt.objects.select_for_update().get(id=attempt.id)
        _ensure_attempt_answer_records(attempt)
        existing_answers = {str(row.question_id): row for row in attempt.answer_records.all()}
        if isinstance(answer_updates, list):
            for row in answer_updates:
                if not isinstance(row, dict):
                    continue
                question_id = str(row.get("question_id") or "")
                answer_record = existing_answers.get(question_id)
                if not answer_record:
                    continue
                previous_score = answer_record.override_score
                if previous_score is None:
                    previous_score = answer_record.manual_score if answer_record.manual_score is not None else answer_record.auto_score

                if "score" in row:
                    score_value = row.get("score")
                    if score_value in ("", None):
                        answer_record.override_score = None
                        answer_record.manual_score = None if str(answer_record.question_type or "").lower() == "essay" else answer_record.manual_score
                    else:
                        try:
                            parsed_score = float(score_value)
                        except (TypeError, ValueError):
                            return Response({"error": f"Invalid score for question {question_id}."}, status=400)
                        if parsed_score < 0 or parsed_score > float(answer_record.max_points or 0):
                            return Response({"error": f"Score for question {question_id} must be between 0 and max points."}, status=400)
                        if str(answer_record.question_type or "").lower() == "essay":
                            answer_record.manual_score = parsed_score
                            answer_record.override_score = None
                        else:
                            answer_record.override_score = parsed_score
                    answer_record.status = QuizAttempt.STATUS_GRADED

                if "feedback" in row:
                    answer_record.feedback = str(row.get("feedback") or "")

                answer_record.save()

                current_score = answer_record.override_score
                if current_score is None:
                    current_score = answer_record.manual_score if answer_record.manual_score is not None else answer_record.auto_score
                if previous_score != current_score or ("feedback" in row):
                    QuizAttemptScoreAudit.objects.create(
                        attempt=attempt,
                        actor=request.user,
                        question_id=question_id,
                        previous_score=previous_score,
                        new_score=current_score,
                        note=note or f"Updated question {question_id}",
                    )

        attempt = _recompute_attempt_from_answer_records(
            attempt,
            actor=request.user,
            mark_override=bool(isinstance(answer_updates, list) and any(str(item.get("question_id") or "") for item in answer_updates if isinstance(item, dict))),
            audit_note=note,
        )
        if override_total_score not in (None, ""):
            try:
                parsed_total_override = float(override_total_score)
            except (TypeError, ValueError):
                return Response({"error": "override_total_score must be numeric."}, status=400)
            if parsed_total_override < 0 or parsed_total_override > float(attempt.total_points or 0):
                return Response({"error": "override_total_score must be between 0 and total points."}, status=400)
            previous_display_score = attempt.override_score if attempt.override_score is not None else attempt.score
            attempt.override_score = parsed_total_override
            attempt.is_overridden = True
            attempt.graded_at = timezone.now()
            attempt.graded_by = request.user
            attempt.save(update_fields=["override_score", "is_overridden", "graded_at", "graded_by"])
            QuizAttemptScoreAudit.objects.create(
                attempt=attempt,
                actor=request.user,
                previous_score=previous_display_score,
                new_score=parsed_total_override,
                note=note or "Updated total override score",
            )
            _update_attempt_submission_record(attempt, feedback_text="Instructor adjusted this submission score.")

    attempt = QuizAttempt.objects.select_related("student", "graded_by").prefetch_related("answer_records", "score_audits__actor").get(id=attempt.id)
    return Response(_serialize_attempt_review_payload(attempt))


@api_view(["GET", "POST"])
@permission_classes([IsAuthenticated])
def classwork_draft(request, course_id):
    course = get_object_or_404(Course, id=course_id)
    if request.user.role != "instructor" or course.instructor_id != request.user.id:
        return Response({"error": "Only instructor can access drafts."}, status=403)

    draft = (
        ClassworkDraft.objects.filter(course=course, instructor=request.user)
        .order_by("-updated_at")
        .first()
    )

    if request.method == "GET":
        if not draft:
            return Response({"detail": "No draft found."}, status=404)
        return Response(ClassworkDraftSerializer(draft).data)

    serializer = ClassworkDraftSerializer(
        draft,
        data=request.data,
        partial=True,
        context={"request": request},
    )
    if serializer.is_valid():
        saved = serializer.save(course=course, instructor=request.user)
        return Response(ClassworkDraftSerializer(saved).data, status=200)
    return Response(serializer.errors, status=400)


@api_view(["POST"])
@permission_classes([IsAuthenticated])
@parser_classes([MultiPartParser, FormParser])
def classwork_import_questions(request, course_id):
    course = get_object_or_404(Course, id=course_id)
    if request.user.role != "instructor" or course.instructor_id != request.user.id:
        return Response({"error": "Only instructor can import questions."}, status=403)

    upload = request.FILES.get("file")
    if not upload:
        return Response({"error": "file is required (.docx or .pdf)"}, status=400)

    file_name = upload.name or ""
    lowered = file_name.lower()
    text = ""

    if lowered.endswith(".docx"):
        if docx is None:
            return Response({"error": "python-docx is not installed on server."}, status=500)
        document = docx.Document(upload)
        text = _extract_docx_text(document)
    elif lowered.endswith(".pdf"):
        if pdfplumber is None:
            return Response({"error": "pdfplumber is not installed on server."}, status=500)
        with pdfplumber.open(upload) as pdf:
            chunks = []
            for page in pdf.pages:
                page_text = page.extract_text() or ""
                if page_text:
                    chunks.append(page_text)
            text = "\n".join(chunks)
    else:
        return Response({"error": "Unsupported file format. Use .docx or .pdf"}, status=400)

    mode = str(request.data.get("mode") or request.query_params.get("mode") or "balanced").strip().lower()
    if mode not in {"strict", "balanced", "tolerant"}:
        mode = "balanced"

    parsed = _parse_questions_from_plain_text(text, mode=mode)
    sections = parsed.get("sections", [])
    parse_warnings = list(parsed.get("warnings") or [])
    total_questions = sum(len(section.get("questions", [])) for section in sections)
    if total_questions == 0:
        parse_warnings.append("No questions detected automatically. Please use manual correction.")
    confidence_score = int(parsed.get("confidence_score", 0))

    if mode == "strict" and parsed.get("errors"):
        return Response(
            {
                "error": "Strict import validation failed.",
                "mode": mode,
                "errors": parsed.get("errors", []),
                "warnings": parse_warnings,
                "data": parsed.get("data", {"sections": sections}),
                "debug": parsed.get("debug", {}),
                "confidence_score": confidence_score,
            },
            status=400,
        )

    logger.debug(
        "Exam import summary: source=%s sections=%s questions=%s answers=%s warnings=%s",
        file_name,
        len(sections),
        total_questions,
        len(parsed.get("answer_map", {})),
        len(parse_warnings),
    )
    return Response(
        {
            "sections": sections,
            "source": file_name,
            "warnings": parse_warnings,
            "info": list(parsed.get("notes") or []),
            "raw_preview": {
                "text_excerpt": text[:5000],
                "line_count": parsed.get("line_count", 0),
                "answer_map": parsed.get("answer_map", {}),
            },
            "confidence": round(confidence_score / 100.0, 2),
            "confidence_score": confidence_score,
            "debug": parsed.get("debug", {}),
            "mode": mode,
        },
        status=200,
    )


@api_view(["GET", "POST"])
@permission_classes([IsAuthenticated])
def question_bank_items(request, course_id):
    course = get_object_or_404(Course, id=course_id)
    if request.user.role != "instructor" or course.instructor_id != request.user.id:
        return Response({"error": "Only instructor can manage question bank."}, status=403)

    if request.method == "GET":
        queryset = QuestionBankItem.objects.filter(instructor=request.user).order_by("-updated_at")
        query = str(request.query_params.get("query", "") or "").strip()
        topic = str(request.query_params.get("topic", "") or "").strip()
        difficulty = str(request.query_params.get("difficulty", "") or "").strip()

        if topic:
            queryset = queryset.filter(topic__icontains=topic)
        if difficulty:
            queryset = queryset.filter(difficulty__iexact=difficulty)
        if query:
            queryset = queryset.filter(
                Q(topic__icontains=query)
                | Q(question_data__question_text__icontains=query)
            )
        serializer = QuestionBankItemSerializer(queryset[:300], many=True)
        return Response(serializer.data)

    serializer = QuestionBankItemSerializer(data=request.data, context={"request": request})
    if serializer.is_valid():
        item = serializer.save(instructor=request.user, course=course)
        return Response(QuestionBankItemSerializer(item).data, status=201)
    return Response(serializer.errors, status=400)


@api_view(["PATCH", "DELETE"])
@permission_classes([IsAuthenticated])
def question_bank_item_detail(request, course_id, item_id):
    course = get_object_or_404(Course, id=course_id)
    if request.user.role != "instructor" or course.instructor_id != request.user.id:
        return Response({"error": "Only instructor can manage question bank."}, status=403)

    item = get_object_or_404(QuestionBankItem, id=item_id, instructor=request.user)

    if request.method == "PATCH":
        serializer = QuestionBankItemSerializer(item, data=request.data, partial=True, context={"request": request})
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data)
        return Response(serializer.errors, status=400)

    item.delete()
    return Response(status=204)


@api_view(["GET", "POST"])
@permission_classes([IsAuthenticated])
def quiz_security_events(request, course_id, activity_id):
    course = _get_course_with_access(course_id, request.user)
    if not course:
        return Response({"error": "Course not found or access denied"}, status=404)
    quiz = get_object_or_404(CourseActivity, id=activity_id, course=course, activity_type__name__iexact="quiz")

    if request.method == "GET":
        if request.user.role != "instructor" or quiz.course.instructor_id != request.user.id:
            return Response({"error": "Only instructor can view security events."}, status=403)
        events = QuizSecurityEvent.objects.filter(quiz=quiz).select_related("student", "attempt")[:500]
        return Response(QuizSecurityEventSerializer(events, many=True).data)

    if request.user.role != "student":
        return Response({"error": "Only students can submit security events."}, status=403)

    event_type = request.data.get("event_type")
    details = request.data.get("details") if isinstance(request.data.get("details"), dict) else {}
    attempt_id = request.data.get("attempt_id") or request.data.get("attempt")
    if not attempt_id:
        return Response({"error": "Missing attempt_id"}, status=400)
    with transaction.atomic():
        locked_attempt = QuizAttempt.objects.select_for_update().filter(
            id=attempt_id,
            quiz=quiz,
            student=request.user,
        ).first()
        if not locked_attempt:
            return Response({"error": "Attempt not found"}, status=404)

        if locked_attempt.is_locked:
            logger.info(
                "Rejected security event for locked attempt.",
                extra={"attempt_id": locked_attempt.id, "quiz_id": quiz.id, "user_id": request.user.id},
            )
            return Response(
                {
                    "error": "Attempt is locked; security events are closed.",
                    "attempt_id": locked_attempt.id,
                    "attempt_locked": True,
                    "attempt_submitted": bool(locked_attempt.submitted_at),
                },
                status=409,
            )

        if locked_attempt.submitted_at:
            submitted_age = (
                int(max((timezone.now() - locked_attempt.submitted_at).total_seconds(), 0))
                if locked_attempt.submitted_at
                else 0
            )
            # Race reconciliation:
            # If a threshold event arrives concurrently with submit, reconcile to a locked force-submit
            # so non-zero grading cannot survive the threshold boundary.
            if (
                not locked_attempt.is_locked
                and int(locked_attempt.suspicious_events or 0) >= (QUIZ_SECURITY_FORCE_SUBMIT_THRESHOLD - 1)
                and submitted_age <= QUIZ_SECURITY_RACE_RECONCILE_SECONDS
            ):
                _force_submit_locked_attempt(locked_attempt, quiz, reason="security_threshold_race_reconcile")
                logger.warning(
                    "Reconciled concurrent submit/security race by force-locking attempt.",
                    extra={
                        "attempt_id": locked_attempt.id,
                        "quiz_id": quiz.id,
                        "user_id": request.user.id,
                        "violation_count": int(locked_attempt.suspicious_events or 0),
                    },
                )
                return Response(
                    {
                        "error": "Attempt locked after concurrent security threshold event.",
                        "attempt_id": locked_attempt.id,
                        "attempt_locked": True,
                        "attempt_submitted": True,
                        "force_submit": True,
                    },
                    status=409,
                )
            return Response(
                {
                    "error": "Attempt already submitted",
                    "attempt_id": locked_attempt.id,
                    "attempt_submitted": True,
                },
                status=409,
            )

        current_event_count = QuizSecurityEvent.objects.filter(attempt=locked_attempt).count()
        if current_event_count >= QUIZ_SECURITY_MAX_EVENTS_PER_ATTEMPT:
            logger.warning(
                "Security event rate cap reached; dropping event.",
                extra={
                    "attempt_id": locked_attempt.id,
                    "quiz_id": quiz.id,
                    "user_id": request.user.id,
                    "event_count": current_event_count,
                },
            )
            return Response(
                {
                    "error": "Security event limit reached for this attempt.",
                    "attempt_id": locked_attempt.id,
                    "attempt_locked": bool(locked_attempt.is_locked),
                    "attempt_submitted": bool(locked_attempt.submitted_at),
                    "violation_count": int(locked_attempt.suspicious_events or 0),
                },
                status=429,
            )

        serializer = QuizSecurityEventSerializer(
            data={
                "quiz": quiz.id,
                "attempt": locked_attempt.id,
                "student": request.user.id,
                "event_type": event_type,
                "details": details,
            }
        )
        serializer.is_valid(raise_exception=True)
        event = serializer.save(student=request.user)

        now_ts = timezone.now()
        gap_seconds = int(max((now_ts - locked_attempt.last_activity_at).total_seconds(), 0)) if locked_attempt.last_activity_at else 0
        if gap_seconds > QUIZ_SECURITY_INACTIVITY_SECONDS:
            details = dict(details or {})
            details["inactivity_gap_seconds"] = gap_seconds
            details["inactivity_violation"] = True
            event.details = details
            event.save(update_fields=["details"])
            locked_attempt.suspicious_events = int(locked_attempt.suspicious_events or 0) + 1

        locked_attempt.suspicious_events = int(locked_attempt.suspicious_events or 0) + 1
        locked_attempt.last_activity_at = now_ts
        locked_attempt.save(update_fields=["suspicious_events", "last_activity_at"])
        violation_count = int(locked_attempt.suspicious_events or 0)

        attempt_locked = False
        attempt_submitted = False
        if violation_count >= QUIZ_SECURITY_FORCE_SUBMIT_THRESHOLD:
            _force_submit_locked_attempt(locked_attempt, quiz, reason="security_threshold")
            logger.warning(
                "Violation threshold reached; attempt force-submitted.",
                extra={
                    "attempt_id": locked_attempt.id,
                    "quiz_id": quiz.id,
                    "user_id": request.user.id,
                    "violation_count": violation_count,
                },
            )
            attempt_locked = True
            attempt_submitted = True
        else:
            attempt_locked = bool(locked_attempt.is_locked)
            attempt_submitted = bool(locked_attempt.submitted_at)

        force_submit = bool(attempt_locked or violation_count >= QUIZ_SECURITY_FORCE_SUBMIT_THRESHOLD)
        payload = QuizSecurityEventSerializer(event).data
        payload.update(
            {
                "violation_count": violation_count,
                "force_submit": force_submit,
                "force_submit_threshold": QUIZ_SECURITY_FORCE_SUBMIT_THRESHOLD,
                "attempt_locked": attempt_locked,
                "attempt_submitted": attempt_submitted,
            }
        )
        return Response(payload, status=201)

@api_view(["POST"])
@permission_classes([IsAuthenticated])
@parser_classes([MultiPartParser, FormParser])
def submit_task(request, course_id, activity_id):
    try:
        activity = CourseActivity.objects.get(id=activity_id, course_id=course_id)
    except CourseActivity.DoesNotExist:
        return Response({"error": "Activity not found"}, status=404)
    if not _allow(STUDENT_ROLE_PERMISSION, request.user):
        return Response({"error": "Only students can submit classwork."}, status=403)
    if not _allow(COURSE_ENROLLED_PERMISSION, request.user, activity):
        return Response({"error": "Activity not found or access denied"}, status=404)

    logger.debug(
        "submit_task request received.",
        extra={
            "course_id": course_id,
            "activity_id": activity_id,
            "user_id": getattr(request.user, "id", None),
            "payload_keys": list(request.data.keys()),
            "file_count": len(request.FILES.getlist("files")) if hasattr(request.FILES, "getlist") else len(request.FILES or []),
        },
    )

    data = request.data.copy()
    files = request.FILES.getlist("files")
    is_late_now = bool(activity.due_date and timezone.now() > activity.due_date)

    if is_late_now and not bool(activity.allow_late_submissions):
        return Response(
            {"error": "Late submissions are disabled for this classwork."},
            status=400,
        )

    if _is_quiz_activity(activity):
        raw_answers = data.get("answers")
        if raw_answers in (None, "") and data.get("text_answer"):
            try:
                parsed_text = json.loads(data.get("text_answer"))
                raw_answers = parsed_text.get("answers")
            except (json.JSONDecodeError, TypeError, ValueError):
                raw_answers = None
        if raw_answers in (None, "", [], {}):
            return Response(
                {"error": "Quiz submissions require structured answers. Use the quiz submit endpoint."},
                status=400,
            )

    # check if submission exists FIRST
    existing = ActivitySubmission.objects.filter(
        activity=activity,
        student=request.user
    ).first()

    # -------------------------
    # UPDATE EXISTING
    # -------------------------
    if existing:
        serializer = ActivitySubmissionSerializer(
            existing,
            data=data,
            partial=True,
            context={"request": request}
        )

        if serializer.is_valid():
            submission = serializer.save()
            if submission.status != "submitted":
                submission.status = "submitted"
                submission.save(update_fields=["status"])
            if is_late_now and not submission.is_late:
                submission.is_late = True
                submission.save(update_fields=["is_late"])

            # replace files
            if files:
                submission.attachments.all().delete()
                for f in files:
                    SubmissionAttachment.objects.create(
                        submission=submission,
                        file=f
                    )

            dispatch_event("assignment_submitted", submission=submission, actor=request.user)

            return Response(
                ActivitySubmissionSerializer(submission, context={"request": request}).data
            )

        return Response(serializer.errors, status=400)


    # -------------------------
    # CREATE NEW
    # -------------------------
    data["student"] = request.user.id
    data["activity"] = activity.id

    serializer = ActivitySubmissionSerializer(data=data, context={"request": request})

    if serializer.is_valid():
        submission = serializer.save()
        if submission.status != "submitted":
            submission.status = "submitted"
            submission.save(update_fields=["status"])
        if is_late_now and not submission.is_late:
            submission.is_late = True
            submission.save(update_fields=["is_late"])

        for f in files:
            SubmissionAttachment.objects.create(submission=submission, file=f)

        dispatch_event("assignment_submitted", submission=submission, actor=request.user)

        return Response(
            ActivitySubmissionSerializer(submission, context={"request": request}).data,
            status=201
        )

    return Response(serializer.errors, status=400)    

@api_view(["DELETE"])
@permission_classes([IsAuthenticated])
def unsubmit_task(request, submission_id):
    try:
        submission = ActivitySubmission.objects.get(
            id=submission_id,
            student=request.user
        )
    except ActivitySubmission.DoesNotExist:
        return Response({"error": "Submission not found"}, status=404)

    submission.delete()
    return Response({"message": "Unsubmitted successfully"})


@api_view(["POST"])
@permission_classes([IsAuthenticated])
@parser_classes([MultiPartParser, FormParser])
def upload_attachment(request, submission_id):

    try:
        submission = ActivitySubmission.objects.get(
            id=submission_id,
            student=request.user
        )
    except ActivitySubmission.DoesNotExist:
        return Response({"error": "Submission not found"}, status=404)

    file = request.FILES.get("file")
    if not file:
        return Response({"error": "No file provided"}, status=400)

    attachment = SubmissionAttachment.objects.create(
        submission=submission,
        file=file
    )

    return Response({"message": "File uploaded"}, status=201)


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def submissions_list(request, course_id, activity_id):
    try:
        activity = CourseActivity.objects.get(id=activity_id, course_id=course_id)
    except CourseActivity.DoesNotExist:
        return Response({"error": "Activity not found"}, status=404)

    if request.user.role == "student":
        submissions = ActivitySubmission.objects.filter(activity=activity, student=request.user)
    else:  # instructor
        submissions = ActivitySubmission.objects.filter(activity=activity)

    serializer = ActivitySubmissionSerializer(submissions, many=True, context={"request": request})
    return Response(serializer.data)


@api_view(["PUT", "PATCH"])
@permission_classes([IsAuthenticated])
def grade_submission(request, course_id, activity_id, submission_id):
    try:
        submission = ActivitySubmission.objects.get(id=submission_id, activity__id=activity_id, activity__course_id=course_id)
    except ActivitySubmission.DoesNotExist:
        return Response({"error": "Submission not found"}, status=404)
    if not _allow(INSTRUCTOR_ROLE_PERMISSION, request.user):
        return Response({"error": "Only instructor can grade submissions"}, status=403)
    if not _allow(COURSE_OWNER_PERMISSION, request.user, submission.activity):
        return Response({"error": "Submission not found or access denied"}, status=404)

    serializer = ActivitySubmissionSerializer(submission, data=request.data, partial=True, context={"request": request})
    if serializer.is_valid():
        updated = serializer.save()
        if updated.grade is not None and updated.status != "graded":
            ActivitySubmission.objects.filter(pk=updated.pk).update(status="graded")
            updated.refresh_from_db()
        dispatch_event("grade_posted", submission=updated, actor=request.user)
        return Response(ActivitySubmissionSerializer(updated, context={"request": request}).data)
    return Response(serializer.errors, status=400)


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def student_enrolled_courses(request):

    if request.user.role != "student":
        return Response({"error": "Only students can access this."}, status=403)

    courses = request.user.enrolled_courses.all()

    serializer = CourseSerializer(
        courses,
        many=True,
        context={"request": request}
    )

    return Response(serializer.data)


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def student_course_detail(request, course_id):
    """
    Allows a student to view course details if they are enrolled.
    """
    try:
        course = Course.objects.get(id=course_id, students=request.user)
    except Course.DoesNotExist:
        return Response({"error": "Course not found or access denied"}, status=404)

    serializer = CourseSerializer(course, context={"request": request})
    return Response(serializer.data)


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def student_course_lessons(request, course_id):
    try:
        course = Course.objects.get(id=course_id, students=request.user)
    except Course.DoesNotExist:
        return Response({"error": "Course not found or access denied"}, status=404)

    lessons = Lesson.objects.filter(course=course)
    serializer = LessonSerializer(lessons, many=True, context=_build_lesson_serializer_context(request, lessons))
    return Response(serializer.data)


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def course_progress(request, course_id):
    try:
        course = Course.objects.get(id=course_id, students=request.user)
    except Course.DoesNotExist:
        return Response({"error": "Course not found or access denied"}, status=404)

    total_lessons = Lesson.objects.filter(course=course).count()
    completed_lesson_ids = list(
        LessonCompletion.objects.filter(student=request.user, lesson__course=course).values_list("lesson_id", flat=True)
    )
    completed_lessons = len(completed_lesson_ids)
    percentage = round((completed_lessons / total_lessons) * 100) if total_lessons > 0 else 0

    return Response(
        {
            "course_id": course.id,
            "completed_lessons": completed_lessons,
            "total_lessons": total_lessons,
            "completed_lesson_ids": completed_lesson_ids,
            "percentage": percentage,
            "progress": percentage,
        }
    )


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def complete_lesson(request, course_id, lesson_id):
    try:
        course = Course.objects.get(id=course_id, students=request.user)
    except Course.DoesNotExist:
        return Response({"error": "Course not found or access denied"}, status=404)

    lesson = get_object_or_404(Lesson, id=lesson_id, course=course)
    completion, created = LessonCompletion.objects.get_or_create(lesson=lesson, student=request.user)

    return Response(
        {
            "lesson_id": lesson.id,
            "course_id": course.id,
            "completed": True,
            "created": created,
            "completed_at": completion.completed_at,
        },
        status=201 if created else 200,
    )


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def student_course_activities(request, course_id):
    try:
        course = Course.objects.get(id=course_id, students=request.user)
    except Course.DoesNotExist:
        return Response({"error": "Course not found or access denied"}, status=404)

    activities = CourseActivity.objects.filter(course=course).exclude(
        Q(activity_type__name__iexact="quiz") & Q(publish_state=CourseActivity.PUBLISH_STATE_DRAFT)
    )
    serializer = CourseActivitySerializer(activities, many=True, context={"request": request})
    return Response(serializer.data)


def _is_quiz_activity(activity):
    return str(getattr(activity.activity_type, "name", "") or "").lower() == "quiz"


def _normalize_submitted_answers(raw_answers):
    if isinstance(raw_answers, dict):
        normalized = []
        for question_id, answer in raw_answers.items():
            normalized.append(
                {
                    "question_id": question_id,
                    "answer": answer,
                }
            )
        raw_answers = normalized

    if not isinstance(raw_answers, list):
        raise ValueError("answers must be a list or object")

    normalized = []
    for item in raw_answers:
        if not isinstance(item, dict):
            raise ValueError("each answer must be an object")
        question_id = item.get("question_id")
        answer = item.get("answer")
        answer_items = item.get("answer_items")
        if question_id in (None, ""):
            raise ValueError("question_id is required")
        if isinstance(answer, list):
            parsed_items = [str(value or "").strip() for value in answer if str(value or "").strip()]
            if not parsed_items:
                raise ValueError("answer cannot be empty")
            normalized.append(
                {
                    "question_id": str(question_id),
                    "answer": ", ".join(parsed_items),
                    "answer_items": parsed_items,
                }
            )
            continue
        if answer is None or str(answer).strip() == "":
            raise ValueError("answer cannot be empty")
        parsed_answer_items = []
        if isinstance(answer_items, list):
            parsed_answer_items = [str(value or "").strip() for value in answer_items if str(value or "").strip()]
        normalized.append(
            {
                "question_id": str(question_id),
                "answer": str(answer).strip(),
                "answer_items": parsed_answer_items,
            }
        )

    if not normalized:
        raise ValueError("at least one answer is required")
    return normalized


def _load_quiz_questions_for_runtime(activity):
    questions = activity.quiz_questions if isinstance(activity.quiz_questions, list) else []
    if not questions and isinstance(getattr(activity, "quiz_sections", None), list):
        questions = _flatten_sections_to_questions(activity.quiz_sections)
    if not questions and activity.description:
        try:
            parsed = json.loads(activity.description)
            questions = parsed.get("questions", parsed) if isinstance(parsed, dict) else parsed
        except (json.JSONDecodeError, TypeError, ValueError):
            questions = []
    return questions if isinstance(questions, list) else []


def _flatten_sections_to_questions(sections):
    if not isinstance(sections, list):
        return []
    flat = []
    question_id = 1
    for section in sections:
        if not isinstance(section, dict):
            continue
        section_id = section.get("id")
        section_title = section.get("title")
        for question in section.get("questions", []):
            if not isinstance(question, dict):
                continue
            item = dict(question)
            item["id"] = question_id
            item["section_id"] = section_id
            item["section_title"] = section_title
            flat.append(item)
            question_id += 1
    return flat


def _normalize_text_token(value):
    return re.sub(r"\s+", " ", str(value or "").strip()).lower()


def _split_enumeration_answer_parts(value):
    if isinstance(value, list):
        raw_parts = value
    else:
        raw_parts = re.split(r"[,;\n]+", str(value or ""))
    return [str(part or "").strip() for part in raw_parts if str(part or "").strip()]


def _normalize_enumeration_items(item):
    raw_items = item.get("enumeration_items")
    normalized = []

    if isinstance(raw_items, list):
        for index, raw_item in enumerate(raw_items):
            if isinstance(raw_item, dict):
                answer = str(raw_item.get("answer") or raw_item.get("text") or "").strip()
                alternatives = raw_item.get("alternatives") or raw_item.get("synonyms") or []
                if not isinstance(alternatives, list):
                    alternatives = []
                alt_values = [str(value or "").strip() for value in alternatives if str(value or "").strip()]
                try:
                    points = float(raw_item.get("points", 0) or 0)
                except (TypeError, ValueError):
                    points = 0.0
            else:
                answer = str(raw_item or "").strip()
                alt_values = []
                points = 0.0

            if answer:
                normalized.append(
                    {
                        "id": index + 1,
                        "answer": answer,
                        "alternatives": alt_values,
                        "points": max(points, 0.0),
                    }
                )

    if normalized:
        return normalized

    fallback_answers = item.get("enumeration_answers")
    if not isinstance(fallback_answers, list):
        fallback_answers = _split_enumeration_answer_parts(item.get("correct_answer", ""))

    for index, answer in enumerate(fallback_answers):
        answer_text = str(answer or "").strip()
        if answer_text:
            normalized.append(
                {
                    "id": index + 1,
                    "answer": answer_text,
                    "alternatives": [],
                    "points": 0.0,
                }
            )
    return normalized


def _grade_enumeration_question(question, submitted_answer, submitted_answer_items=None):
    enumeration_items = _normalize_enumeration_items(question)
    expected_count = len(enumeration_items)
    max_points = float(question.get("points", 1) or 1)
    scoring_mode = str(question.get("enumeration_scoring_mode") or "partial").strip().lower()
    if scoring_mode not in {"strict", "partial", "percentage"}:
        scoring_mode = "partial"
    points_mode = str(question.get("enumeration_points_mode") or "equal").strip().lower()
    if points_mode not in {"equal", "custom"}:
        points_mode = "equal"

    if submitted_answer_items:
        submitted_parts = _split_enumeration_answer_parts(submitted_answer_items)
    else:
        submitted_parts = _split_enumeration_answer_parts(submitted_answer)
    evaluated_parts = submitted_parts[:expected_count] if expected_count > 0 else submitted_parts

    if expected_count == 0:
        return {
            "is_correct": False,
            "points_earned": 0.0,
            "submitted_items": evaluated_parts,
            "answer_feedback": [],
            "matched_count": 0,
            "expected_count": 0,
            "limit_applied": 0,
        }

    if points_mode == "custom":
        item_points = [max(float(item.get("points", 0) or 0), 0.0) for item in enumeration_items]
        if sum(item_points) <= 0:
            item_points = [max_points / expected_count] * expected_count
    else:
        item_points = [max_points / expected_count] * expected_count

    remaining_expected = list(range(expected_count))
    matched_expected = set()
    answer_feedback = []
    matched_count = 0
    matched_points = 0.0

    for submitted in evaluated_parts:
        normalized_submitted = _normalize_text_token(submitted)
        matched_index = None
        for expected_index in remaining_expected:
            item = enumeration_items[expected_index]
            accepted_tokens = [_normalize_text_token(item.get("answer", ""))]
            accepted_tokens.extend(_normalize_text_token(value) for value in (item.get("alternatives") or []))
            accepted_tokens = [token for token in accepted_tokens if token]
            if normalized_submitted and normalized_submitted in accepted_tokens:
                matched_index = expected_index
                break

        if matched_index is not None:
            remaining_expected.remove(matched_index)
            matched_expected.add(matched_index)
            matched_count += 1
            awarded_points = float(item_points[matched_index] or 0)
            matched_points += awarded_points
            answer_feedback.append(
                {
                    "submitted": submitted,
                    "is_correct": True,
                    "expected": enumeration_items[matched_index].get("answer", ""),
                    "accepted_variations": enumeration_items[matched_index].get("alternatives", []),
                    "awarded_points": round(awarded_points, 2),
                }
            )
        else:
            answer_feedback.append(
                {
                    "submitted": submitted,
                    "is_correct": False,
                    "expected": "",
                    "accepted_variations": [],
                    "awarded_points": 0.0,
                }
            )

    for expected_index in range(expected_count):
        if expected_index in matched_expected:
            continue
        answer_feedback.append(
            {
                "submitted": "",
                "is_correct": False,
                "expected": enumeration_items[expected_index].get("answer", ""),
                "accepted_variations": enumeration_items[expected_index].get("alternatives", []),
                "awarded_points": 0.0,
                "missing": True,
            }
        )

    if scoring_mode == "strict":
        points_earned = max_points if matched_count == expected_count and len(evaluated_parts) == expected_count else 0.0
    elif scoring_mode == "percentage":
        points_earned = (matched_count / expected_count) * max_points
    else:
        points_earned = matched_points

    return {
        "is_correct": matched_count == expected_count and len(evaluated_parts) == expected_count,
        "points_earned": round(points_earned, 2),
        "submitted_items": evaluated_parts,
        "answer_feedback": answer_feedback,
        "matched_count": matched_count,
        "expected_count": expected_count,
        "limit_applied": expected_count,
    }


def _validate_and_normalize_quiz_questions(questions):
    if not isinstance(questions, list) or len(questions) == 0:
        return []

    normalized = []
    for index, item in enumerate(questions):
        if not isinstance(item, dict):
            return []

        question_id = item.get("id", index + 1)
        question_text = str(item.get("question_text") or item.get("text") or "").strip()
        question_type = _route_question_type(
            str(item.get("type") or "identification").strip().lower(),
            question_text,
            has_options=bool(item.get("options") or item.get("choices")),
        )
        if question_type in {"mcq", "multiple choice"}:
            question_type = "multiple_choice"
        if question_type in {"true_false", "truefalse", "tf"}:
            question_type = "true_false"
        if question_type in {"short", "short_answer"}:
            question_type = "short_answer"
        if question_type == "matching_type":
            question_type = "matching"
        if question_type not in {
            "multiple_choice",
            "true_false",
            "short_answer",
            "identification",
            "essay",
            "coding",
            "file_upload",
            "matching",
            "enumeration",
        }:
            return []
        if not question_text:
            return []

        try:
            points = float(item.get("points", 1) or 1)
        except (TypeError, ValueError):
            points = 1.0
        if points <= 0:
            points = 1.0

        raw_options = item.get("options") or []
        if not raw_options and isinstance(item.get("choices"), list):
            raw_options = item.get("choices") or []
        normalized_options = []
        for option_index, option in enumerate(raw_options):
            if isinstance(option, str):
                option_text = option.strip()
                if option_text:
                    normalized_options.append({"id": option_index + 1, "text": option_text})
            elif isinstance(option, dict):
                option_text = str(option.get("text", "")).strip()
                if option_text:
                    normalized_options.append(
                        {
                            "id": option.get("id", option_index + 1),
                            "text": option_text,
                        }
                    )

        correct_answer = str(item.get("correct_answer", "") or "").strip()
        correct_answer_index = item.get("correct_answer_index")
        if correct_answer_index in ("", None):
            correct_answer_index = -1
        try:
            correct_answer_index = int(correct_answer_index)
        except (TypeError, ValueError):
            correct_answer_index = -1
        if question_type == "multiple_choice" and len(normalized_options) < 2:
            return []
        if question_type == "multiple_choice" and correct_answer:
            letter_match = re.match(r"^([A-Ha-h])(?:[\)\.\-:]|$)", correct_answer)
            if letter_match:
                idx = ord(letter_match.group(1).upper()) - ord("A")
                if 0 <= idx < len(normalized_options):
                    correct_answer = str(normalized_options[idx].get("text", "")).strip()
                    correct_answer_index = idx
            else:
                idx = next(
                    (i for i, opt in enumerate(normalized_options) if opt.get("text", "").strip().lower() == correct_answer.strip().lower()),
                    -1,
                )
                if idx >= 0:
                    correct_answer = str(normalized_options[idx].get("text", "")).strip()
                    correct_answer_index = idx
        if question_type == "multiple_choice" and not correct_answer and 0 <= correct_answer_index < len(normalized_options):
            correct_answer = str(normalized_options[correct_answer_index].get("text", "")).strip()
        if question_type == "true_false":
            normalized_options = [{"id": 1, "text": "True"}, {"id": 2, "text": "False"}]
            normalized_tf = _normalize_true_false_answer(correct_answer)
            if not normalized_tf:
                return []
            correct_answer = "true" if normalized_tf == "TRUE" else "false"
        if question_type in {"essay", "coding", "file_upload"}:
            normalized_options = []
            correct_answer = correct_answer or ""
        if question_type == "matching":
            pair_map = _parse_matching_pairs(item.get("matching_pairs") or correct_answer)
            if not pair_map and len(normalized_options) < 2:
                return []
            if pair_map:
                correct_answer = ",".join(f"{k}:{v}" for k, v in pair_map.items())
        acceptable_answers = item.get("acceptable_answers")
        if not isinstance(acceptable_answers, list):
            acceptable_answers = []
        if question_type == "short_answer" and not acceptable_answers and correct_answer:
            acceptable_answers = [part.strip() for part in re.split(r"[|;\n]+", correct_answer) if part.strip()]
        if question_type == "identification" and not acceptable_answers and correct_answer:
            acceptable_answers = [part.strip() for part in re.split(r"[|,;\n]+", correct_answer) if part.strip()]
        enumeration_items = []
        enumeration_scoring_mode = "partial"
        enumeration_points_mode = "equal"
        expected_count = item.get("expected_count")
        if question_type == "enumeration":
            enumeration_items = _normalize_enumeration_items(item)
            if not expected_count:
                expected_count = len(enumeration_items)
            if not expected_count:
                match = re.search(r"\b(?:give|list|name)\s+(\d+)\b", question_text, flags=re.IGNORECASE)
                if match:
                    expected_count = int(match.group(1))
            enumeration_scoring_mode = str(item.get("enumeration_scoring_mode") or "partial").strip().lower()
            if enumeration_scoring_mode not in {"strict", "partial", "percentage"}:
                enumeration_scoring_mode = "partial"
            enumeration_points_mode = str(item.get("enumeration_points_mode") or "equal").strip().lower()
            if enumeration_points_mode not in {"equal", "custom"}:
                enumeration_points_mode = "equal"
            if enumeration_points_mode == "custom":
                custom_total = sum(max(float(entry.get("points", 0) or 0), 0.0) for entry in enumeration_items)
                if custom_total > 0:
                    points = custom_total
                else:
                    enumeration_points_mode = "equal"
            if enumeration_items and enumeration_points_mode == "equal":
                equal_share = points / len(enumeration_items)
                enumeration_items = [
                    {
                        **entry,
                        "points": round(equal_share, 4),
                    }
                    for entry in enumeration_items
                ]

        normalized.append(
            {
                "id": question_id,
                "question_text": question_text,
                "type": question_type,
                "options": normalized_options if question_type in {"multiple_choice", "true_false", "matching"} else [],
                "correct_answer": correct_answer,
                "correct_answer_index": correct_answer_index,
                "acceptable_answers": acceptable_answers,
                "expected_count": expected_count or 0,
                "points": points,
                "section_id": item.get("section_id"),
                "section_title": item.get("section_title"),
                "instructions": item.get("instructions", ""),
                "starter_code": item.get("starter_code", ""),
                "language": item.get("language", ""),
                "expected_output": item.get("expected_output", ""),
                "test_cases": item.get("test_cases", ""),
                "allowed_file_types": item.get("allowed_file_types", ""),
                "max_file_size": item.get("max_file_size", ""),
                "matching_pairs": item.get("matching_pairs", []),
                "enumeration_answers": [entry.get("answer", "") for entry in enumeration_items] if enumeration_items else item.get("enumeration_answers", []),
                "enumeration_items": enumeration_items,
                "enumeration_scoring_mode": enumeration_scoring_mode,
                "enumeration_points_mode": enumeration_points_mode,
                "accepted_answers": item.get("accepted_answers", []),
                "correct_formula": item.get("correct_formula", ""),
                "formula_input": item.get("formula_input", ""),
            }
        )

    return normalized


def _grade_quiz_questions(questions, submitted_answers, frozen_total_points=None):
    submitted_by_question = {
        str(item["question_id"]): {
            "answer": str(item["answer"]).strip(),
            "answer_items": item.get("answer_items") or [],
        }
        for item in submitted_answers
    }

    total_points = 0.0
    score = 0.0
    correct_answers = 0
    incorrect_answers = 0
    breakdown = []

    for question in questions:
        question_id = str(question.get("id"))
        question_text = question.get("question_text", "")
        max_points = float(question.get("points", 1) or 1)
        total_points += max_points

        if question_id not in submitted_by_question:
            raise ValueError(f"Missing answer for question {question_id}.")

        submitted_payload = submitted_by_question[question_id]
        submitted_answer = submitted_payload.get("answer", "")
        correct_answer = str(question.get("correct_answer", "") or "").strip()
        enumeration_grade = {}

        question_type = str(question.get("type") or "").lower()
        requires_manual_review = question_type in {"essay", "coding", "file_upload"}
        if requires_manual_review:
            is_correct = False
        elif question_type == "multiple_choice":
            option_texts = [str(option.get("text", "")).strip() for option in (question.get("options") or []) if str(option.get("text", "")).strip()]
            normalized_submitted = submitted_answer.strip().lower()
            normalized_correct = correct_answer.strip().lower()
            if re.match(r"^[A-Z]$", correct_answer.strip(), flags=re.IGNORECASE):
                idx = ord(correct_answer.strip().upper()) - ord("A")
                if 0 <= idx < len(option_texts):
                    normalized_correct = option_texts[idx].strip().lower()
            if not normalized_correct and option_texts:
                raw_index = question.get("correct_answer_index")
                try:
                    idx = int(raw_index)
                except (TypeError, ValueError):
                    idx = -1
                if 0 <= idx < len(option_texts):
                    normalized_correct = option_texts[idx].strip().lower()
            is_correct = bool(normalized_submitted and normalized_correct and normalized_submitted == normalized_correct)
        elif question_type == "true_false":
            normalized_submitted = _normalize_true_false_answer(submitted_answer)
            normalized_correct = _normalize_true_false_answer(correct_answer)
            is_correct = bool(normalized_submitted and normalized_correct and normalized_submitted == normalized_correct)
        elif question_type == "short_answer":
            acceptable = [str(value).strip().lower() for value in (question.get("acceptable_answers") or []) if str(value).strip()]
            if not acceptable and correct_answer:
                acceptable = [part.strip().lower() for part in re.split(r"[|;\n]+", correct_answer) if part.strip()]
            submitted_norm = submitted_answer.strip().lower()
            is_correct = bool(submitted_norm and acceptable and any(submitted_norm == option for option in acceptable))
        elif question_type == "identification":
            keywords = [str(value).strip().lower() for value in (question.get("acceptable_answers") or []) if str(value).strip()]
            if not keywords and correct_answer:
                keywords = [part.strip().lower() for part in re.split(r"[|,;\n]+", correct_answer) if part.strip()]
            submitted_norm = submitted_answer.strip().lower()
            if keywords:
                is_correct = any(keyword in submitted_norm for keyword in keywords)
            else:
                is_correct = bool(submitted_norm and correct_answer and submitted_norm == correct_answer.strip().lower())
        elif question_type == "enumeration":
            enumeration_grade = _grade_enumeration_question(
                question,
                submitted_answer,
                submitted_payload.get("answer_items") or [],
            )
            is_correct = bool(enumeration_grade.get("is_correct"))
        elif question_type == "matching":
            expected_pairs = _parse_matching_pairs(correct_answer)
            submitted_pairs = _parse_matching_pairs(submitted_answer)
            is_correct = bool(expected_pairs) and expected_pairs == submitted_pairs
        else:
            is_correct = submitted_answer.lower() == correct_answer.lower()
        points_earned = enumeration_grade.get("points_earned", 0.0) if question_type == "enumeration" else (max_points if is_correct else 0.0)
        if requires_manual_review:
            incorrect_answers += 1
        else:
            score += points_earned
            if is_correct:
                correct_answers += 1
            else:
                incorrect_answers += 1

        breakdown.append(
            {
                "question_id": question.get("id"),
                "question_text": question_text,
                "question_type": question_type,
                "submitted_answer": submitted_answer,
                "submitted_answer_items": enumeration_grade.get("submitted_items", []) if question_type == "enumeration" else [],
                "is_correct": is_correct,
                "points_earned": points_earned,
                "max_points": max_points,
                "auto_score": None if requires_manual_review else points_earned,
                "manual_score": None,
                "override_score": None,
                "feedback": "",
                "answer_feedback": enumeration_grade.get("answer_feedback", []) if question_type == "enumeration" else [],
                "matched_count": enumeration_grade.get("matched_count", 0) if question_type == "enumeration" else None,
                "expected_count": enumeration_grade.get("expected_count", 0) if question_type == "enumeration" else None,
                "correct_answer_items": question.get("enumeration_answers", []) if question_type == "enumeration" else [],
                "status": QuizAttempt.STATUS_PENDING_REVIEW if requires_manual_review else QuizAttempt.STATUS_GRADED,
            }
        )

    effective_total_points = round(float(frozen_total_points), 2) if frozen_total_points not in (None, "") else round(total_points, 2)
    has_pending_review = any(item.get("status") == QuizAttempt.STATUS_PENDING_REVIEW for item in breakdown)

    return {
        "score": round(score, 2),
        "total_points": effective_total_points,
        "correct_answers": correct_answers,
        "incorrect_answers": incorrect_answers,
        "status": QuizAttempt.STATUS_PENDING_REVIEW if has_pending_review else QuizAttempt.STATUS_GRADED,
        "breakdown": breakdown,
    }


def _force_submit_locked_attempt(attempt, activity, reason="security_threshold"):
    if attempt.submitted_at and attempt.is_locked:
        return

    source_questions = (
        attempt.question_snapshot
        if isinstance(attempt.question_snapshot, list) and attempt.question_snapshot
        else _load_quiz_questions_for_runtime(activity)
    )
    questions = _validate_and_normalize_quiz_questions(source_questions)
    now = timezone.now()
    elapsed_seconds = int(max((now - attempt.started_at).total_seconds(), 0)) if attempt.started_at else 0
    frozen_total = float(attempt.total_points or 0)
    if frozen_total <= 0:
        frozen_total = round(sum(float(question.get("points", 1) or 1) for question in questions), 2)

    breakdown = []
    for question in questions:
        max_points = float(question.get("points", 1) or 1)
        breakdown.append(
            {
                "question_id": question.get("id"),
                "question_text": question.get("question_text", ""),
                "submitted_answer": "",
                "is_correct": False,
                "points_earned": 0.0,
                "max_points": max_points,
            }
        )

    attempt.answers = attempt.answers if isinstance(attempt.answers, list) else []
    attempt.score = 0.0
    attempt.total_points = round(frozen_total, 2)
    attempt.result_breakdown = breakdown
    attempt.correct_answers = 0
    attempt.incorrect_answers = len(questions)
    attempt.time_spent = elapsed_seconds
    attempt.last_activity_at = now
    attempt.submitted_at = now
    attempt.status = QuizAttempt.STATUS_GRADED
    attempt.graded_at = now
    attempt.is_locked = True
    attempt.lock_reason = reason
    attempt.force_submitted_at = now
    if attempt.is_locked and attempt.submitted_at is None:
        attempt.submitted_at = now
    attempt.save(
        update_fields=[
            "answers",
            "score",
            "total_points",
            "result_breakdown",
            "correct_answers",
            "incorrect_answers",
            "time_spent",
            "last_activity_at",
            "submitted_at",
            "status",
            "graded_at",
            "is_locked",
            "lock_reason",
            "force_submitted_at",
        ]
    )
    logger.warning(
        "Attempt force-submitted and locked.",
        extra={
            "attempt_id": attempt.id,
            "quiz_id": activity.id,
            "user_id": attempt.student_id,
            "reason": reason,
            "violation_count": int(attempt.suspicious_events or 0),
        },
    )

    ActivitySubmission.objects.update_or_create(
        activity=activity,
        student=attempt.student,
        defaults={
            "text_answer": json.dumps({"answers": attempt.answers, "forced": True}),
            "status": "graded",
            "grade": 0,
            "feedback": "Attempt force-submitted due to security violations.",
        },
    )


def _public_quiz_questions(questions):
    if not isinstance(questions, list):
        return []
    output = []
    for question in questions:
        if not isinstance(question, dict):
            continue
        q_type = question.get("type", "multiple_choice")
        options = [
            (
                {"id": option.get("id"), "text": option.get("text", "")}
                if isinstance(option, dict)
                else {"id": option_index + 1, "text": str(option).strip()}
            )
            for option_index, option in enumerate(question.get("options") or [])
            if (isinstance(option, dict) and str(option.get("text", "")).strip())
            or (isinstance(option, str) and option.strip())
        ]
        if q_type == "true_false" and not options:
            options = [{"id": 1, "text": "True"}, {"id": 2, "text": "False"}]
        output.append(
            {
                "id": question.get("id"),
                "question_text": question.get("question_text", ""),
                "type": q_type,
                "options": options,
                "points": question.get("points", 1),
                "section_id": question.get("section_id"),
                "section_title": question.get("section_title"),
                "instructions": question.get("instructions", ""),
                "starter_code": question.get("starter_code", ""),
                "language": question.get("language", ""),
                "expected_output": question.get("expected_output", ""),
                "test_cases": question.get("test_cases", ""),
                "allowed_file_types": question.get("allowed_file_types", ""),
                "max_file_size": question.get("max_file_size", ""),
                "matching_pairs": question.get("matching_pairs", []),
                "enumeration_answers": question.get("enumeration_answers", []),
                "accepted_answers": question.get("accepted_answers", []),
                "formula_input": question.get("formula_input", ""),
                "correct_formula": question.get("correct_formula", ""),
            }
        )
    return output


def _sanitize_attempt_for_student_visibility(attempt_data, activity):
    cleaned = dict(attempt_data or {})
    visibility = _resolve_attempt_visibility(activity, attempt_data=attempt_data)
    if cleaned.get("status") == QuizAttempt.STATUS_PENDING_REVIEW:
        visibility["show_score_immediately"] = False
        visibility["allow_answer_review"] = False
    cleaned["show_score_immediately"] = bool(visibility["show_score_immediately"])
    cleaned["allow_answer_review"] = bool(visibility["allow_answer_review"])
    if not visibility["show_score_immediately"]:
        cleaned["score"] = None
        cleaned["total_points"] = None
        cleaned["correct_answers"] = None
        cleaned["incorrect_answers"] = None
        cleaned["override_score"] = None
    if not visibility["allow_answer_review"]:
        cleaned["answers"] = []
        cleaned["result_breakdown"] = []
    elif cleaned.get("id"):
        cleaned["review_url"] = f"/api/courses/{activity.course_id}/activities/{activity.id}/quiz/review/?attempt_id={cleaned['id']}"
    return cleaned


def _resolve_attempt_visibility(activity, attempt=None, attempt_data=None):
    snapshot = {}
    if attempt is not None and isinstance(getattr(attempt, "visibility_snapshot", None), dict):
        snapshot = attempt.visibility_snapshot or {}
    elif isinstance(attempt_data, dict) and isinstance(attempt_data.get("visibility_snapshot"), dict):
        snapshot = attempt_data.get("visibility_snapshot") or {}
    return {
        "show_score_immediately": bool(snapshot.get("show_score_immediately", getattr(activity, "show_score_immediately", False))),
        "allow_answer_review": bool(snapshot.get("allow_answer_review", getattr(activity, "allow_answer_review", False))),
    }


def _compute_attempt_grade_totals(answer_rows):
    total_score = 0.0
    total_points = 0.0
    correct_answers = 0
    incorrect_answers = 0
    pending_manual_review = False
    breakdown = []

    for row in answer_rows:
        max_points = round(float(row.get("max_points", 0) or 0), 2)
        total_points += max_points
        auto_score = row.get("auto_score")
        manual_score = row.get("manual_score")
        override_score = row.get("override_score")
        status_value = str(row.get("status") or QuizAttempt.STATUS_GRADED).lower()
        final_score = override_score if override_score is not None else manual_score if manual_score is not None else auto_score
        final_score = round(float(final_score or 0), 2)
        is_pending = status_value == QuizAttempt.STATUS_PENDING_REVIEW
        if is_pending:
            pending_manual_review = True
        total_score += final_score
        if not is_pending:
            if max_points > 0 and final_score >= max_points:
                correct_answers += 1
            else:
                incorrect_answers += 1

        breakdown.append(
            {
                "question_id": row.get("question_id"),
                "question_text": row.get("question_text", ""),
                "question_type": row.get("question_type", ""),
                "submitted_answer": row.get("student_answer", ""),
                "auto_score": auto_score,
                "manual_score": manual_score,
                "override_score": override_score,
                "feedback": row.get("feedback", ""),
                "points_earned": final_score,
                "max_points": max_points,
                "is_correct": (not is_pending) and max_points > 0 and final_score >= max_points,
                "status": QuizAttempt.STATUS_PENDING_REVIEW if is_pending else QuizAttempt.STATUS_GRADED,
            }
        )

    return {
        "score": round(total_score, 2),
        "total_points": round(total_points, 2),
        "correct_answers": correct_answers,
        "incorrect_answers": incorrect_answers,
        "status": QuizAttempt.STATUS_PENDING_REVIEW if pending_manual_review else QuizAttempt.STATUS_GRADED,
        "breakdown": breakdown,
    }


def _sync_attempt_answer_records(attempt, questions, submitted_answers, grading):
    submitted_map = {str(item.get("question_id")): str(item.get("answer", "")).strip() for item in (submitted_answers or []) if isinstance(item, dict)}
    grading_map = {str(item.get("question_id")): item for item in (grading.get("breakdown") or []) if isinstance(item, dict)}
    existing_records = {str(item.question_id): item for item in attempt.answer_records.all()}
    seen_ids = set()

    for question in questions:
        question_id = str(question.get("id"))
        grading_row = grading_map.get(question_id, {})
        record = existing_records.get(question_id)
        payload = {
            "question_text": question.get("question_text", ""),
            "question_type": str(question.get("type") or "").lower(),
            "student_answer": submitted_map.get(question_id, str(grading_row.get("submitted_answer", "") or "")),
            "max_points": float(grading_row.get("max_points", question.get("points", 0)) or 0),
            "auto_score": grading_row.get("auto_score"),
            "manual_score": grading_row.get("manual_score"),
            "override_score": grading_row.get("override_score"),
            "feedback": str(grading_row.get("feedback", "") or ""),
            "status": grading_row.get("status") or grading.get("status") or QuizAttempt.STATUS_GRADED,
        }
        if record:
            for key, value in payload.items():
                setattr(record, key, value)
            record.save()
        else:
            QuizAttemptAnswer.objects.create(attempt=attempt, question_id=question_id, **payload)
        seen_ids.add(question_id)

    for question_id, record in existing_records.items():
        if question_id not in seen_ids:
            record.delete()


def _ensure_attempt_answer_records(attempt):
    if attempt.answer_records.exists():
        return

    source_questions = (
        attempt.question_snapshot
        if isinstance(attempt.question_snapshot, list) and attempt.question_snapshot
        else _load_quiz_questions_for_runtime(attempt.quiz)
    )
    questions = _validate_and_normalize_quiz_questions(source_questions)
    if not questions:
        return

    submitted_map = {
        str(item.get("question_id")): str(item.get("answer", "")).strip()
        for item in (attempt.answers or [])
        if isinstance(item, dict)
    }
    breakdown_map = {
        str(item.get("question_id")): item
        for item in (attempt.result_breakdown or [])
        if isinstance(item, dict) and item.get("question_id") is not None
    }
    fallback_grading = {"status": attempt.status, "breakdown": []}
    for question in questions:
        question_id = str(question.get("id"))
        breakdown_row = breakdown_map.get(question_id, {})
        fallback_grading["breakdown"].append(
            {
                "question_id": question.get("id"),
                "question_text": question.get("question_text", ""),
                "question_type": question.get("type", ""),
                "submitted_answer": submitted_map.get(question_id, str(breakdown_row.get("submitted_answer", "") or "")),
                "max_points": breakdown_row.get("max_points", question.get("points", 0)),
                "auto_score": breakdown_row.get("auto_score", breakdown_row.get("points_earned", 0)),
                "manual_score": breakdown_row.get("manual_score"),
                "override_score": breakdown_row.get("override_score"),
                "feedback": breakdown_row.get("feedback", ""),
                "status": breakdown_row.get("status") or attempt.status,
            }
        )
    _sync_attempt_answer_records(attempt, questions, attempt.answers or [], fallback_grading)


def _update_attempt_submission_record(attempt, feedback_text=""):
    ActivitySubmission.objects.update_or_create(
        activity=attempt.quiz,
        student=attempt.student,
        defaults={
            "text_answer": json.dumps({"answers": attempt.answers}),
            "status": "graded" if attempt.status == QuizAttempt.STATUS_GRADED else "submitted",
            "grade": attempt.override_score if attempt.override_score is not None else attempt.score,
            "feedback": feedback_text,
        },
    )


def _serialize_attempt_review_payload(attempt):
    _ensure_attempt_answer_records(attempt)
    return {
        "attempt_id": attempt.id,
        "student_id": attempt.student_id,
        "student_name": getattr(attempt.student, "username", "Student"),
        "status": attempt.status,
        "score": float(attempt.score or 0),
        "override_score": float(attempt.override_score) if attempt.override_score is not None else None,
        "display_score": float(attempt.override_score if attempt.override_score is not None else attempt.score or 0),
        "total_points": float(attempt.total_points or 0),
        "submitted_at": attempt.submitted_at,
        "graded_at": attempt.graded_at,
        "graded_by": attempt.graded_by_id,
        "graded_by_name": getattr(attempt.graded_by, "username", None),
        "is_overridden": bool(attempt.is_overridden),
        "answers": [
            {
                "id": row.id,
                "question_id": row.question_id,
                "question_text": row.question_text,
                "question_type": row.question_type,
                "student_answer": row.student_answer,
                "max_points": float(row.max_points or 0),
                "auto_score": float(row.auto_score) if row.auto_score is not None else None,
                "manual_score": float(row.manual_score) if row.manual_score is not None else None,
                "override_score": float(row.override_score) if row.override_score is not None else None,
                "final_score": float(row.override_score if row.override_score is not None else row.manual_score if row.manual_score is not None else row.auto_score or 0),
                "feedback": row.feedback,
                "status": row.status,
            }
            for row in attempt.answer_records.all().order_by("id")
        ],
        "audit_log": [
            {
                "id": audit.id,
                "question_id": audit.question_id,
                "previous_score": float(audit.previous_score) if audit.previous_score is not None else None,
                "new_score": float(audit.new_score) if audit.new_score is not None else None,
                "note": audit.note,
                "actor_id": audit.actor_id,
                "actor_name": getattr(audit.actor, "username", None),
                "created_at": audit.created_at,
            }
            for audit in attempt.score_audits.all()[:30]
        ],
    }


def _recompute_attempt_from_answer_records(attempt, actor=None, mark_override=False, audit_note=""):
    answer_rows = [
        {
            "question_id": row.question_id,
            "question_text": row.question_text,
            "question_type": row.question_type,
            "student_answer": row.student_answer,
            "max_points": row.max_points,
            "auto_score": row.auto_score,
            "manual_score": row.manual_score,
            "override_score": row.override_score,
            "feedback": row.feedback,
            "status": row.status,
        }
        for row in attempt.answer_records.all().order_by("id")
    ]
    computed = _compute_attempt_grade_totals(answer_rows)
    previous_display_score = attempt.override_score if attempt.override_score is not None else attempt.score
    attempt.score = computed["score"]
    attempt.total_points = computed["total_points"]
    attempt.correct_answers = computed["correct_answers"]
    attempt.incorrect_answers = computed["incorrect_answers"]
    attempt.result_breakdown = computed["breakdown"]
    attempt.status = computed["status"]
    attempt.graded_at = timezone.now()
    attempt.graded_by = actor
    if attempt.status == QuizAttempt.STATUS_PENDING_REVIEW:
        attempt.override_score = None
        attempt.is_overridden = False
    elif mark_override:
        attempt.is_overridden = True
    attempt.save(
        update_fields=[
            "score",
            "total_points",
            "correct_answers",
            "incorrect_answers",
            "result_breakdown",
            "status",
            "graded_at",
            "graded_by",
            "override_score",
            "is_overridden",
        ]
    )
    current_display_score = attempt.override_score if attempt.override_score is not None else attempt.score
    if previous_display_score != current_display_score or audit_note:
        QuizAttemptScoreAudit.objects.create(
            attempt=attempt,
            actor=actor,
            previous_score=previous_display_score,
            new_score=current_display_score,
            note=audit_note,
        )

    feedback_text = f"Correct answers: {attempt.correct_answers} | Incorrect answers: {attempt.incorrect_answers}"
    if attempt.status == QuizAttempt.STATUS_PENDING_REVIEW:
        feedback_text = "Submission received and pending manual review."
    _update_attempt_submission_record(attempt, feedback_text=feedback_text)
    return attempt


def _randomize_questions_for_attempt(
    questions,
    student_id,
    quiz_id,
    attempt_id=None,
    shuffle_questions=False,
    shuffle_choices=False,
    subset_size=0,
):
    cloned = [dict(question) for question in (questions or []) if isinstance(question, dict)]
    seed = f"{student_id}-{quiz_id}-{attempt_id or 'preview'}"
    rng = random.Random(seed)

    if shuffle_questions:
        rng.shuffle(cloned)

    if subset_size and subset_size > 0:
        try:
            subset_size = int(subset_size)
        except (TypeError, ValueError):
            subset_size = 0
        if subset_size > 0 and subset_size < len(cloned):
            cloned = cloned[:subset_size]

    if shuffle_choices:
        for question in cloned:
            options = question.get("options")
            if isinstance(options, list) and len(options) > 1:
                rng.shuffle(options)

    return cloned


def _get_course_with_access(course_id, user):
    course = get_object_or_404(Course, id=course_id)
    if user.role == "instructor" and course.instructor_id == user.id:
        return course
    if user.role == "student" and course.students.filter(id=user.id).exists():
        return course
    return None


@api_view(["GET", "PUT"])
@permission_classes([IsAuthenticated])
def grading_scheme_detail(request, course_id):
    course = _get_course_with_access(course_id, request.user)
    if not course:
        return Response({"error": "Course not found or access denied"}, status=404)

    scheme = GradingScheme.objects.filter(course=course).prefetch_related("components").first()

    if request.method == "GET":
        if not scheme:
            data = {
                "id": None,
                "course": course.id,
                "grading_type": GradingScheme.TYPE_ZERO_BASED,
                "passing_grade": 75,
                "custom_config": {},
                "components": [],
            }
        else:
            serializer = GradingSchemeSerializer(scheme, context={"request": request, "course": course})
            data = serializer.data
        activities = list(CourseActivity.objects.filter(course=course).select_related("activity_type").order_by("created_at", "id"))
        detected_activities = []
        category_counts = OrderedDict((key, 0) for key in ACTIVITY_CATEGORY_LABELS.keys())
        for activity in activities:
            activity_type_name = str(getattr(activity.activity_type, "name", "") or "").strip().lower()
            if activity_type_name == "quiz" and str(getattr(activity, "assessment_type", "") or "").strip().lower() == "exam":
                category_key = "exam"
            elif activity_type_name in {"assignment", "task", "homework"}:
                category_key = "assignment"
            elif activity_type_name in {"project", "material", "attendance"}:
                category_key = activity_type_name
            elif activity_type_name == "quiz":
                category_key = "quiz"
            else:
                category_key = "other"
            category_counts[category_key] += 1
            detected_activities.append(
                {
                    "id": int(activity.id),
                    "title": activity.title,
                    "category_key": category_key,
                    "category_label": ACTIVITY_CATEGORY_LABELS.get(category_key, ACTIVITY_CATEGORY_LABELS["other"]),
                    "activity_type": str(getattr(activity.activity_type, "name", "") or ""),
                    "assessment_type": str(getattr(activity, "assessment_type", "") or ""),
                    "grading_type": str(getattr(activity, "grading_type", "") or ""),
                    "points": float(getattr(activity, "points", 0) or 0),
                }
            )

        detected_categories = [
            {"key": key, "label": label, "count": int(category_counts.get(key, 0))}
            for key, label in ACTIVITY_CATEGORY_LABELS.items()
            if category_counts.get(key, 0) > 0
        ]
        default_categories = detected_categories or [{"key": "assignment", "label": "Assignments", "count": 0}]
        suggested_weight = round(100.0 / len(default_categories), 2) if default_categories else 100.0
        suggested_components = []
        running_total = 0.0
        for index, category in enumerate(default_categories):
            weight = suggested_weight
            if index == len(default_categories) - 1:
                weight = round(100.0 - running_total, 2)
            running_total += weight
            suggested_components.append(
                {
                    "name": category["label"],
                    "weight": weight,
                    "activity_ids": [],
                    "category_key": category["key"],
                    "drop_lowest_count": 0,
                }
            )

        data["detected_activities"] = detected_activities
        data["available_categories"] = detected_categories
        data["suggested_components"] = suggested_components
        return Response(data)

    if course.instructor_id != request.user.id:
        return Response({"error": "Only instructor can update grading scheme."}, status=403)

    payload = request.data.copy()
    with transaction.atomic():
        if scheme:
            serializer = GradingSchemeSerializer(scheme, data=payload, context={"request": request, "course": course})
        else:
            serializer = GradingSchemeSerializer(data=payload, context={"request": request, "course": course})
        serializer.is_valid(raise_exception=True)
        if scheme:
            serializer.save()
        else:
            serializer.save(course=course)
    return Response(serializer.data)


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def grade_sheet(request, course_id):
    course = _get_course_with_access(course_id, request.user)
    if not course:
        return Response({"error": "Course not found or access denied"}, status=404)

    scheme = GradingScheme.objects.filter(course=course).prefetch_related("components").first()
    if not scheme:
        return Response({"error": "Grading scheme is not configured for this course."}, status=400)
    try:
        scheme.validate_component_weights()
    except Exception as exc:
        return Response({"error": str(exc)}, status=400)

    if request.user.role == "student":
        students = course.students.filter(id=request.user.id, role="student")
    else:
        students = course.students.filter(role="student").order_by("last_name", "first_name", "username")

    rows = _build_grade_sheet_rows(course, scheme, students)
    serializer = GradeSheetSerializer(rows, many=True)
    return Response(serializer.data)


def _build_grade_sheet_rows(course, scheme, students):
    rows = []
    grade_details_map = compute_grade_details_for_students(course, list(students))
    for student in students:
        grade_details = grade_details_map.get(student.id, {})
        student_name = (f"{student.first_name} {student.last_name}").strip() or student.username

        if grade_details.get("error"):
            rows.append(
                {
                    "student_id": student.id,
                    "student_name": student_name,
                    "components": {},
                    "activities": [],
                    "uncovered_activities": [],
                    "weighted_total": 0.0,
                    "final_grade": 0.0,
                    "status": "Failed",
                    "remarks": f"Error: {grade_details['error']}",
                    "formula": "Unavailable due to grading configuration error.",
                    "formula_text": "Unavailable due to grading configuration error.",
                }
            )
            continue

        final_grade = float(grade_details["final_grade"])
        status_text = "Passed" if final_grade >= float(scheme.passing_grade) else "Failed"
        remarks = "Good Standing" if status_text == "Passed" else "At Risk"
        rows.append(
            {
                "student_id": student.id,
                "student_name": student_name,
                "components": grade_details["components"],
                "activities": grade_details.get("activities", []),
                "uncovered_activities": grade_details.get("uncovered_activities", []),
                "weighted_total": float(grade_details.get("weighted_total", 0.0)),
                "final_grade": final_grade,
                "status": status_text,
                "remarks": remarks,
                "formula": grade_details["formula"],
                "formula_text": grade_details.get("formula_text", grade_details["formula"]),
            }
        )
    return rows


def _excel_escape_sheet_formula(value):
    return str(value or "").replace('"', '""')


def _excel_activity_percent_formula(score_cell, max_cell, grading_type, passfail_threshold, treat_missing_as_zero):
    if grading_type == "percent":
        if treat_missing_as_zero:
            return f'=MIN(100,MAX(0,{score_cell}))'
        return f'=IF({score_cell}="", "", MIN(100,MAX(0,{score_cell})))'
    if grading_type == "passfail":
        if treat_missing_as_zero:
            return f'=IF({max_cell}<=0,0,IF(({score_cell}/{max_cell})*100>={passfail_threshold},100,0))'
        return f'=IF(OR({score_cell}="",{max_cell}=""),"",IF({max_cell}<=0,0,IF(({score_cell}/{max_cell})*100>={passfail_threshold},100,0)))'
    if grading_type == "none":
        return '=""'
    if treat_missing_as_zero:
        return f'=IF({max_cell}<=0,0,MIN(100,MAX(0,({score_cell}/{max_cell})*100)))'
    return f'=IF(OR({score_cell}="",{max_cell}=""),"",IF({max_cell}<=0,0,MIN(100,MAX(0,({score_cell}/{max_cell})*100))))'


def _excel_component_average_formula(refs, drop_lowest_count):
    if not refs:
        return "=0"
    joined = ",".join(refs)
    if drop_lowest_count <= 0:
        return f"=IF(COUNT({joined})=0,0,AVERAGE({joined}))"
    if drop_lowest_count == 1:
        return f"=IF(COUNT({joined})<=1,0,(SUM({joined})-MIN({joined}))/(COUNT({joined})-1))"
    return None


def _excel_transmutation_formula(weighted_total_cell, table):
    segments = []
    for row in reversed(table):
        minimum = float(row.get("min", 0) or 0)
        value = float(row.get("value", 0) or 0)
        segments.append(f'IF({weighted_total_cell}>={minimum},{value}')
    closing = ")" * len(segments)
    return f"={' '.join(segments)},0{closing}".replace(" ", "")


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def grade_sheet_export(request, course_id):
    course = _get_course_with_access(course_id, request.user)
    if not course:
        return Response({"error": "Course not found or access denied"}, status=404)
    if request.user.role != "instructor" or course.instructor_id != request.user.id:
        return Response({"error": "Only instructor can export Excel grade sheets."}, status=403)
    if xlsxwriter is None:
        return Response({"error": "Excel export is unavailable because xlsxwriter is not installed."}, status=500)

    scheme = GradingScheme.objects.filter(course=course).prefetch_related("components").first()
    if not scheme:
        return Response({"error": "Grading scheme is not configured for this course."}, status=400)

    students = list(course.students.filter(role="student").order_by("last_name", "first_name", "username"))
    try:
        scheme.validate_component_weights()
    except Exception as exc:
        return Response({"error": str(exc)}, status=400)
    grade_rows = _build_grade_sheet_rows(course, scheme, students)

    activity_map = OrderedDict()
    for row in grade_rows:
        for activity in row.get("activities") or []:
            activity_map.setdefault(int(activity["activity_id"]), activity)
    component_names = [component.name for component in scheme.components.all()]
    treat_missing_as_zero = bool((scheme.custom_config or {}).get("treat_missing_as_zero", True))
    passfail_threshold = float((scheme.custom_config or {}).get("passfail_threshold", 60) or 60)
    formula_expression = str((scheme.custom_config or {}).get("formula_expression") or "").strip()
    transmutation_table = (scheme.custom_config or {}).get("transmutation_table") or []

    output = io.BytesIO()
    workbook = xlsxwriter.Workbook(output, {"in_memory": True})
    worksheet = workbook.add_worksheet("Grades")
    header_format = workbook.add_format({"bold": True, "bg_color": "#D9F2E6", "border": 1})
    text_format = workbook.add_format({"border": 1})
    number_format = workbook.add_format({"border": 1, "num_format": "0.00"})

    headers = ["Student Name"]
    for activity in activity_map.values():
        headers.extend(
            [
                f'{activity["title"]} Score',
                f'{activity["title"]} Max',
                f'{activity["title"]} %',
            ]
        )
    for component_name in component_names:
        headers.extend([f"{component_name} %", f"{component_name} Weight", f"{component_name} Weighted"])
    headers.extend(["Weighted Total", "Final Grade", "Status", "Remarks", "Formula"])

    for col_index, header in enumerate(headers):
        worksheet.write(0, col_index, header, header_format)

    activity_percent_cols = {}
    col_index = 1
    for activity_id, activity in activity_map.items():
        activity_percent_cols[activity_id] = col_index + 2
        col_index += 3

    component_raw_cols = {}
    component_weight_cols = {}
    component_weighted_cols = {}
    for component_name in component_names:
        component_raw_cols[component_name] = col_index
        component_weight_cols[component_name] = col_index + 1
        component_weighted_cols[component_name] = col_index + 2
        col_index += 3
    weighted_total_col = col_index
    final_grade_col = col_index + 1
    status_col = col_index + 2
    remarks_col = col_index + 3
    formula_col = col_index + 4

    for row_index, row in enumerate(grade_rows, start=1):
        worksheet.write(row_index, 0, row.get("student_name"), text_format)
        activity_lookup = {int(item["activity_id"]): item for item in (row.get("activities") or [])}
        col = 1
        for activity_id, activity in activity_map.items():
            row_activity = activity_lookup.get(activity_id, {})
            score_value = row_activity.get("score")
            max_value = row_activity.get("max_score", activity.get("max_score", 0))
            score_cell = xlsxwriter.utility.xl_rowcol_to_cell(row_index, col)
            max_cell = xlsxwriter.utility.xl_rowcol_to_cell(row_index, col + 1)

            if score_value is None and not treat_missing_as_zero:
                worksheet.write_blank(row_index, col, None, text_format)
            else:
                worksheet.write_number(row_index, col, float(score_value or 0), number_format)
            worksheet.write_number(row_index, col + 1, float(max_value or 0), number_format)
            percent_formula = _excel_activity_percent_formula(
                score_cell,
                max_cell,
                str(activity.get("grading_type") or "points").strip().lower(),
                passfail_threshold,
                treat_missing_as_zero,
            )
            worksheet.write_formula(row_index, col + 2, percent_formula, number_format)
            col += 3

        for component_name in component_names:
            component = (row.get("components") or {}).get(component_name, {})
            refs = []
            for activity in component.get("activities") or []:
                if activity.get("dropped") or activity.get("excluded"):
                    continue
                percent_col = activity_percent_cols.get(int(activity["activity_id"]))
                if percent_col is not None:
                    refs.append(xlsxwriter.utility.xl_rowcol_to_cell(row_index, percent_col))
            component_formula = _excel_component_average_formula(refs, int(component.get("drop_lowest_count") or 0))
            if component_formula:
                worksheet.write_formula(row_index, component_raw_cols[component_name], component_formula, number_format)
            else:
                worksheet.write_number(row_index, component_raw_cols[component_name], float(component.get("raw", 0) or 0), number_format)
            worksheet.write_number(row_index, component_weight_cols[component_name], float(component.get("weight", 0) or 0), number_format)
            raw_cell = xlsxwriter.utility.xl_rowcol_to_cell(row_index, component_raw_cols[component_name])
            weight_cell = xlsxwriter.utility.xl_rowcol_to_cell(row_index, component_weight_cols[component_name])
            worksheet.write_formula(row_index, component_weighted_cols[component_name], f"={raw_cell}*({weight_cell}/100)", number_format)

        weighted_refs = [
            xlsxwriter.utility.xl_rowcol_to_cell(row_index, component_weighted_cols[component_name])
            for component_name in component_names
        ]
        worksheet.write_formula(row_index, weighted_total_col, f"=SUM({','.join(weighted_refs)})", number_format)
        weighted_total_cell = xlsxwriter.utility.xl_rowcol_to_cell(row_index, weighted_total_col)

        if scheme.grading_type == GradingScheme.TYPE_ZERO_BASED:
            final_formula = f"={weighted_total_cell}"
        elif scheme.grading_type == GradingScheme.TYPE_TRANSMUTED:
            final_formula = f"=50+({weighted_total_cell}*0.5)"
        elif formula_expression:
            final_formula = formula_expression
            replacement_pairs = {"weighted_total": weighted_total_cell}
            for component_name in component_names:
                replacement_pairs[_slugify_identifier(component_name, "component")] = xlsxwriter.utility.xl_rowcol_to_cell(
                    row_index, component_raw_cols[component_name]
                )
                replacement_pairs[f"{_slugify_identifier(component_name, 'component')}_weighted"] = xlsxwriter.utility.xl_rowcol_to_cell(
                    row_index, component_weighted_cols[component_name]
                )
            for token, cell_ref in sorted(replacement_pairs.items(), key=lambda item: len(item[0]), reverse=True):
                final_formula = re.sub(rf"\b{re.escape(token)}\b", cell_ref, final_formula)
            final_formula = f"={final_formula}"
        elif transmutation_table:
            final_formula = _excel_transmutation_formula(weighted_total_cell, transmutation_table)
        else:
            final_formula = f"={weighted_total_cell}"

        worksheet.write_formula(row_index, final_grade_col, final_formula, number_format)
        final_grade_cell = xlsxwriter.utility.xl_rowcol_to_cell(row_index, final_grade_col)
        worksheet.write_formula(row_index, status_col, f'=IF({final_grade_cell}>={float(scheme.passing_grade)},"Passed","Failed")', text_format)
        worksheet.write_formula(row_index, remarks_col, f'=IF({final_grade_cell}>={float(scheme.passing_grade)},"Good Standing","At Risk")', text_format)
        worksheet.write(row_index, formula_col, row.get("formula"), text_format)

    worksheet.freeze_panes(1, 1)
    worksheet.autofilter(0, 0, max(len(grade_rows), 1), len(headers) - 1)
    worksheet.set_column(0, 0, 28)
    worksheet.set_column(1, len(headers) - 1, 14)
    workbook.close()
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="course-{course_id}-gradesheet.xlsx"'
    return response


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def upload_grades(request, course_id):
    course = _get_course_with_access(course_id, request.user)
    if not course:
        return Response({"error": "Course not found or access denied"}, status=404)
    if course.instructor_id != request.user.id:
        return Response({"error": "Only instructor can upload grades."}, status=403)

    scheme = GradingScheme.objects.filter(course=course).prefetch_related("components").first()
    if not scheme:
        return Response({"error": "Grading scheme is not configured for this course."}, status=400)

    uploaded_file = request.FILES.get("file")
    if not uploaded_file:
        return Response({"error": "CSV file is required."}, status=400)

    try:
        decoded = uploaded_file.read().decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(decoded))
    except Exception:
        return Response({"error": "Invalid CSV format."}, status=400)

    if "student_id" not in (reader.fieldnames or []):
        return Response({"error": "CSV must contain student_id column."}, status=400)

    component_map = {component.name.strip().lower(): component for component in scheme.components.all()}
    updated_rows = 0
    validation_errors = []
    parsed_updates = []

    def _parse_score(row_number, component_name, raw_value):
        try:
            score_value = float(raw_value)
        except (TypeError, ValueError):
            validation_errors.append(
                {"row": row_number, "component": component_name, "error": "Score must be numeric."}
            )
            return None
        if score_value != score_value or score_value in (float("inf"), float("-inf")):
            validation_errors.append(
                {"row": row_number, "component": component_name, "error": "Score must be finite."}
            )
            return None
        if score_value < 0 or score_value > 100:
            validation_errors.append(
                {"row": row_number, "component": component_name, "error": "Score must be between 0 and 100."}
            )
            return None
        return score_value

    row_number = 1
    for row in reader:
        row_number += 1
        student_id = row.get("student_id")
        if not student_id:
            continue
        student = course.students.filter(id=student_id, role="student").first()
        if not student:
            continue

        for column_name, column_value in row.items():
            if column_name is None or column_name.lower() == "student_id":
                continue
            component = component_map.get(str(column_name).strip().lower())
            if not component:
                continue
            score_value = _parse_score(row_number, column_name, column_value)
            if score_value is None:
                continue
            parsed_updates.append((component, student, score_value))

    if validation_errors:
        return Response(
            {"error": "CSV contains invalid score values.", "details": validation_errors},
            status=400,
        )

    with transaction.atomic():
        for component, student, score_value in parsed_updates:
            GradingComponentScore.objects.update_or_create(
                component=component,
                student=student,
                defaults={"raw_score": score_value},
            )
            updated_rows += 1

    return Response({"message": "Grades uploaded successfully.", "updated_scores": updated_rows})


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def quiz_detail(request, course_id, activity_id):
    course = _get_course_with_access(course_id, request.user)
    if not course:
        return Response({"error": "Course not found or access denied"}, status=404)

    activity = get_object_or_404(CourseActivity, id=activity_id, course=course)
    if not _is_quiz_activity(activity):
        return Response({"error": "This activity is not a quiz"}, status=400)
    if str(activity.publish_state or "").lower() != CourseActivity.PUBLISH_STATE_PUBLISHED:
        return Response({"error": "This assessment is not currently available."}, status=400)

    attempts_qs = QuizAttempt.objects.filter(quiz=activity, student=request.user).order_by("-started_at")
    latest_attempt = attempts_qs.first()
    if latest_attempt and latest_attempt.is_locked:
        return Response(
            {
                "attempt_locked": True,
                "attempt_id": latest_attempt.id,
                "lock_reason": latest_attempt.lock_reason or "",
                "score": float(latest_attempt.score or 0),
                "total_points": float(latest_attempt.total_points or 0),
                "submitted_at": latest_attempt.submitted_at,
            },
            status=423,
        )
    attempts_data = QuizAttemptSerializer(attempts_qs, many=True).data
    if request.user.role == "student":
        attempts_data = [_sanitize_attempt_for_student_visibility(item, activity) for item in attempts_data]
    requested_attempt_id = request.query_params.get("attempt_id")
    active_attempt = None
    if requested_attempt_id:
        active_attempt = attempts_qs.filter(id=requested_attempt_id).first()
    if active_attempt is None:
        active_attempt = attempts_qs.filter(submitted_at__isnull=True).first()
    if active_attempt and active_attempt.is_locked:
        return Response(
            {
                "error": "This attempt is locked due to security violations.",
                "attempt_locked": True,
                "attempt_id": active_attempt.id,
                "score": float(active_attempt.score or 0),
                "total_points": float(active_attempt.total_points or 0),
                "submitted_at": active_attempt.submitted_at,
            },
            status=423,
        )

    source_questions = None
    if active_attempt and isinstance(active_attempt.question_snapshot, list) and active_attempt.question_snapshot:
        source_questions = active_attempt.question_snapshot
    else:
        source_questions = _load_quiz_questions_for_runtime(activity)

    questions = _validate_and_normalize_quiz_questions(source_questions)
    if not questions:
        return Response({"error": "Quiz has no questions configured"}, status=400)
    randomized_questions = _randomize_questions_for_attempt(
        questions,
        student_id=request.user.id,
        quiz_id=activity.id,
        attempt_id=active_attempt.id if active_attempt else None,
        shuffle_questions=bool(activity.randomize_questions),
        shuffle_choices=bool(getattr(activity, "randomize_choices", False)),
        subset_size=int(getattr(activity, "random_subset_size", 0) or 0),
    )
    total_points = (
        float(active_attempt.total_points)
        if active_attempt and float(active_attempt.total_points or 0) > 0
        else round(sum(float(question.get("points", 1) or 1) for question in questions), 2)
    )

    data = {
        "id": activity.id,
        "title": activity.title,
        "assessment_type": activity.assessment_type,
        "publish_state": activity.publish_state,
        "time_limit": int(activity.quiz_time_limit_seconds or 600),
        "total_points": total_points,
        "max_attempts": int(activity.max_attempts or 3),
        "randomize_questions": bool(activity.randomize_questions),
        "randomize_choices": bool(getattr(activity, "randomize_choices", False)),
        "random_subset_size": int(getattr(activity, "random_subset_size", 0) or 0),
        "require_answer_to_advance": bool(getattr(activity, "require_answer_to_advance", False)),
        "anti_cheat_enabled": bool(getattr(activity, "anti_cheat_enabled", False)),
        "anti_cheat_tab_switch": bool(getattr(activity, "anti_cheat_tab_switch", False)),
        "anti_cheat_multi_tab": bool(getattr(activity, "anti_cheat_multi_tab", False)),
        "anti_cheat_disable_copy_paste": bool(getattr(activity, "anti_cheat_disable_copy_paste", False)),
        "anti_cheat_fullscreen_required": bool(getattr(activity, "anti_cheat_fullscreen_required", False)),
        "availability_start": activity.availability_start,
        "availability_end": activity.availability_end,
        "questions": _public_quiz_questions(randomized_questions),
        "attempt_id": active_attempt.id if active_attempt else None,
        "current_attempt": active_attempt.id if active_attempt else None,
        "attempts": attempts_data,
        "pre_exam_message": str((activity.classwork_metadata or {}).get("pre_exam_message", "")).strip(),
        "requires_consent": True,
        "show_score_immediately": bool(getattr(activity, "show_score_immediately", False)),
        "allow_answer_review": bool(getattr(activity, "allow_answer_review", False)),
    }
    return Response(data)


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def quiz_review(request, course_id, activity_id):
    course = _get_course_with_access(course_id, request.user)
    if not course:
        return Response({"error": "Course not found or access denied"}, status=404)

    if request.user.role != "student":
        return Response({"error": "Only students can access exam review."}, status=403)

    activity = get_object_or_404(CourseActivity, id=activity_id, course=course)
    if not _is_quiz_activity(activity):
        return Response({"error": "This activity is not a quiz"}, status=400)

    submitted_attempts = QuizAttempt.objects.filter(
        quiz=activity,
        student=request.user,
        submitted_at__isnull=False,
    ).order_by("-submitted_at", "-started_at")

    requested_attempt_id = request.query_params.get("attempt_id")
    if requested_attempt_id:
        attempt = submitted_attempts.filter(id=requested_attempt_id).first()
    else:
        attempt = submitted_attempts.first()
    if not attempt:
        return Response({"error": "No submitted attempt found for this exam."}, status=404)
    attempt_visibility = _resolve_attempt_visibility(activity, attempt=attempt)
    if attempt.status == QuizAttempt.STATUS_PENDING_REVIEW:
        return Response({"error": "This submission is still pending instructor review."}, status=409)
    if not attempt_visibility["allow_answer_review"]:
        return Response({"error": "Detailed review is not available for this exam."}, status=403)

    source_questions = (
        attempt.question_snapshot
        if isinstance(attempt.question_snapshot, list) and attempt.question_snapshot
        else _load_quiz_questions_for_runtime(activity)
    )
    questions = _validate_and_normalize_quiz_questions(source_questions)
    answers_by_question = {
        str(item.get("question_id")): {
            "answer": str(item.get("answer", "")).strip(),
            "answer_items": item.get("answer_items") or [],
        }
        for item in (attempt.answers or [])
        if isinstance(item, dict)
    }
    breakdown_by_question = {
        str(item.get("question_id")): item
        for item in (attempt.result_breakdown or [])
        if isinstance(item, dict) and item.get("question_id") is not None
    }

    review_questions = []
    for question in questions:
        question_id = str(question.get("id"))
        breakdown_item = breakdown_by_question.get(question_id) or {}
        points_earned = float(breakdown_item.get("points_earned", 0) or 0)
        is_correct = bool(breakdown_item.get("is_correct", False))
        review_questions.append(
            {
                "question_id": question.get("id"),
                "question": question.get("question_text", ""),
                "question_type": question.get("type", ""),
                "student_answer": (answers_by_question.get(question_id) or {}).get("answer", str(breakdown_item.get("submitted_answer", "") or "")),
                "student_answer_items": (answers_by_question.get(question_id) or {}).get("answer_items", breakdown_item.get("submitted_answer_items", []) or []),
                "correct_answer": str(question.get("correct_answer", "") or ""),
                "correct_answer_items": breakdown_item.get("correct_answer_items", question.get("enumeration_answers", []) or []),
                "is_correct": is_correct,
                "points": points_earned,
                "max_points": float(question.get("points", 1) or 1),
                "answer_feedback": breakdown_item.get("answer_feedback", []) or [],
            }
        )

    score = float(attempt.score or 0)
    total_points = float(attempt.total_points or 0)
    percentage = round((score / total_points) * 100, 2) if total_points > 0 else 0.0
    return Response(
        {
            "attempt_id": attempt.id,
            "score": score,
            "total_points": total_points,
            "percentage": percentage,
            "questions": review_questions,
        }
    )


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def quiz_start(request, course_id, activity_id):
    logger.debug(
        "Quiz start requested.",
        extra={"course_id": course_id, "activity_id": activity_id, "user_id": getattr(request.user, "id", None)},
    )

    if request.user.role != "student":
        return Response({"error": "Only students can start quizzes."}, status=403)

    try:
        course = Course.objects.filter(id=course_id, students=request.user).first()
        if not course:
            return Response({"error": "Course not found or access denied"}, status=404)

        enrollment_exists = course.students.filter(id=request.user.id).exists()
        if not enrollment_exists:
            return Response({"error": "Course not found or access denied"}, status=404)

        activity = CourseActivity.objects.filter(id=activity_id, course=course).first()
        if not activity:
            return Response({"error": "Activity not found"}, status=404)

        if not _is_quiz_activity(activity):
            return Response({"error": "Activity is not a quiz"}, status=400)
        if str(activity.publish_state or "").lower() != CourseActivity.PUBLISH_STATE_PUBLISHED:
            return Response({"error": "This assessment is not currently available."}, status=400)

        now = timezone.now()
        if activity.availability_start and now < activity.availability_start:
            return Response({"error": "This assessment is not yet available."}, status=400)
        if activity.availability_end and now > activity.availability_end:
            return Response({"error": "This assessment has already closed."}, status=400)

        raw_questions = _load_quiz_questions_for_runtime(activity)
        questions = _validate_and_normalize_quiz_questions(raw_questions)
        if not questions:
            return Response({"error": "Quiz has no questions configured"}, status=400)
        frozen_total_points = round(sum(float(question.get("points", 1) or 1) for question in questions), 2)
        requires_consent = True
        consent_accepted = bool(request.data.get("acknowledged"))
        consent_message = str(request.data.get("ack_message") or (activity.classwork_metadata or {}).get("pre_exam_message", "")).strip()
        if requires_consent and not consent_accepted:
            return Response(
                {
                    "error": "Pre-exam acknowledgment is required before starting.",
                    "requires_consent": True,
                    "pre_exam_message": str((activity.classwork_metadata or {}).get("pre_exam_message", "")).strip(),
                },
                status=400,
            )

        max_attempts = int(activity.max_attempts or 3)
        if max_attempts < 1:
            max_attempts = 1

        existing_attempt = QuizAttempt.objects.filter(
            quiz=activity,
            student=request.user,
            submitted_at__isnull=True,
            is_locked=False,
        ).order_by("-started_at").first()
        logger.debug(
            "Quiz start active attempt lookup complete.",
            extra={
                "course_id": course_id,
                "activity_id": activity_id,
                "user_id": request.user.id,
                "existing_attempt_id": existing_attempt.id if existing_attempt else None,
            },
        )
        if existing_attempt:
            return Response(
                {
                    "attempt_id": existing_attempt.id,
                    "resume": True,
                    "started_at": existing_attempt.started_at,
                    "time_limit": int(activity.quiz_time_limit_seconds or 600),
                    "max_attempts": max_attempts,
                    "total_points": float(existing_attempt.total_points or frozen_total_points),
                    "anti_cheat": _anti_cheat_runtime_config(activity),
                    "requires_consent": True,
                },
                status=200,
            )

        submitted_attempts_count = QuizAttempt.objects.filter(
            quiz=activity,
            student=request.user,
            submitted_at__isnull=False,
        ).count()
        logger.debug(
            "Quiz start submitted attempt count computed.",
            extra={
                "course_id": course_id,
                "activity_id": activity_id,
                "user_id": request.user.id,
                "submitted_attempts_count": submitted_attempts_count,
            },
        )
        if submitted_attempts_count >= max_attempts:
            return Response(
                {"error": "Maximum quiz attempts reached"},
                status=400,
            )

        with transaction.atomic():
            now_ts = timezone.now()
            attempt = QuizAttempt.objects.create(
                student=request.user,
                quiz=activity,
                started_at=now_ts,
                last_activity_at=now_ts,
                total_points=frozen_total_points,
                question_snapshot=questions,
                visibility_snapshot={
                    "show_score_immediately": bool(activity.show_score_immediately),
                    "allow_answer_review": bool(activity.allow_answer_review),
                },
            )
            QuizAttemptAcknowledgement.objects.update_or_create(
                attempt=attempt,
                student=request.user,
                defaults={
                    "quiz": activity,
                    "ack_timestamp": now_ts,
                    "ack_message": consent_message,
                },
            )
        logger.info(
            "Quiz attempt created.",
            extra={
                "attempt_id": attempt.id,
                "course_id": course_id,
                "activity_id": activity_id,
                "user_id": request.user.id,
            },
        )

        return Response(
            {
                "attempt_id": attempt.id,
                "resume": False,
                "started_at": attempt.started_at,
                "time_limit": int(activity.quiz_time_limit_seconds or 600),
                "max_attempts": max_attempts,
                "total_points": float(attempt.total_points or 0),
                "anti_cheat": _anti_cheat_runtime_config(activity),
                "requires_consent": requires_consent,
            },
            status=201,
        )
    except Exception as exc:
        logger.exception(
            "Quiz start failed.",
            extra={"course_id": course_id, "activity_id": activity_id, "user_id": getattr(request.user, "id", None)},
        )
        return Response(
            {"error": "Failed to start quiz"},
            status=500,
        )


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def quiz_submit(request, course_id, activity_id):
    course = _get_course_with_access(course_id, request.user)
    if not course:
        return Response({"error": "Course not found or access denied"}, status=404)

    if request.user.role != "student":
        return Response({"error": "Only students can submit quizzes."}, status=403)

    activity = get_object_or_404(CourseActivity, id=activity_id, course=course)
    if not _is_quiz_activity(activity):
        return Response({"error": "This activity is not a quiz"}, status=400)
    if str(activity.publish_state or "").lower() != CourseActivity.PUBLISH_STATE_PUBLISHED:
        return Response({"error": "This assessment is not currently available."}, status=400)

    attempt_id = request.data.get("attempt_id")
    if not attempt_id:
        return Response({"error": "attempt_id is required"}, status=400)

    try:
        answers = _normalize_submitted_answers(request.data.get("answers"))
    except ValueError as exc:
        return Response({"error": str(exc)}, status=400)

    with transaction.atomic():
        attempt = get_object_or_404(
            QuizAttempt.objects.select_for_update(),
            id=attempt_id,
            quiz=activity,
            student=request.user,
        )
        # Concurrency guard: submit endpoint is final authority once this row lock is held.
        if attempt.is_locked:
            if attempt.submitted_at is None:
                now_fix = timezone.now()
                attempt.submitted_at = now_fix
                attempt.last_activity_at = now_fix
                attempt.save(update_fields=["submitted_at", "last_activity_at"])
            return Response(
                {
                    "error": "This attempt was force-submitted due to security violations.",
                    "attempt_locked": True,
                    "attempt_id": attempt.id,
                    "score": float(attempt.score or 0),
                    "total_points": float(attempt.total_points or 0),
                    "submitted_at": attempt.submitted_at,
                },
                status=409,
            )
        if int(attempt.suspicious_events or 0) >= QUIZ_SECURITY_FORCE_SUBMIT_THRESHOLD:
            # Critical race hardening: never grade if threshold is reached at submit time.
            _force_submit_locked_attempt(attempt, activity, reason="threshold_reached_at_submit")
            logger.warning(
                "Submit blocked by threshold and attempt force-submitted.",
                extra={
                    "attempt_id": attempt.id,
                    "quiz_id": activity.id,
                    "user_id": request.user.id,
                    "violation_count": int(attempt.suspicious_events or 0),
                },
            )
            return Response(
                {
                    "error": "Attempt was force-submitted due to security violations.",
                    "attempt_locked": True,
                    "attempt_id": attempt.id,
                    "score": float(attempt.score or 0),
                    "total_points": float(attempt.total_points or 0),
                    "submitted_at": attempt.submitted_at,
                },
                status=409,
            )
        if attempt.submitted_at:
            return Response({"error": "This attempt was already submitted."}, status=409)
        if attempt.started_at is None:
            return Response({"error": "Attempt is invalid."}, status=400)
        if not QuizAttemptAcknowledgement.objects.filter(attempt=attempt, student=request.user).exists():
            return Response(
                {"error": "Attempt is missing required pre-exam acknowledgment."},
                status=409,
            )
        submitted_attempts_count = QuizAttempt.objects.filter(
            quiz=activity,
            student=request.user,
            submitted_at__isnull=False,
        ).exclude(id=attempt.id).count()
        if submitted_attempts_count >= int(activity.max_attempts or 3):
            return Response({"error": "Maximum quiz attempts reached"}, status=400)

        source_questions = (
            attempt.question_snapshot
            if isinstance(attempt.question_snapshot, list) and attempt.question_snapshot
            else _load_quiz_questions_for_runtime(activity)
        )
        questions = _validate_and_normalize_quiz_questions(source_questions)
        if not questions:
            return Response({"error": "Quiz has no questions configured"}, status=400)

        now = timezone.now()
        elapsed_seconds = int(max((now - attempt.started_at).total_seconds(), 0))
        time_limit = int(activity.quiz_time_limit_seconds or 600)
        if elapsed_seconds > time_limit:
            return Response(
                {
                    "error": "Quiz time limit exceeded.",
                    "time_spent": elapsed_seconds,
                    "time_limit": time_limit,
                },
                status=400,
            )
        if attempt.last_activity_at:
            inactivity_gap = int(max((now - attempt.last_activity_at).total_seconds(), 0))
            if inactivity_gap > QUIZ_SECURITY_INACTIVITY_SECONDS:
                attempt.suspicious_events = int(attempt.suspicious_events or 0) + 1
                attempt.last_activity_at = now
                attempt.save(update_fields=["suspicious_events", "last_activity_at"])
                if int(attempt.suspicious_events or 0) >= QUIZ_SECURITY_FORCE_SUBMIT_THRESHOLD:
                    _force_submit_locked_attempt(attempt, activity, reason="inactivity_threshold")
                    logger.warning(
                        "Attempt force-submitted due to inactivity threshold.",
                        extra={
                            "attempt_id": attempt.id,
                            "quiz_id": activity.id,
                            "user_id": request.user.id,
                            "violation_count": int(attempt.suspicious_events or 0),
                        },
                    )
                    return Response(
                        {
                            "error": "Attempt was force-submitted due to security violations.",
                            "attempt_locked": True,
                            "attempt_id": attempt.id,
                            "score": float(attempt.score or 0),
                            "total_points": float(attempt.total_points or 0),
                            "submitted_at": attempt.submitted_at,
                        },
                        status=409,
                    )

        try:
            grading = _grade_quiz_questions(
                questions,
                answers,
                frozen_total_points=float(attempt.total_points or 0) if float(attempt.total_points or 0) > 0 else None,
            )
        except ValueError as exc:
            return Response({"error": str(exc)}, status=400)

        attempt.answers = answers
        attempt.score = grading["score"]
        attempt.total_points = grading["total_points"]
        attempt.result_breakdown = grading["breakdown"]
        attempt.correct_answers = grading["correct_answers"]
        attempt.incorrect_answers = grading["incorrect_answers"]
        attempt.status = grading["status"]
        attempt.time_spent = elapsed_seconds
        attempt.last_activity_at = now
        attempt.submitted_at = now
        attempt.graded_at = now if grading["status"] == QuizAttempt.STATUS_GRADED else None
        attempt.graded_by = None
        attempt.save()
        _sync_attempt_answer_records(attempt, questions, answers, grading)

    feedback = f"Correct answers: {grading['correct_answers']} | Incorrect answers: {grading['incorrect_answers']}"
    if grading["status"] == QuizAttempt.STATUS_PENDING_REVIEW:
        feedback = "Submission received and pending manual review."
    _update_attempt_submission_record(attempt, feedback_text=feedback)
    dispatch_event("quiz_completed", attempt=attempt, actor=request.user)

    attempt.refresh_from_db()
    if attempt.is_locked:
        return Response(
            {
                "error": "This attempt was force-submitted due to security violations.",
                "attempt_locked": True,
                "attempt_id": attempt.id,
                "score": float(attempt.score or 0),
                "total_points": float(attempt.total_points or 0),
                "submitted_at": attempt.submitted_at,
            },
            status=409,
        )

    attempts_qs = QuizAttempt.objects.filter(quiz=activity, student=request.user).order_by("-started_at")
    attempts_data = QuizAttemptSerializer(attempts_qs, many=True).data
    attempts_data = [_sanitize_attempt_for_student_visibility(item, activity) for item in attempts_data]

    visibility = _resolve_attempt_visibility(activity, attempt=attempt)
    show_score_immediately = visibility["show_score_immediately"] and attempt.status == QuizAttempt.STATUS_GRADED
    allow_answer_review = visibility["allow_answer_review"] and attempt.status == QuizAttempt.STATUS_GRADED
    response_payload = {
        "attempt_id": attempt.id,
        "show_score_immediately": show_score_immediately,
        "allow_answer_review": allow_answer_review,
        "status": attempt.status,
        "pending_manual_review": attempt.status == QuizAttempt.STATUS_PENDING_REVIEW,
        "time_spent": elapsed_seconds,
        "attempts": attempts_data,
        "max_attempts": int(activity.max_attempts or 3),
    }
    if show_score_immediately:
        response_payload["score"] = float(attempt.score or 0)
        response_payload["total_points"] = float(attempt.total_points or 0)
        response_payload["correct_answers"] = int(attempt.correct_answers or 0)
        response_payload["incorrect_answers"] = int(attempt.incorrect_answers or 0)
        response_payload["breakdown"] = attempt.result_breakdown or []
        if not allow_answer_review:
            response_payload["review_message"] = "Detailed review is not available."
    else:
        response_payload["message"] = (
            "Your exam has been submitted and is pending instructor review."
            if attempt.status == QuizAttempt.STATUS_PENDING_REVIEW
            else "Your exam has been submitted. Results will be released by your instructor."
        )
    if allow_answer_review:
        response_payload["review_url"] = f"/api/courses/{course_id}/activities/{activity_id}/quiz/review/"

    return Response(response_payload, status=200)


@api_view(["GET", "POST"])
@permission_classes([IsAuthenticated])
def attendance_sessions(request, course_id):
    course = _get_course_with_access(course_id, request.user)
    if not course:
        return Response({"error": "Course not found or access denied"}, status=404)

    if request.method == "GET":
        sessions = AttendanceSession.objects.filter(course=course).select_related("created_by").prefetch_related("records__student", "records__marked_by")
        serializer = AttendanceSessionSerializer(sessions, many=True, context={"request": request})
        return Response(serializer.data)

    if course.instructor_id != request.user.id:
        return Response({"error": "Only instructor can create attendance sessions"}, status=403)

    date_raw = request.data.get("date")
    topic = (request.data.get("topic") or "").strip()
    if not date_raw:
        return Response({"error": "date is required (YYYY-MM-DD)"}, status=400)
    if not topic:
        return Response({"error": "topic is required"}, status=400)

    try:
        parsed_date = datetime.strptime(str(date_raw), "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return Response({"error": "Invalid date format. Use YYYY-MM-DD."}, status=400)

    session = AttendanceSession.objects.create(
        course=course,
        date=parsed_date,
        topic=topic,
        created_by=request.user,
    )
    return Response(
        AttendanceSessionSerializer(session, context={"request": request}).data,
        status=201,
    )


@api_view(["GET", "DELETE"])
@permission_classes([IsAuthenticated])
def attendance_session_detail(request, course_id, session_id):
    course = _get_course_with_access(course_id, request.user)
    if not course:
        return Response({"error": "Course not found or access denied"}, status=404)

    session = get_object_or_404(
        AttendanceSession.objects.select_related("created_by").prefetch_related("records__student", "records__marked_by"),
        id=session_id,
        course=course,
    )
    if request.method == "GET":
        serializer = AttendanceSessionSerializer(session, context={"request": request})
        return Response(serializer.data)

    if course.instructor_id != request.user.id:
        return Response({"error": "Only instructor can delete attendance sessions"}, status=403)

    session.delete()
    return Response({"message": "Attendance session deleted."}, status=200)


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def attendance_records(request, course_id, session_id):
    course = _get_course_with_access(course_id, request.user)
    if not course:
        return Response({"error": "Course not found or access denied"}, status=404)

    session = get_object_or_404(AttendanceSession, id=session_id, course=course)

    if request.user.role == "student":
        status_value = request.data.get("status", AttendanceRecord.STATUS_PRESENT)
        if status_value not in dict(AttendanceRecord.STATUS_CHOICES):
            return Response({"error": "Invalid attendance status"}, status=400)

        record, _ = AttendanceRecord.objects.update_or_create(
            session=session,
            student=request.user,
            defaults={
                "status": status_value,
                "marked_by": request.user,
            },
        )
        return Response(AttendanceRecordSerializer(record).data, status=200)

    if course.instructor_id != request.user.id:
        return Response({"error": "Only instructor can mark attendance"}, status=403)

    allowed_status = dict(AttendanceRecord.STATUS_CHOICES)
    bulk_rows = request.data.get("records")
    if isinstance(bulk_rows, list):
        if not bulk_rows:
            return Response({"error": "records must not be empty"}, status=400)

        # Bulk save path used by the instructor attendance modal to avoid request-per-student failures.
        saved_records = []
        with transaction.atomic():
            for index, row in enumerate(bulk_rows):
                if not isinstance(row, dict):
                    return Response({"error": f"records[{index}] must be an object"}, status=400)
                student_id = row.get("student_id")
                if not student_id:
                    return Response({"error": f"records[{index}].student_id is required"}, status=400)

                student = get_object_or_404(User, id=student_id, role="student")
                if not course.students.filter(id=student.id).exists():
                    return Response({"error": f"Student {student.id} is not enrolled in this course"}, status=400)

                status_value = row.get("status", AttendanceRecord.STATUS_PRESENT)
                if status_value not in allowed_status:
                    return Response({"error": f"records[{index}].status is invalid"}, status=400)

                points_value = row.get("points_earned", 0)
                try:
                    points_value = float(points_value)
                except (TypeError, ValueError):
                    return Response({"error": f"records[{index}].points_earned must be numeric"}, status=400)

                record, _ = AttendanceRecord.objects.update_or_create(
                    session=session,
                    student=student,
                    defaults={
                        "status": status_value,
                        "points_earned": points_value,
                        "marked_by": request.user,
                    },
                )
                saved_records.append(record)

        dispatch_event("attendance_marked", course=course, session=session, actor=request.user, records=saved_records)

        serializer = AttendanceRecordSerializer(saved_records, many=True)
        return Response({"saved": len(saved_records), "records": serializer.data}, status=200)

    student_id = request.data.get("student_id")
    if not student_id:
        return Response({"error": "student_id is required"}, status=400)

    student = get_object_or_404(User, id=student_id, role="student")
    if not course.students.filter(id=student.id).exists():
        return Response({"error": "Student is not enrolled in this course"}, status=400)

    status_value = request.data.get("status", AttendanceRecord.STATUS_PRESENT)
    if status_value not in allowed_status:
        return Response({"error": "Invalid attendance status"}, status=400)

    points_value = request.data.get("points_earned", 0)
    try:
        points_value = float(points_value)
    except (TypeError, ValueError):
        return Response({"error": "points_earned must be numeric"}, status=400)

    record, _ = AttendanceRecord.objects.update_or_create(
        session=session,
        student=student,
        defaults={
            "status": status_value,
            "points_earned": points_value,
            "marked_by": request.user,
        },
    )
    dispatch_event("attendance_marked", course=course, session=session, actor=request.user, records=[record])
    return Response(AttendanceRecordSerializer(record).data, status=200)


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def attendance_create_session(request):
    course_id = request.data.get("course_id")
    if not course_id:
        return Response({"error": "course_id is required"}, status=400)
    course = _get_course_with_access(course_id, request.user)
    if not course:
        return Response({"error": "Course not found or access denied"}, status=404)
    if course.instructor_id != request.user.id:
        return Response({"error": "Only instructor can create attendance sessions"}, status=403)

    date_raw = request.data.get("date")
    topic = (request.data.get("topic") or "").strip() or "Attendance Session"
    if not date_raw:
        return Response({"error": "date is required (YYYY-MM-DD)"}, status=400)

    try:
        parsed_date = datetime.strptime(str(date_raw), "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return Response({"error": "Invalid date format. Use YYYY-MM-DD."}, status=400)

    session = AttendanceSession.objects.create(
        course=course,
        date=parsed_date,
        topic=topic,
        created_by=request.user,
    )
    return Response(AttendanceSessionSerializer(session, context={"request": request}).data, status=201)


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def attendance_by_course(request, course_id):
    return attendance_sessions(request, course_id)


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def attendance_record_by_session(request, session_id):
    session = get_object_or_404(AttendanceSession, id=session_id)
    return attendance_records(request, session.course_id, session_id)


@api_view(["GET", "POST"])
@permission_classes([IsAuthenticated])
def course_meetings(request, course_id):
    if request.method == "GET":
        try:
            meetings = list_course_meetings(course_id=course_id, user=request.user)
        except PermissionError:
            return Response({"error": "Course not found or access denied"}, status=404)

        serializer = MeetingSerializer(meetings, many=True, context={"request": request})
        return Response(serializer.data, status=200)

    serializer = MeetingSerializer(data=request.data, context={"request": request})
    serializer.is_valid(raise_exception=True)

    try:
        meeting = create_meeting(
            course_id=course_id,
            title=serializer.validated_data["title"],
            scheduled_time=serializer.validated_data["scheduled_time"],
            meeting_link=serializer.validated_data["meeting_link"],
            created_by=request.user,
        )
    except PermissionError:
        return Response({"error": "Only the course instructor can create meetings"}, status=403)

    return Response(MeetingSerializer(meeting, context={"request": request}).data, status=201)


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def meeting_join(request, meeting_id):
    try:
        meeting, attendance = join_meeting(meeting_id=meeting_id, student=request.user)
    except PermissionError:
        return Response({"error": "Only enrolled students can join meetings"}, status=403)

    return Response(
        {
            "meeting": MeetingSerializer(meeting, context={"request": request}).data,
            "meeting_link": meeting.meeting_link,
            "joined_at": attendance.joined_at,
        },
        status=200,
    )


@api_view(["GET", "POST"])
@permission_classes([IsAuthenticated])
def activity_attendance(request, course_id, activity_id):
    """
    Compatibility endpoint for existing frontend:
    /api/courses/{course_id}/activities/{activity_id}/attendance/
    """
    course = _get_course_with_access(course_id, request.user)
    if not course:
        return Response({"error": "Course not found or access denied"}, status=404)

    activity = get_object_or_404(CourseActivity, id=activity_id, course=course)
    activity_type_name = (activity.activity_type.name or "").lower()
    if activity_type_name != "attendance":
        return Response({"error": "This activity is not an attendance activity"}, status=400)

    if request.method == "GET":
        sessions = AttendanceSession.objects.filter(course=course).select_related("created_by").prefetch_related("records__student", "records__marked_by")
        serializer = AttendanceSessionSerializer(sessions, many=True, context={"request": request})
        return Response(serializer.data)

    session_id = request.data.get("session_id")
    if session_id:
        session = get_object_or_404(AttendanceSession, id=session_id, course=course)
    else:
        session, _ = AttendanceSession.objects.get_or_create(
            course=course,
            date=timezone.now().date(),
            topic=(activity.topic or activity.title or "Attendance"),
            defaults={"created_by": course.instructor},
        )

    if request.user.role == "student":
        status_value = request.data.get("status", AttendanceRecord.STATUS_PRESENT)
        if status_value not in dict(AttendanceRecord.STATUS_CHOICES):
            return Response({"error": "Invalid attendance status"}, status=400)
        record, _ = AttendanceRecord.objects.update_or_create(
            session=session,
            student=request.user,
            defaults={
                "status": status_value,
                "marked_by": request.user,
            },
        )
        return Response(
            {
                "session": AttendanceSessionSerializer(session, context={"request": request}).data,
                "record": AttendanceRecordSerializer(record).data,
            },
            status=200,
        )

    if course.instructor_id != request.user.id:
        return Response({"error": "Only instructor can mark attendance"}, status=403)

    student_id = request.data.get("student_id")
    if not student_id:
        return Response({"error": "student_id is required for instructor attendance marking"}, status=400)

    student = get_object_or_404(User, id=student_id, role="student")
    if not course.students.filter(id=student.id).exists():
        return Response({"error": "Student is not enrolled in this course"}, status=400)

    status_value = request.data.get("status", AttendanceRecord.STATUS_PRESENT)
    if status_value not in dict(AttendanceRecord.STATUS_CHOICES):
        return Response({"error": "Invalid attendance status"}, status=400)

    points_value = request.data.get("points_earned", 0)
    try:
        points_value = float(points_value)
    except (TypeError, ValueError):
        return Response({"error": "points_earned must be numeric"}, status=400)

    record, _ = AttendanceRecord.objects.update_or_create(
        session=session,
        student=student,
        defaults={
            "status": status_value,
            "points_earned": points_value,
            "marked_by": request.user,
        },
    )

    return Response(
        {
            "session": AttendanceSessionSerializer(session, context={"request": request}).data,
            "record": AttendanceRecordSerializer(record).data,
        },
        status=200,
    )

@api_view(["GET", "POST"])
@permission_classes([IsAuthenticated])
def activity_comments(request, activity_id):
    try:
        activity = CourseActivity.objects.get(id=activity_id)
    except CourseActivity.DoesNotExist:
        return Response({"error": "Activity not found"}, status=404)

    if request.method == "GET":
        comments = activity.comments.all().order_by("created_at")
        serializer = ActivityCommentSerializer(comments, many=True, context={"request": request})
        return Response(serializer.data)

    elif request.method == "POST":
        data = request.data.copy()
        data["activity"] = activity.id
        data["user"] = request.user.id
        serializer = ActivityCommentSerializer(data=data, context={"request": request})
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data, status=201)
        return Response(serializer.errors, status=400)
    

class ActivitySubmissionViewSet(viewsets.ModelViewSet):
    queryset = ActivitySubmission.objects.all()
    serializer_class = ActivitySubmissionSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        user = self.request.user
        if user.role == 'instructor':
            # Instructor sees submissions only for their courses
            return ActivitySubmission.objects.filter(activity__course__instructor=user)
        return ActivitySubmission.objects.filter(student=user)
    

@api_view(["GET", "POST"])
@permission_classes([IsAuthenticated])
@parser_classes([MultiPartParser, FormParser])
def course_announcements(request, course_id):
    if request.method == "GET":
        announcements = CourseActivity.objects.filter(
            course_id=course_id,
            activity_type__name="announcement"  # filter by ActivityType name
        ).order_by("-created_at")
        serializer = CourseActivitySerializer(announcements, many=True, context={"request": request})
        return Response(serializer.data)

    elif request.method == "POST":
        data = request.data.copy()
        data["course"] = course_id

        # 🔹 Get the ActivityType object for "announcement"
        try:
            announcement_type = ActivityType.objects.get(name="announcement")
        except ActivityType.DoesNotExist:
            return Response({"error": "ActivityType 'announcement' does not exist."}, status=400)

        data["activity_type"] = announcement_type.id  # ✅ use ID, not string

        serializer = CourseActivitySerializer(data=data, context={"request": request})
        if serializer.is_valid():
            announcement = serializer.save()

            # Handle multiple files
            files = request.FILES.getlist("files")
            for f in files:
                SubmissionAttachment.objects.create(
                    submission=None,
                    announcement=announcement,
                    file=f
                )

            dispatch_event("announcement_created", announcement=announcement, actor=request.user)

            return Response(serializer.data, status=201)
        return Response(serializer.errors, status=400)


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def activity_types(request):
    allowed_names = {"assignment", "quiz", "material", "project"}
    types = (
        ActivityType.objects.filter(name__iregex=r"^(assignment|quiz|material|project)$")
        .order_by("name")
        .values("id", "name")
    )
    return Response(types)


class CourseViewSet(viewsets.ModelViewSet):
    queryset = Course.objects.all()
    serializer_class = CourseSerializer

    @action(detail=True, methods=['get', 'post'])
    def announcements(self, request, pk=None):
        course = self.get_object()
        if request.method == "GET":
            announcements = CourseActivity.objects.filter(course=course, activity_type="announcement")
            serializer = CourseActivitySerializer(announcements, many=True, context={"request": request})
            return Response(serializer.data)
        
@api_view(["POST"])
@permission_classes([IsAuthenticated])
def join_course(request):
    code = str(request.data.get("code", "") or "").strip().upper()

    if not code:
        return Response({"error": "Join code required"}, status=400)

    try:
        course = Course.objects.get(join_code=code, join_code_enabled=True)
    except Course.DoesNotExist:
        return Response({"error": "Invalid or disabled code"}, status=404)

    if request.user.role != "student":
        return Response({"error": "Only students can join courses"}, status=403)

    if course.students.filter(id=request.user.id).exists():
        return Response({"message": "You are already enrolled in this course."}, status=200)

    if EnrollmentRequest.objects.filter(
        course=course,
        student=request.user,
        status=EnrollmentRequest.STATUS_PENDING,
    ).exists():
        return Response({"message": "Enrollment request already pending approval."}, status=200)

    enrollment_request = EnrollmentRequest.objects.create(
        course=course,
        student=request.user,
        status=EnrollmentRequest.STATUS_PENDING,
    )
    dispatch_event("enrollment_request_created", enrollment_request=enrollment_request, actor=request.user)

    serializer = EnrollmentRequestSerializer(enrollment_request, context={"request": request})
    return Response(
        {
            "message": "Enrollment request sent. Please wait for instructor approval.",
            "request": serializer.data,
        },
        status=201,
    )


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def enrollment_requests_list(request):
    if not _allow(INSTRUCTOR_ROLE_PERMISSION, request.user):
        return Response({"error": "Only instructors can view enrollment requests."}, status=403)

    queryset = _pending_enrollment_requests_for_user(request.user)
    course_id = request.query_params.get("course_id")
    if course_id not in (None, ""):
        queryset = queryset.filter(course_id=course_id)

    serializer = EnrollmentRequestSerializer(queryset.order_by("-created_at"), many=True, context={"request": request})
    return Response(serializer.data)


def _review_enrollment_request(request, enrollment_request_id, next_status):
    enrollment_request = get_object_or_404(
        EnrollmentRequest.objects.select_related("course", "student", "reviewed_by"),
        id=enrollment_request_id,
    )
    course = enrollment_request.course
    if not _can_manage_course(request.user, course):
        return Response({"error": "Course not found or access denied"}, status=404)

    if enrollment_request.status != EnrollmentRequest.STATUS_PENDING:
        return Response({"error": "This enrollment request has already been reviewed."}, status=400)

    with transaction.atomic():
        enrollment_request.status = next_status
        enrollment_request.reviewed_at = timezone.now()
        enrollment_request.reviewed_by = request.user
        enrollment_request.save(update_fields=["status", "reviewed_at", "reviewed_by", "updated_at"])

        if next_status == EnrollmentRequest.STATUS_APPROVED:
            course.students.add(enrollment_request.student)
            dispatch_event("student_joined_course", course=course, student=enrollment_request.student, actor=request.user)

    serializer = EnrollmentRequestSerializer(enrollment_request, context={"request": request})
    return Response(serializer.data, status=200)


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def approve_enrollment_request(request, enrollment_request_id):
    if not _allow(INSTRUCTOR_ROLE_PERMISSION, request.user):
        return Response({"error": "Only instructors can approve enrollment requests."}, status=403)
    return _review_enrollment_request(request, enrollment_request_id, EnrollmentRequest.STATUS_APPROVED)


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def reject_enrollment_request(request, enrollment_request_id):
    if not _allow(INSTRUCTOR_ROLE_PERMISSION, request.user):
        return Response({"error": "Only instructors can reject enrollment requests."}, status=403)
    return _review_enrollment_request(request, enrollment_request_id, EnrollmentRequest.STATUS_REJECTED)

@api_view(['PUT'])
@permission_classes([IsAuthenticated])
def toggle_join_code(request, course_id):
    course = Course.objects.get(id=course_id)
    if not _can_manage_course(request.user, course):
        return Response({"error": "Unauthorized"}, status=403)
    
    enabled = request.data.get("enabled", True)
    course.join_code_enabled = enabled
    course.save()
    return Response({"join_code_enabled": course.join_code_enabled})


@api_view(['PUT'])
@permission_classes([IsAuthenticated])
def toggle_archive(request, course_id):
    course = Course.objects.get(id=course_id)
    if not _can_manage_course(request.user, course):
        return Response({"error": "Unauthorized"}, status=403)

    archived = bool(request.data.get("archived", True))
    course.is_archived = archived
    course.save(update_fields=["is_archived"])
    if _is_admin_user(request.user):
        _log_admin_action(
            "Course archived" if course.is_archived else "Course restored",
            performed_by=request.user,
            target_user=course.instructor,
            description=f"Set course '{course.title}' archive state to {course.is_archived}.",
        )

    return Response(
        {
            "id": course.id,
            "is_archived": course.is_archived,
            "status": "archived" if course.is_archived else "active",
            "state": "archived" if course.is_archived else "active",
        }
    )

@api_view(['POST'])
@permission_classes([IsAuthenticated])
def regenerate_join_code(request, course_id):
    course = Course.objects.get(id=course_id)
    if not _can_manage_course(request.user, course):
        return Response({"error": "Unauthorized"}, status=403)
    
    course.join_code = secrets.token_hex(4).upper()  # 8-char code
    course.join_code_expiration = datetime.now() + timedelta(days=7)  # optional default 7 days
    course.save()
    return Response({
        "join_code": course.join_code,
        "join_code_expiration": course.join_code_expiration
    })

class CourseActivityViewSet(viewsets.ModelViewSet):
    queryset = CourseActivity.objects.all()
    serializer_class = CourseActivitySerializer

    def get_serializer_context(self):
        context = super().get_serializer_context()
        context["request"] = self.request
        return context
